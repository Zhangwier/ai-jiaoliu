# -*- coding: utf-8 -*-
"""一键生成检查补充材料（只读原文件；全部产物写入 _检查补充_20261017）。
运行：python 复现_检查补充.py
"""
import os, re, glob, csv, json, zipfile, subprocess, pathlib
from collections import defaultdict

ROOT = r"D:\MG\program\ai交流"
OUT = os.path.join(ROOT, "_检查补充_20261017")
os.makedirs(OUT, exist_ok=True)
DRAW = os.path.join(ROOT, "AI对比", "ep1 识图")
PPTX = os.path.join(ROOT, "AI交流分享-修改版.pptx")

# ══════════ 1. 鹈鹕 10 份：逐字节一致性 ══════════
def md5(p):
    import hashlib
    return hashlib.md5(open(p, "rb").read()).hexdigest()

blind = {}
for code in "ABCDEFGHIJ":
    f = os.path.join(ROOT, "完善版01配套", "鹈鹕原始结果", f"{code}.html")
    if os.path.exists(f):
        blind[code] = md5(f)
mapping = json.load(open(os.path.join(ROOT, "完善版01配套", "鹈鹕配置映射.json"), encoding="utf-8"))
rows = []
for item in mapping:
    d = item["configuration"]
    hits = glob.glob(os.path.join(DRAW, d, "*.html"))
    src_hash = md5(hits[0]) if hits else ""
    rows.append({
        "编号": item["code"], "配置目录": d,
        "原目录HTML": os.path.basename(hits[0]) if hits else "缺失",
        "与盲测副本一致": src_hash == blind.get(item["code"], "x"),
        "映射source路径可用": os.path.exists(item["source"]),
        "KB": round(os.path.getsize(hits[0]) / 1024, 1) if hits else 0,
    })
with open(os.path.join(OUT, "补充_鹈鹕10份逐字节一致性.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
print("鹈鹕一致性:", sum(r["与盲测副本一致"] for r in rows), "/10 一致；路径可用:",
      sum(r["映射source路径可用"] for r in rows), "/10")

# ══════════ 2. 修编变更（源文件独立复算，两级口径） ══════════
import openpyxl
SRC = os.path.join(DRAW, "任务3源文件", "分项工程量清单.xlsx")

def num(v):
    try:
        return float(v)
    except Exception:
        return None

wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)

# 2a 清单子目级（工程量清单 sheet）
ws = wb["工程量清单"]
rows = list(ws.iter_rows(values_only=True))
sub = [r for r in rows[3:] if r[0] and "-" in str(r[0])]
items = []
for r in sub:
    amt, qty = num(r[14]), num(r[12])
    if amt in (None, 0) and qty in (None, 0):
        continue
    items.append(dict(层级="清单子目", 编码=r[0], 名称=(r[1] or "").strip(), 单位=r[2] or "",
                      原数量=r[3], 修编数量=r[9], 增减数量=qty, 增减金额=amt))
# 2b 分项级（分项清单 sheet）
ws2 = wb["分项清单"]
rows2 = list(ws2.iter_rows(values_only=True))
sub2 = [r for r in rows2[4:] if r[1] not in (None, "")]
for r in sub2:
    amt, qty = num(r[15]), num(r[13])
    if amt in (None, 0) and qty in (None, 0):
        continue
    items.append(dict(层级="分项", 编码=r[1], 名称=(r[2] or "").strip(), 单位=r[3] or "",
                      原数量=r[4], 修编数量=r[10], 增减数量=qty, 增减金额=amt))
with open(os.path.join(OUT, "补充_修编变更全量明细.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(items[0].keys())); w.writeheader(); w.writerows(items)

lvl = defaultdict(lambda: [0, 0, 0, 0.0])
for i in items:
    k = i["层级"]
    a = i["增减金额"] or 0
    lvl[k][0] += 1
    lvl[k][1] += 1 if a > 0 else 0
    lvl[k][2] += 1 if a < 0 else 0
    lvl[k][3] += a
with open(os.path.join(OUT, "补充_修编变更_两级口径.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f); w.writerow(["层级", "变动行数", "核增", "核减", "增减金额合计(元)"])
    for k, v in lvl.items():
        w.writerow([k, v[0], v[1], v[2], round(v[3], 2)])

# 章节汇总（清单子目级）
chap = defaultdict(lambda: [0.0, 0, 0])
for i in items:
    if i["层级"] != "清单子目":
        continue
    ch = str(i["编码"]).split("-")[0]
    a = i["增减金额"] or 0
    chap[ch][0] += a
    chap[ch][1] += 1 if a > 0 else 0
    chap[ch][2] += 1 if a < 0 else 0
with open(os.path.join(OUT, "补充_修编变更_按章汇总.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f); w.writerow(["章", "增减金额(元)", "核增项", "核减项"])
    for ch in sorted(chap):
        v = chap[ch]; w.writerow([ch, round(v[0], 2), v[1], v[2]])

sm = [[c for c in r if c is not None] for r in wb["汇总表"].iter_rows(values_only=True)]
wb.close()
with open(os.path.join(OUT, "补充_修编汇总表_源文件.json"), "w", encoding="utf-8") as f:
    json.dump(sm, f, ensure_ascii=False, indent=1)
print("两级口径:", {k: (v[0], round(v[3], 2)) for k, v in lvl.items()})

# ══════════ 3. 图纸审查 10 配置交付形态 ══════════
rows = []
for d in sorted(os.listdir(DRAW)):
    p = os.path.join(DRAW, d)
    if not os.path.isdir(p) or "源文件" in d or d.startswith("提示词3"):
        continue
    ans = os.path.join(p, "回答.txt")
    txt = open(ans, encoding="utf-8", errors="ignore").read()
    first = txt.splitlines()[0].strip() if txt else ""
    xls = glob.glob(os.path.join(p, "*.xlsx"))
    sheets = nrows = nform = 0
    for f in xls:
        z = zipfile.ZipFile(f)
        for n in z.namelist():
            if re.match(r"xl/worksheets/sheet\d+\.xml$", n):
                c = z.read(n).decode("utf8", "ignore")
                sheets += 1
                nrows += len(re.findall(r"<(?:x:)?row[ >]", c))
                nform += len(re.findall(r"<(?:x:)?f[ >/]", c))
        z.close()
    rows.append({
        "配置目录": d, "记录耗时": re.sub(r"^提示词1\s*", "", first).strip(),
        "回答正文字数": len(txt), "工作簿": os.path.basename(xls[0]) if xls else "无",
        "工作表数": sheets, "表内总行数": nrows, "真实公式单元格": nform,
        "自述提到公式/条件格式": ("公式" in txt or "条件格式" in txt),
        "自述视觉未启用": ("不支持视觉" in txt or "未启用视觉" in txt),
        "引出图签TJ5标段": ("TJ5" in txt),
    })
rows.sort(key=lambda r: r["配置目录"])
with open(os.path.join(OUT, "补充_图纸审查10配置交付形态.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
print("交付形态:", len(rows), "行")

# ══════════ 4. 鹈鹕技术栈 ══════════
def main_vb(s):
    best, area = "", 0
    for m in re.finditer(r'<svg[^>]*viewBox="([^"]+)"', s):
        n = [float(x) for x in re.findall(r"-?\d+(?:\.\d+)?", m.group(1))]
        if len(n) == 4 and n[2] * n[3] > area:
            area, best = n[2] * n[3], m.group(1)
    return best

pel = []
for d in sorted(os.listdir(DRAW)):
    p = os.path.join(DRAW, d)
    if not os.path.isdir(p):
        continue
    for h in glob.glob(os.path.join(p, "*.html")):
        s = open(h, encoding="utf-8", errors="ignore").read()
        smil = len(re.findall(r"<animate(?:Transform|Motion)?\b", s))
        kf = len(re.findall(r"@keyframes", s))
        raf = len(re.findall(r"requestAnimationFrame", s))
        stack = "SMIL" if smil and not kf and not raf else ("CSS 关键帧" if kf and not smil else ("JS + rAF" if raf else "混合/静态"))
        pel.append({
            "配置目录": d, "文件": os.path.basename(h), "KB": round(len(s.encode()) / 1024, 1),
            "主画布viewBox": main_vb(s), "动画技术栈": stack,
            "SMIL元素数": smil, "CSS关键帧数": kf, "rAF调用": raf,
            "外部网络依赖": bool(re.search(r'<(?:script|link|img)[^>]+(?:src|href)="https?://', s)),
        })
with open(os.path.join(OUT, "补充_鹈鹕10结果技术栈.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(pel[0].keys())); w.writeheader(); w.writerows(pel)

# ══════════ 5. 设计稿 vs 实际 PPT ══════════
design = json.load(open(os.path.join(ROOT, "逐页内容设计v3.json"), encoding="utf-8"))
z = zipfile.ZipFile(PPTX)
pres = z.read("ppt/presentation.xml").decode("utf8")
rels = z.read("ppt/_rels/presentation.xml.rels").decode("utf8")
rmap = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
slds = re.findall(r'<p:sldId[^>]*r:id="([^"]+)"', pres)
actual = []
for i, rid in enumerate(slds, 1):
    x = z.read("ppt/" + rmap[rid]).decode("utf8")
    actual.append((i, rmap[rid].split("/")[-1], "".join(re.findall(r"<a:t>(.*?)</a:t>", x))))

mrows = []
for d in design:
    key = re.sub(r"\s+", "", d["title"])[:8]
    hit = [a for a in actual if key and key in re.sub(r"\s+", "", a[2])]
    mrows.append(dict(设计页=d["page"], 来源=d["source"], 设计标题=d["title"], 版式=d["layout"],
                      实际页=hit[0][0] if hit else "未落地", 实际文件=hit[0][1] if hit else "-",
                      结论="" if hit else "设计稿有、PPT 无"))
with open(os.path.join(OUT, "补充_设计稿57页与实际PPT54页对照.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(mrows[0].keys())); w.writeheader(); w.writerows(mrows)
missing = [m for m in mrows if m["实际页"] == "未落地"]
print("设计 57 页 / 实际 54 页；未落地:", [(m["设计页"], m["设计标题"]) for m in missing])

# ══════════ 6. PPT 体检 ══════════
import collections
sz = collections.Counter()
per = []
for i, rid in enumerate(slds, 1):
    x = z.read("ppt/" + rmap[rid]).decode("utf8")
    sizes = [int(v) / 100 for v in re.findall(r'sz="(\d+)"', x)]
    per.append((i, len(sizes), len([v for v in sizes if v < 14])))
    for s in sizes:
        sz[s] += 1
tot = sum(sz.values())
small = sum(v for k, v in sz.items() if k < 14)
media = [n for n in z.namelist() if n.startswith("ppt/media/")]
av = [n for n in z.namelist() if re.search(r"\.(mp4|m4v|mov|wmv|avi|mp3|wav|m4a)$", n, re.I)]
m = re.search(r'<p:sldSz cx="(\d+)" cy="(\d+)"([^/]*)/>', pres)
health = {
    "页数": len(slds),
    "文字run总数": tot,
    "小于14pt的run": small,
    "小于14pt占比": f"{100 * small / tot:.1f}%",
    "最小字号pt": min(sz),
    "含<14pt占比≥80%的页": [dict(页=p, run=n, 小字=s) for p, n, s in per if n and s / n >= 0.8],
    "表格数": sum(len(re.findall(r"<a:tbl>", z.read("ppt/" + rmap[r]).decode("utf8"))) for r in slds),
    "图片数": len(media),
    "音视频部件": av,
    "外部超链接": sum(z.read(n).decode("utf8").count('TargetMode="External"')
                 for n in z.namelist() if n.endswith(".rels") and "slides/_rels" in n),
    "备注页数": len([n for n in z.namelist() if re.match(r"ppt/notesSlides/notesSlide\d+\.xml$", n)]),
    "幻灯片尺寸": dict(EMU=[int(m.group(1)), int(m.group(2))],
                 英寸=[round(int(m.group(1)) / 914400, 3), round(int(m.group(2)) / 914400, 3)],
                 比例=round(int(m.group(1)) / int(m.group(2)), 4),
                 sldSz属性=m.group(3).strip()),
    "生成工具": z.read("docProps/core.xml").decode("utf8") if "docProps/core.xml" in z.namelist() else "",
}
health["生成工具"] = re.findall(r"<dc:description>([^<]*)</dc:description>|<cp:lastModifiedBy>([^<]*)</cp:lastModifiedBy>",
                            health["生成工具"])
with open(os.path.join(OUT, "补充_PPT体检.json"), "w", encoding="utf-8") as f:
    json.dump(health, f, ensure_ascii=False, indent=1)
print("PPT 体检:", health["页数"], "页，小字占比", health["小于14pt占比"], "，图片", health["图片数"], "，音视频", health["音视频部件"])
