import openpyxl, shutil, os
OUT = r'd:\MG\program\ai交流\AI对比\ep3 做表\工作沙箱3\out'
src = OUT + r'\v2_edit.xlsx'
wb = openpyxl.load_workbook(src)
ws = wb['主梁数量']
old = ws['J431'].value
ws['J431'] = '=SUM(J403:J430)+SUM(J432:J3777)'
try:
    wb.calculation.fullCalcOnLoad = True
except Exception as e:
    print('calcflag', e)
wb.save(src)
print('J431', old, '->', ws['J431'].value)
shutil.copy(src, r'C:\temp\ep3recalc\v2_edit.xlsx')
print('copied')
