import openpyxl, shutil, os
OUT = r'd:\MG\program\ai交流\AI对比\ep3 做表\工作沙箱3\out'
src = OUT + r'\v2_edit.xlsx'
wb = openpyxl.load_workbook(src)
ws = wb['合计数量']
changed = []
for r in range(4, 25):
    th = ws.cell(r, 2).value
    c = ws.cell(r, 7)
    old = c.value
    if isinstance(old, str) and old.startswith('=SUMIFS('):
        new = old.replace(f',$B{r})', f',{th})')
        c.value = new
        changed.append((r, th, old, new))
try:
    wb.calculation.fullCalcOnLoad = True
except Exception as e:
    print('calcflag', e)
wb.save(src)
print('changed', len(changed))
for r, th, o, n in changed[:5]:
    print(' ', r, th, n)
shutil.copy(src, r'C:\temp\ep3recalc\v2_edit.xlsx')
print('copied')
