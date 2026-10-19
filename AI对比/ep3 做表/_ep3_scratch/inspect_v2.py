import openpyxl
OUT = r'd:\MG\program\ai交流\AI对比\ep3 做表\工作沙箱3\out'
wb = openpyxl.load_workbook(OUT + r'\v2_edit.xlsx')
ws = wb['主梁数量']
print('max_row', ws.max_row, 'max_col', ws.max_column)
print('J431 =', ws['J431'].value)
# find rows with M formula (block first rows) and the last non-empty data row
mr = [r for r in range(4, ws.max_row + 1)
      if isinstance(ws.cell(r, 13).value, str) and ws.cell(r, 13).value.startswith('=SUMIF')]
print('SUMIF rows', mr)
last = 0
for r in range(4, ws.max_row + 1):
    if any(ws.cell(r, c).value not in (None, '') for c in (1, 2, 5, 6, 7, 9, 10)):
        last = r
print('last data-ish row', last)
print('row431 all:', [ws.cell(431, c).value for c in range(1, 24)])
print('rows 403,430,431,432 J:', [(r, ws.cell(r, 10).value) for r in (403, 430, 431, 432)])
print('row 432,2750,2751,2752,2753 A..K:')
for r in (432, 2750, 2751, 2752, 2753, 2754):
    print(' ', r, [ws.cell(r, c).value for c in range(1, 12)])
