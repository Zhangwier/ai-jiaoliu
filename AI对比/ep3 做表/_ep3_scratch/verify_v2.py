import openpyxl, json
OUT = r'd:\MG\program\ai交流\AI对比\ep3 做表\工作沙箱3\out'
P = OUT + r'\final_v2.xlsx'
ERR = ('#REF!', '#VALUE!', '#DIV/0!', '#NAME?', '#NULL!', '#N/A')
wv = openpyxl.load_workbook(P, data_only=True)
wf = openpyxl.load_workbook(P, data_only=False)

errs = {}
for ws in wv.worksheets:
    h = [(c.coordinate, c.value) for row in ws.iter_rows() for c in row
         if isinstance(c.value, str) and c.value.strip() in ERR]
    if h: errs[ws.title] = h

empty = 0
for sn in wf.sheetnames:
    a, b = wf[sn], wv[sn]
    for row in a.iter_rows():
        for c in row:
            if c.data_type == 'f':
                if b[c.coordinate].value is None:
                    empty += 1

merges = {ws.title: len(ws.merged_cells.ranges) for ws in wf.worksheets}
print('sheets', wf.sheetnames)
print('errors', errs, 'total', sum(len(v) for v in errs.values()))
print('empty-cache formulas', empty)
print('merges', merges, 'total', sum(merges.values()))
hq = wv['合计数量']
print('合计数量 钢梁制造合计 F27/G27/H27 ->', hq['F27'].value, hq['G27'].value, hq['H27'].value)
ws = wv['主梁数量']
jt = sum(v for r in range(4, 2753) for v in [ws.cell(r, 10).value] if isinstance(v, (int, float)))
print('主梁明细 J 合计', round(jt, 2))
print('J431', ws['J431'].value)
json.dump({'errors': errs, 'empty_cache': empty, 'merges': merges,
           'zhuliang_J_total': round(jt, 2)}, open(OUT + r'\final_v2_check.json', 'w', encoding='utf-8'),
          ensure_ascii=False, indent=1)
