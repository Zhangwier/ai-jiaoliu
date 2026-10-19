# -*- coding: utf-8 -*-
"""
E2 v2 重建写入层：主梁明细净化 / 合计数量全厚度覆盖+新沙图纸数量 / 横隔板坏引用修复。
基座 filled.xlsx -> 复制为 v2_edit.xlsx 后再改。绝不改 filled.xlsx / final.xlsx。
"""
import os, json, copy, shutil, re
import pandas as pd
import openpyxl

OUT = r'd:\MG\program\ai交流\AI对比\ep3 做表\工作沙箱3\out'
SCR = r'd:\MG\program\ai交流\AI对比\ep3 做表\_ep3_scratch'
FILLED = os.path.join(OUT, 'filled.xlsx')
V2 = os.path.join(OUT, 'v2_edit.xlsx')

def num(x):
    try:
        if x is None or (isinstance(x, float) and pd.isna(x)):
            return None
        v = float(x)
        return v
    except Exception:
        return None

def numify(v):
    """把数值转成 int(若整) / float，其它原样返回。"""
    if v is None:
        return None
    try:
        f = float(v)
        if f == int(f):
            return int(f)
        return f
    except Exception:
        return v

log = {}

# ---------------------------------------------------------------------------
# 1. 读标准化数据，分类
# ---------------------------------------------------------------------------
df = pd.read_csv(os.path.join(OUT, 'normalized_tables.csv'), encoding='utf-8-sig')
zl = df[df['构件类别'] == '主梁'].copy()
L = pd.to_numeric(zl['长'], errors='coerce')
W = pd.to_numeric(zl['宽'], errors='coerce')
T = pd.to_numeric(zl['厚'], errors='coerce')
Q = pd.to_numeric(zl['数量'], errors='coerce')
H = pd.to_numeric(zl['单件重'], errors='coerce')
TOT = pd.to_numeric(zl['合计重量'], errors='coerce')

missdim = L.isna() | W.isna() | T.isna()

# 明细 = 长/宽/厚 均为数值
detail = zl[~missdim].copy()
excluded = zl[missdim].copy()

# --- 缺尺寸清单.csv ---
rows_miss = []
for _, r in excluded.iterrows():
    miss_fields = ','.join([f for f, s in (('长', r['长']), ('宽', r['宽']), ('厚', r['厚']))
                            if num(s) is None])
    tab = str(r['表名'])
    spec = str(r['规格'])
    if ('汇总' in tab) or re.match(r'^t=\d+$', spec):
        dest = '汇总表行(材质+规格厚+单个重量)→移出明细区；作新沙图纸数量来源'
    else:
        dest = '缺尺寸行→移出明细区'
    rows_miss.append({
        '来源页': r['来源页'], '表名': tab, '材质': r['材质'],
        '厚': r['厚'], '数量': r['数量'], '缺失字段': miss_fields, '建议去向': dest,
    })
pd.DataFrame(rows_miss).to_csv(os.path.join(OUT, '缺尺寸清单.csv'), index=False, encoding='utf-8-sig')

# --- fail 分类 ---
def classify_fail(r):
    if num(r['长']) is None or num(r['宽']) is None or num(r['厚']) is None:
        return '应删'
    q, h, tot = num(r['数量']), num(r['单件重']), num(r['合计重量'])
    if q is None or h is None or tot is None or tot == 0:
        return '待核'
    # sheet 重算总重 = q * (L*W*t*7850/1e9)
    l, w, t = num(r['长']), num(r['宽']), num(r['厚'])
    sheet_tot = q * (l / 1000) * (w / 1000) * (t / 1000) * 7850
    rel = abs(sheet_tot - tot) / tot
    return '可修' if rel <= 0.05 else '待核'

fail_rows = zl[zl['校验状态'] == 'fail'].copy()
fail_csv = []
fail_counts = {'可修': 0, '应删': 0, '待核': 0}
for _, r in fail_rows.iterrows():
    cls = classify_fail(r)
    fail_counts[cls] += 1
    q, h, tot = num(r['数量']), num(r['单件重']), num(r['合计重量'])
    if q and h:
        rel = abs(q * h - tot) / tot if tot else None
    else:
        rel = None
    fail_csv.append({
        '来源页': r['来源页'], '表名': str(r['表名']), '梁板编号': r['梁板编号'],
        '厚': r['厚'], '长': r['长'], '宽': r['宽'], '数量': r['数量'],
        '单件重': r['单件重'], '合计重量': r['合计重量'],
        '相对误差': (f'{rel*100:.3f}%' if rel is not None else ''),
        '结论': cls,
        '依据': ('缺长/宽/厚→汇总表行' if cls == '应删'
                 else ('尺寸齐全且重算后与OCR合计吻合(≤5%)→表内公式自动修正' if cls == '可修'
                       else '尺寸齐全但数量×单件重≠合计且重算偏差>5%→需人工核对')),
    })
pd.DataFrame(fail_csv).to_csv(os.path.join(OUT, 'fail分类.csv'), index=False, encoding='utf-8-sig')

# ---------------------------------------------------------------------------
# 2. 新沙图纸数量：对"汇总表行"(933 ∩ 有厚 ∩ 表名含汇总/规格 t=NN) 聚合 sum(单件重)
# ---------------------------------------------------------------------------
src_mask = missdim & Q.notna() & T.notna()
src = zl[src_mask].copy()
src['is_summary'] = src['表名'].astype(str).str.contains('汇总', na=False) | \
                    src['规格'].astype(str).str.match(r'^t=\d+$', na=False)
sumrows = src[src['is_summary']].copy()
sumrows['_t'] = pd.to_numeric(sumrows['厚'], errors='coerce')
sumrows['_h'] = pd.to_numeric(sumrows['单件重'], errors='coerce')
sumrows['_mat'] = sumrows['材质'].astype(str).str.strip()
agg = {}   # (t, mat) -> {'sum':.., 'pages':set()}
for _, r in sumrows.iterrows():
    key = (numify(r['_t']), ('Q355' if '355' in r['_mat'] else ('Q420' if '420' in r['_mat'] else r['_mat'])))
    d = agg.setdefault(key, {'sum': 0.0, 'pages': set()})
    d['sum'] += float(r['_h']) if pd.notna(r['_h']) else 0.0
    for p in str(r['来源页']).replace('，', ',').split(','):
        p = p.strip()
        if p:
            d['pages'].add(p)

# ---------------------------------------------------------------------------
# 3. 复制 filled -> v2_edit，载入
# ---------------------------------------------------------------------------
shutil.copyfile(FILLED, V2)
wb = openpyxl.load_workbook(V2, data_only=False)
sheet_names = wb.sheetnames

# ---------------------------------------------------------------------------
# 4. 主梁数量明细区净化（压缩重写, 保留合并锚点与公式列）
# ---------------------------------------------------------------------------
ws = wb['主梁数量']
WRITE_COLS = [1, 2, 3, 4, 5, 6, 7, 9, 11]   # A B C D E F G I K
FORMULA_COLS = [8, 10, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23]  # H J L M..W

# 合并成员表
merged_members = {}
for m in ws.merged_cells.ranges:
    for rr in range(m.min_row, m.max_row + 1):
        for cc in range(m.min_col, m.max_col + 1):
            merged_members[(rr, cc)] = (m.min_row, m.min_col)

subtotal_rows = set()
for r in range(4, 432):
    if ws.cell(row=r, column=2).value == '小计':
        subtotal_rows.add(r)

slot_rows = [r for r in range(4, 3778) if r not in subtotal_rows]

# 清空所有数据槽的可写列（跳过小计行与公式列；非锚点合并成员不可写）
cleared = 0
for r in slot_rows:
    for c in WRITE_COLS:
        k = (r, c)
        if k in merged_members and merged_members[k] != k:
            continue
        cell = ws.cell(row=r, column=c)
        if cell.data_type == 'f':
            continue
        cell.value = None
        cleared += 1

# 明细写入
skipped_merge = 0
skipped_formula = 0
written_rows = 0
detail = detail.reset_index(drop=True)
for i in range(len(detail)):
    if i >= len(slot_rows):
        break
    sr = slot_rows[i]
    r = detail.iloc[i]
    q = num(r['数量'])
    pagestr = str(r['来源页'])
    base = f'来源页{pagestr}'
    is_fail = (str(r['校验状态']) == 'fail')
    if is_fail:
        # 复用分类：按行内容重新 classify_fail
        cls = classify_fail(r)
        ktext = f'【{cls}】{base}'
    elif q is None:
        ktext = base + '；【缺数量待核】'
    else:
        ktext = base
    vals = {
        1: r['梁板编号'], 2: r['板件名称'], 3: r['位置'], 4: r['材质'],
        5: numify(num(r['长'])), 6: numify(num(r['宽'])), 7: numify(num(r['厚'])),
        9: numify(q), 11: ktext,
    }
    for c, v in vals.items():
        k = (sr, c)
        if k in merged_members and merged_members[k] != k:
            skipped_merge += 1
            continue
        cell = ws.cell(row=sr, column=c)
        if cell.data_type == 'f':
            skipped_formula += 1
            continue
        cell.value = v
    written_rows += 1

# ---------------------------------------------------------------------------
# 5. 合计数量 重建（全厚度覆盖 + SUMIFS 复核 + 新沙图纸数量）
# ---------------------------------------------------------------------------
TARGET_TH = [0.18, 4, 6, 8, 10, 12, 14, 16, 20, 24, 28, 30, 32, 36, 40, 44, 48, 50, 60, 80, 100]
HZ_BUCKET = {10: 'N139', 12: 'O139', 14: 'P139', 16: 'Q139', 20: 'R139', 24: 'S139',
             28: 'T139', 30: 'U139', 40: 'V139', 50: 'W139'}

wsq = wb['合计数量']
# 捕获原样式
orig = {}
for r in range(1, wsq.max_row + 1):
    for c in range(1, 10):
        cell = wsq.cell(row=r, column=c)
        orig[(r, c)] = (cell.value, copy.copy(cell._style))
# 记录原始 10 个合并
orig_merges = [str(m) for m in wsq.merged_cells.ranges]
for m in list(wsq.merged_cells.ranges):
    wsq.unmerge_cells(str(m))

def setcell(r, c, val, style=None):
    cell = wsq.cell(row=r, column=c)
    cell.value = val
    if style is not None:
        cell._style = copy.copy(style)

# 表头 1..3 原样
for r in (1, 2, 3):
    for c in range(1, 10):
        v, st = orig.get((r, c), (None, None))
        setcell(r, c, v, st)

# 钢梁 21 行：样式取原第 4 行
steel_style = {c: orig[(4, c)][1] for c in range(1, 10)}
draw_cover = 0
draw_todo = 0
draw_log = []
for i, th in enumerate(TARGET_TH):
    r = 4 + i
    q355 = agg.get((numify(th), 'Q355'))
    q420 = agg.get((numify(th), 'Q420'))
    dv = round(q355['sum'], 1) if q355 else None
    ev = round(q420['sum'], 1) if q420 else None
    if dv is not None or ev is not None:
        draw_cover += 1
        pages = sorted(set((q355['pages'] if q355 else set()) | (q420['pages'] if q420 else set())), key=lambda x: int(x) if x.isdigit() else 0)
        note = '来源页' + ','.join(pages)
    else:
        draw_todo += 1
        dv = ev = None
        note = '未完成:无来源页'
    setcell(r, 1, ('钢梁' if i == 0 else None), steel_style[1])
    setcell(r, 2, th, steel_style[2])
    setcell(r, 3, 'kg', steel_style[3])
    setcell(r, 4, dv, steel_style[4])
    setcell(r, 5, ev, steel_style[5])
    setcell(r, 6, f'=D{r}+E{r}', steel_style[6])
    hzc = HZ_BUCKET.get(th)
    gform = f'=SUMIFS(主梁数量!$J$4:$J$3777,主梁数量!$G$4:$G$3777,$B{r})'
    if hzc:
        gform += f'+横隔板!{hzc}'
    setcell(r, 7, gform, steel_style[7])
    setcell(r, 8, f'=G{r}-F{r}', steel_style[8])
    setcell(r, 9, note, steel_style[9])
    draw_log.append({'板厚': th, 'Q355': dv, 'Q420': ev, '备注': note})

# 梁底挡块(原14) / 钢梁制造合计(原15) / 附属预埋(原16,17)
row_map = {25: 14, 26: 15, 27: 16, 28: 17, 29: 18}
for newr, oldr in row_map.items():
    for c in range(1, 10):
        v, st = orig.get((oldr, c), (None, None))
        setcell(newr, c, v, st)
# 补齐新行公式引用
setcell(25, 8, '=G25-F25', orig[(14, 8)][1])
setcell(26, 6, '=SUM(F4:F25)', orig[(15, 6)][1])
setcell(26, 7, '=SUM(G4:G25)', orig[(15, 7)][1])
setcell(26, 8, '=G26-F26', orig[(15, 8)][1])
setcell(27, 8, '=G27-F27', orig[(16, 8)][1])
setcell(28, 8, '=G28-F28', orig[(17, 8)][1])
setcell(28, 1, None, orig[(17, 1)][1])   # 合并在 A27:A28

# 重建合并
rebuild_merges = ['A1:I1', 'A2:A3', 'B2:B3', 'C2:C3', 'D2:F2', 'G2:G3', 'H2:H3', 'I2:I3',
                  'A4:A24', 'A27:A28']
for m in rebuild_merges:
    wsq.merge_cells(m)

# ---------------------------------------------------------------------------
# 6. 横隔板!D140 坏引用修复
# ---------------------------------------------------------------------------
wsh = wb['横隔板']
d140 = wsh['D140']
bad_ref = d140.value
if d140.data_type == 'f' or (isinstance(bad_ref, str) and bad_ref.startswith('=')):
    d140.value = None

# ---------------------------------------------------------------------------
# 7. 保存
# ---------------------------------------------------------------------------
wb.save(V2)

# ---------------------------------------------------------------------------
# 8. 自检 / 日志
# ---------------------------------------------------------------------------
wb2 = openpyxl.load_workbook(V2, data_only=False)
wbf = openpyxl.load_workbook(FILLED, data_only=False)

def count_formulas(wbk):
    per = {}
    tot = 0
    for sh in wbk.worksheets:
        n = sum(1 for row in sh.iter_rows() for c in row if c.data_type == 'f')
        per[sh.title] = n
        tot += n
    return per, tot

f_per, f_tot = count_formulas(wbf)
v_per, v_tot = count_formulas(wb2)
m_filled = {sh.title: len(sh.merged_cells.ranges) for sh in wbf.worksheets}
m_v2 = {sh.title: len(sh.merged_cells.ranges) for sh in wb2.worksheets}

log = {
    'base_file': FILLED,
    'out_file': V2,
    'sheets_order_v2': wb2.sheetnames,
    'sheets_order_ok': wb2.sheetnames == ['合计数量', '主梁数量', '横隔板', '梁底挡块', '附属预埋件'],
    '主梁数量': {
        'source_main_rows': 3758,
        'detail_written_rows': written_rows,
        'missing_dim_rows_total': int(missdim.sum()),
        'missing_dim_with_qty_rows': int((missdim & Q.notna()).sum()),
        'missing_dim_no_qty_rows': int((missdim & Q.isna()).sum()),
        'note_19_no_qty_kept': '长宽厚齐全但数量不可解析的 19 行仍写入明细(数量留空+备注【缺数量待核】)',
        'skipped_by_merge_anchor_cells': skipped_merge,
        'skipped_formula_cells': skipped_formula,
        'cleared_cells': cleared,
        'cleared_slot_rows_tail': [slot_rows[written_rows], 3777] if written_rows < len(slot_rows) else None,
        'subtotal_rows_preserved': sorted(subtotal_rows),
    },
    'fail分类计数': fail_counts,
    'fail分类明细cvs': os.path.join(OUT, 'fail分类.csv'),
    '缺尺寸清单cvs': os.path.join(OUT, '缺尺寸清单.csv'),
    '合计数量': {
        'inserted_rows_count': len(TARGET_TH) - 10,
        'insert_position': '原第14行(梁底挡块)前插入 11 行; 钢梁块由 4..13(10行) 扩为 4..24(21行); 原14~18 下移为 25~29',
        'steel_block_rows': '4..24',
        'thickness_buckets': TARGET_TH,
        'bucket_count': len(TARGET_TH),
        'draw_qty_covered_buckets': draw_cover,
        'draw_qty_todo_buckets': draw_todo,
        'draw_qty_agg_rule': '933 中 表名含"汇总"/规格t=NN 的汇总表行, 按(厚,材质) sum(单件重)',
        'draw_qty_by_bucket': draw_log,
        'rebuilt_merges': rebuild_merges,
        'orig_merges': orig_merges,
        'sum_formula_example': 'G4=SUMIFS(主梁数量!$J$4:$J$3777,主梁数量!$G$4:$G$3777,$B4)',
    },
    '横隔板': {
        'bad_ref_cell': 'D140',
        'was': bad_ref,
        'fixed_to': 'cleared(None)',
    },
    '公式格数': {
        'filled_total': f_tot,
        'v2_total': v_tot,
        'v2_ge_filled': v_tot >= f_tot,
        'filled_per_sheet': f_per,
        'v2_per_sheet': v_per,
        'formula_cols_written_times': 0,
    },
    '合并单元格': {
        'filled_total': sum(m_filled.values()),
        'filled_per_sheet': m_filled,
        'v2_total': sum(m_v2.values()),
        'v2_per_sheet': m_v2,
        'target_735': sum(m_v2.values()) == 735,
    },
}

with open(os.path.join(OUT, 'e2_log.json'), 'w', encoding='utf-8') as f:
    json.dump(log, f, ensure_ascii=False, indent=2)

print('DONE')
print('detail_written_rows =', written_rows)
print('missing_dim total =', int(missdim.sum()), ' with_qty =', int((missdim & Q.notna()).sum()))
print('fail counts =', fail_counts)
print('skipped_merge =', skipped_merge, ' skipped_formula =', skipped_formula)
print('draw_cover =', draw_cover, ' draw_todo =', draw_todo)
print('formula filled/v2 =', f_tot, '/', v_tot, ' ge:', v_tot >= f_tot)
print('merges filled/v2 =', sum(m_filled.values()), '/', sum(m_v2.values()))
print('sheets =', wb2.sheetnames)
