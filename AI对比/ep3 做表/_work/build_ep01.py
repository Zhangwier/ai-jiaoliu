# -*- coding: utf-8 -*-
"""中桁节段01 计算底稿生成脚本（试验）
沿用模板 主梁数量 / 合计数量 两个 sheet 的列结构、公式与格式。
"""
import re
from copy import copy
import openpyxl

TPL = r'd:\MG\program\ai交流\AI对比\ep3 做表\工作沙箱4\模板原件.xlsx'
OUT = r'd:\MG\program\ai交流\AI对比\ep3 做表\工作沙箱4\中桁节段01_计算底稿.xlsx'

# ---------------- 图纸数据（来源：中桁节段01设计图 第5张 节段01材料数量表） ----------------
# (部件名称, 件号, 材质, 规格原文, 数量, 图纸单件重, 备注)
PARTS = [
    ("上弦顶板", "N11", "Q355D", "t=20mm,A=27m2", 1, 4239, ""),
    ("上弦底板", "N12", "Q355D", "□13350×1200×20", 1, 2515.1, ""),
    ("上弦节点板", "N14", "Q355D", "t=28mm,A=13.05m2", 2, 2868.4, ""),
    ("上弦节点板", "N14a", "Q355D", "t=40mm,A=25m2", 2, 7850, ""),
    ("上弦横隔板", "N15", "Q355D", "□1980×1200×16", 6, 298.9, ""),
    ("上弦横隔板", "N15'", "Q355D", "□2057×200×16", 4, 51.5, ""),
    ("上弦横隔板", "N15b", "Q355B", "□2057×40×4", 4, 2.5, "材质Q355B，与主材不同"),
    ("上弦横隔板", "N15c", "Q355D", "□800×400×16", 4, 40.2, ""),
    ("上弦加劲肋", "N16a", "Q355D", "□5651×240×20", 4, 213.5, ""),
    ("上弦加劲肋", "N16b", "Q355D", "□7700×240×20", 4, 290.5, ""),
    ("上弦加劲肋", "N16c", "Q355D", "□7700×240×20", 0, 290.5, "图纸数量为0（共重0），需设计核实"),
    ("上弦加劲肋", "N16d", "Q355D", "□13300×240×20", 1, 500.8, ""),
    ("上弦加劲肋", "N16e", "Q355D", "□13300×240×20", 1, 500.8, ""),
    ("下弦顶板", "N21", "Q355D", "□4790×2500×20", 1, 1880.9, ""),
    ("下弦顶板", "N21a", "Q355D", "□10105×2500×20", 1, 3965.8, ""),
    ("下弦底板", "N22", "Q355D", "t=36mm,A=9.36m2", 1, 2645.1, ""),
    ("下弦底板", "N22a", "Q355D", "t=20mm,A=15.68m2", 1, 2461.8, ""),
    ("下弦腹板", "N23", "Q355D", "□10105×1880×24", 2, 3579.6, ""),
    ("下弦节点板", "N24", "Q355D", "t=40mm,A=17.33m2", 2, 5441.6, "Z35钢"),
    ("下弦横隔板", "N25", "Q355D", "□1800×1200×16", 1, 271.3, ""),
    ("下弦横隔板", "N25'", "Q355D", "□2057×200×16", 3, 51.5, ""),
    ("下弦横隔板", "N25a", "Q355D", "□1800×1200×16", 3, 271.3, ""),
    ("下弦加劲肋", "N26a", "Q355D", "□2792×240×20", 2, 105.2, ""),
    ("下弦加劲肋", "N26b", "Q355D", "□10105×240×20", 2, 381.5, ""),
    ("下弦加劲肋", "N26c", "Q355D", "□12897×240×20", 1, 486.7, ""),
    ("下弦加劲肋", "N26d", "Q355D", "□10105×240×20", 1, 381.5, ""),
    ("下弦加劲肋", "N26d'", "Q355D", "□2792×320×32", 1, 223.6, ""),
    ("下弦加劲肋", "N26e", "Q355D", "□12897×170×14", 2, 240.7, ""),
    ("斜腹杆", "N31", "Q355D", "□14038×1200×36", 1, 4761.8, ""),
    ("斜腹杆", "N31a", "Q355D", "□14038×1200×36", 1, 4761.8, ""),
    ("斜腹杆", "N32", "Q355D", "□9838×1360×36", 2, 3781.2, ""),
    ("斜腹杆", "N33", "Q355D", "□9838×320×32", 2, 791.3, ""),
    ("斜腹杆", "N33a", "Q355D", "□12638×320×32", 1, 1014.8, ""),
    ("斜腹杆", "N33b", "Q355D", "□12638×320×32", 1, 1014.8, ""),
    ("斜腹杆", "N34", "Q355D", "□1200×1200×16", 4, 180.9, ""),
    ("斜腹杆", "N41", "Q355D", "□6725×1200×36", 1, 2280.6, ""),
    ("斜腹杆", "N41a", "Q355D", "□6925×1200×36", 1, 2348.4, ""),
    ("斜腹杆", "N42", "Q355D", "□4925×1360×36", 2, 1893.4, ""),
    ("斜腹杆", "N43", "Q355D", "□4925×320×32", 2, 396.9, ""),
    ("斜腹杆", "N43a", "Q355D", "□6025×320×32", 1, 484.8, ""),
    ("斜腹杆", "N43b", "Q355D", "□6225×320×32", 1, 499.9, ""),
    ("斜腹杆", "N44", "Q355D", "□1200×1200×16", 2, 180.9, ""),
    ("竖腹杆顶底板", "N51", "Q355D", "□13418×1200×20", 1, 2527.7, ""),
    ("竖腹杆顶底板", "N51a", "Q355D", "□11818×1200×20", 1, 2226.3, ""),
    ("竖腹杆腹板", "N52", "Q355D", "□9133×1500×20", 2, 2150.9, ""),
    ("竖腹杆横隔板", "N54", "Q355D", "□1400×1200×16", 3, 211, ""),
    ("竖腹杆加劲肋", "N53", "Q355D", "□9133×190×16", 2, 218.5, ""),
    ("竖腹杆加劲肋", "N53a", "Q355D", "□12018×190×16", 1, 286.4, ""),
    ("竖腹杆加劲肋", "N53b", "Q355D", "□10418×190×16", 1, 248.7, ""),
    ("节点板加劲", "N61", "Q355D", "□2800×320×32", 2, 226.1, ""),
    ("节点板加劲", "N62", "Q355D", "□2600×320×32", 2, 208.5, ""),
    ("节点板加劲", "N63", "Q355D", "□750×320×28", 2, 52.8, ""),
    ("节点板加劲", "N64", "Q355D", "□2100×320×32", 2, 168.3, ""),
    ("节点板加劲", "N65", "Q355D", "□2100×320×32", 2, 168.3, ""),
    # ---- 附属杆件 ----
    ("横梁接头", "N71", "Q355D", "□1980×610×16", 6, 152, ""),
    ("横梁接头", "N71'", "Q355D", "□1980×650×16", 2, 162, ""),
    ("横梁接头", "N71a", "Q355D", "□1980×622×16", 4, 154.5, ""),
    ("横梁接头", "N72", "Q355D", "t=24mm,A=0.47m2", 6, 88.4, ""),
    ("横梁接头", "N72'", "Q355D", "t=24mm,A=0.48m2", 6, 90.4, ""),
    ("横梁接头", "N72a", "Q355D", "t=24mm,A=0.97m2", 2, 182, ""),
    ("横梁接头", "N72b", "Q355D", "t=24mm,A=0.48m2", 2, 90, ""),
    ("横梁接头", "N73", "Q355D", "□1800×626×16", 8, 141.9, ""),
    ("横梁接头", "N73a", "Q355D", "□1800×610×28", 4, 241.8, ""),
    ("支承加劲肋", "N81", "Q355D", "□1800×1200×32", 4, 542.6, ""),
    ("支承加劲肋", "N82", "Q355D", "□502×240×20", 2, 18.8, ""),
    ("支承加劲肋", "N83", "Q355D", "□1800×1800×100", 1, 2543.4, ""),
    ("支承加劲肋", "N84", "Q355D", "□1000×568×32", 3, 143.2, ""),
    ("支承加劲肋", "N85", "Q355D", "□1350×570×32", 4, 193.4, ""),
]

BUCKETS = [4, 14, 16, 20, 24, 28, 32, 36, 40, 100]
# 图纸《节段01材料数量汇总表》——图纸数量（kg）
DRAWING_BY_T = {100: 2543.4, 40: 26583.2, 36: 28146.9, 32: 10530.3, 28: 6809.6,
                24: 8776.0, 20: 29015.2, 16: 9079.6, 14: 481.4, 4: 10.0}
DRAWING_SUBTOTAL = 121975.6
DRAWING_WELD = 1829.6
DRAWING_TOTAL = 123805.2


def parse_spec(spec):
    """返回 (长mm, 宽mm, 厚mm, 面积m2)；矩形返回长宽，面积型返回面积。"""
    s = spec.replace('□', '').replace('x', '×')
    m = re.match(r'^t=(\d+(?:\.\d+)?)mm,A=([\d.]+)m2$', s)
    if m:
        return (None, None, float(m.group(1)), float(m.group(2)))
    nums = re.split(r'×', s)
    if len(nums) == 3:
        return (float(nums[0]), float(nums[1]), float(nums[2]), None)
    raise ValueError(spec)


def cp(dst, src):
    dst._style = copy(src._style)


wb = openpyxl.load_workbook(TPL)
src = wb['主梁数量']
sum_src = wb['合计数量']

# ============ Sheet 1: 主梁数量-节段01（沿用 主梁数量 列结构与公式） ============
ws = wb.create_sheet('主梁数量-节段01')
for col in 'ABCDEFGHIJKLMNOPQRSTUVWXYZ':
    if col in src.column_dimensions and src.column_dimensions[col].width:
        ws.column_dimensions[col].width = src.column_dimensions[col].width

# 行高
for r in (1, 2, 3):
    if r in src.row_dimensions and src.row_dimensions[r].height:
        ws.row_dimensions[r].height = src.row_dimensions[r].height

# 表头
ws.merge_cells('A1:L1')
ws['A1'] = '中桁节段01材料数量表（计算底稿）'
cp(ws['A1'], src['A1'])

hdr = ['梁板编号', '板件名称', '位置', '材质', '长', '宽', '厚', '单件重', '数量', '合计重量', '规格（图纸）', '备注']
for i, t in enumerate(hdr):
    c = ws.cell(row=2, column=1 + i, value=t)
    cp(c, src.cell(row=2, column=1 + i))
ws['E3'] = 'mm'; ws['F3'] = 'mm'; ws['G3'] = 'mm'; ws['H3'] = 'kg'; ws['J3'] = 'kg'
for col in 'EFGHJ':
    cp(ws[col + '3'], src[col + '3'])

ws.merge_cells('M2:V2')
ws['M2'] = '复核数量（按板厚，kg）'
cp(ws['M2'], src['M2'])
ws['W2'] = '合计'; ws['X2'] = '图纸数量'; ws['Y2'] = '差异'; ws['Z2'] = '备注'
for col in 'WXYZ':
    cp(ws[col + '2'], src[col + '2'])
for i, b in enumerate(BUCKETS):
    c = ws.cell(row=3, column=13 + i, value=b)
    cp(c, src.cell(row=3, column=13 + i))
ws['W3'] = '合计'
cp(ws['W3'], src['W3'])

r = 4
start_data = r
group_rows = []          # (部件名称, 起始行, 结束行)
subtotal_cells = []      # 各部件小计所在行的 J 单元格
cur = None
g_start = r
for (grp, part, mat, spec, qty, w, note) in PARTS:
    if cur is None:
        cur = grp; g_start = r
    if grp != cur:
        group_rows.append((cur, g_start, r - 1))
        # 小计行
        c = ws.cell(row=r, column=2, value='小计'); cp(c, src['B32'])
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        jc = ws.cell(row=r, column=10, value='=SUM(J%d:J%d)' % (g_start, r - 1))
        cp(jc, src['J32'])
        subtotal_cells.append(r)
        r += 1
        cur = grp; g_start = r
    L, W, T, A = parse_spec(spec)
    ws.cell(row=r, column=1, value='中桁节段01')
    ws.cell(row=r, column=2, value=grp)
    ws.cell(row=r, column=3, value=part)
    ws.cell(row=r, column=4, value=mat)
    if L is not None:
        ws.cell(row=r, column=5, value=L)
        ws.cell(row=r, column=6, value=W)
        wt = '=(E%d/1000)*(F%d/1000)*(G%d/1000)*7850' % (r, r, r)
    else:
        wt = w   # 图纸按面积给出（模板同类单元格亦为图纸值）
        if not note:
            note = ''
        note = ('%s图纸按面积计：%s' % (note + '；' if note else '', spec))
    ws.cell(row=r, column=7, value=T)
    ws.cell(row=r, column=8, value=wt)
    ws.cell(row=r, column=9, value=qty)
    ws.cell(row=r, column=10, value='=H%d*I%d' % (r, r))
    ws.cell(row=r, column=11, value=spec)
    ws.cell(row=r, column=12, value=note if note else None)
    for col in range(1, 13):
        cp(ws.cell(row=r, column=col), src.cell(row=4, column=col if col <= 11 else 11))
    ws.row_dimensions[r].height = src.row_dimensions[4].height
    r += 1
group_rows.append((cur, g_start, r - 1))
c = ws.cell(row=r, column=2, value='小计'); cp(c, src['B32'])
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
jc = ws.cell(row=r, column=10, value='=SUM(J%d:J%d)' % (g_start, r - 1)); cp(jc, src['J32'])
subtotal_cells.append(r)
r += 1
last_data = r - 1

# A 列合并（梁板编号）
ws.merge_cells(start_row=start_data, start_column=1, end_row=last_data, end_column=1)
# B 列按部件合并
for (grp, a, b) in group_rows:
    ws.merge_cells(start_row=a, start_column=2, end_row=b, end_column=2)

# 节段小计
c = ws.cell(row=r, column=2, value='节段小计'); cp(c, src['B32'])
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
jc = ws.cell(row=r, column=10, value='=' + '+'.join('J%d' % i for i in subtotal_cells))
cp(jc, src['J32'])
row_seg_sub = r
sumif = '$G$%d:$G$%d' % (start_data, last_data)
rng_j = '$J$%d:$J$%d' % (start_data, last_data)

# ---- 底部：复核数量 / 图纸数量 / 差异（按板厚） ----
def label_row(row, text):
    c = ws.cell(row=row, column=2, value=text); cp(c, src['B139'])
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)

r += 2
row_recheck = r
label_row(r, '复核数量（按板厚）')
for i in range(len(BUCKETS)):
    col = 13 + i
    cell = ws.cell(row=r, column=col, value='=SUMIF(%s,%s$3,%s)' % (sumif, openpyxl.utils.get_column_letter(col), rng_j))
    cp(cell, src.cell(row=139, column=col))
tot = ws.cell(row=r, column=23, value='=SUM(M%d:V%d)' % (r, r)); cp(tot, src['X139'])
r += 1
row_draw = r
label_row(r, '图纸数量（节段01材料数量汇总表）')
for i, b in enumerate(BUCKETS):
    cell = ws.cell(row=r, column=13 + i, value=DRAWING_BY_T[b])
    cp(cell, src.cell(row=139, column=13 + i))
tot = ws.cell(row=r, column=23, value='=SUM(M%d:V%d)' % (r, r)); cp(tot, src['X139'])
r += 1
row_diff = r
label_row(r, '差异（复核-图纸）')
for i in range(len(BUCKETS)):
    col = 13 + i
    l = openpyxl.utils.get_column_letter(col)
    cell = ws.cell(row=r, column=col, value='=%s%d-%s%d' % (l, row_recheck, l, row_draw))
    cp(cell, src.cell(row=141, column=13 + i))
diff_tot = ws.cell(row=r, column=23, value='=W%d-W%d' % (row_recheck, row_draw)); cp(diff_tot, src['X141'])
r += 1
row_weld = r
label_row(r, '焊缝重（1.5%）')
c = ws.cell(row=r, column=23, value='=ROUND(W%d*0.015,1)' % row_recheck); cp(c, src['X139'])
c2 = ws.cell(row=r, column=24, value=DRAWING_WELD); cp(c2, src['Y139'])
c3 = ws.cell(row=r, column=25, value='=W%d-X%d' % (r, r)); cp(c3, src['Z139'])
r += 1
row_tot = r
label_row(r, '合计（含焊缝）')
c = ws.cell(row=r, column=23, value='=W%d+W%d' % (row_recheck, row_weld)); cp(c, src['X139'])
c2 = ws.cell(row=r, column=24, value=DRAWING_TOTAL); cp(c2, src['Y139'])
c3 = ws.cell(row=r, column=25, value='=W%d-X%d' % (r, r)); cp(c3, src['Z139'])
r += 2
ws.cell(row=r, column=1, value='注：单件重=长×宽×厚×7850/10⁹；面积型板件（t=…,A=…）图纸按面积给出，单件重取图纸值并在“备注”注明；'
                               '图纸数量取自《节段01材料数量汇总表》（图纸第5张）。')

# 删除多余的空白行（保留表头与已写内容，去掉下方模板遗留行）
if ws.max_row > r + 1:
    ws.delete_rows(r + 1, ws.max_row - (r + 1) + 1)

# ============ Sheet 2: 合计数量-节段01（沿用 合计数量 结构：按厚度汇总） ============
ws2 = wb.create_sheet('合计数量-节段01')
for col in 'ABCDEFGHI':
    if col in sum_src.column_dimensions and sum_src.column_dimensions[col].width:
        ws2.column_dimensions[col].width = sum_src.column_dimensions[col].width
ws2.merge_cells('A1:I1')
ws2['A1'] = '中桁节段01 按板厚汇总（计算底稿）'
cp(ws2['A1'], sum_src['A1'])
h2 = ['部位', '板厚', '单位', '图纸数量', None, None, '复核数量', '差异', '备注']
for i, t in enumerate(h2):
    cell = ws2.cell(row=2, column=1 + i, value=t)
    cp(cell, sum_src.cell(row=2, column=1 + i))
ws2.merge_cells('D2:F2')
cp(ws2['D3'], sum_src['D3'])
ws2['D3'] = 'kg'
from openpyxl.utils import get_column_letter as GL
row = 4
first_row = row
for i, b in enumerate(BUCKETS):
    ws2.cell(row=row, column=1, value='中桁节段01')
    ws2.cell(row=row, column=2, value=b)
    ws2.cell(row=row, column=3, value='kg')
    ws2.cell(row=row, column=4, value=DRAWING_BY_T[b])
    ws2.cell(row=row, column=7, value="='主梁数量-节段01'!%s%d" % (GL(13 + i), row_recheck))
    ws2.cell(row=row, column=8, value='=G%d-D%d' % (row, row))
    if b == 4:
        ws2.cell(row=row, column=9, value='图纸t=4（N15b，Q355B）与主材厚度不同，建议核实')
    for col in range(1, 10):
        cp(ws2.cell(row=row, column=col), sum_src.cell(row=(4 if col <= 4 else 4), column=col))
    row += 1
ws2.cell(row=row, column=2, value='焊缝重（1.5%）')
ws2.cell(row=row, column=4, value=DRAWING_WELD)
ws2.cell(row=row, column=7, value="='主梁数量-节段01'!W%d" % row_weld)
ws2.cell(row=row, column=8, value='=G%d-D%d' % (row, row))
for col in range(1, 10):
    cp(ws2.cell(row=row, column=col), sum_src.cell(row=4, column=col))
row += 1
ws2.cell(row=row, column=1, value='钢梁制造合计')
ws2.cell(row=row, column=2, value='合计')
ws2.cell(row=row, column=3, value='kg')
ws2.cell(row=row, column=4, value='=SUM(D%d:D%d)' % (first_row, row - 1))
ws2.cell(row=row, column=7, value='=SUM(G%d:G%d)' % (first_row, row - 1))
ws2.cell(row=row, column=8, value='=G%d-D%d' % (row, row))
for col in range(1, 10):
    cp(ws2.cell(row=row, column=col), sum_src.cell(row=15, column=col if col <= 9 else 9))
row += 2
ws2.cell(row=row, column=1, value='注：图纸数量取自《节段01材料数量汇总表》（图纸第5张）；'
                                  '复核数量按主梁数量-节段01 明细逐件按板厚用SUMIF汇总。')
ws2.cell(row=row + 1, column=1, value='图纸合计 123805.2 = 单个合计 121975.6 + 焊缝重1.5% 1829.6（图纸数）')

wb.save(OUT)

# ---------------- 校验：复核 vs 图纸 ----------------
print('明细行数：', len(PARTS), '  明细区：', start_data, '-', last_data)
wg = openpyxl.load_workbook(OUT)
w2 = wg['主梁数量-节段01']
print('复核(按板厚):', [w2.cell(row=row_recheck, column=13 + i).value for i in range(10)])
# 本地重算（不依赖Excel引擎）
from collections import defaultdict
acc = defaultdict(float)
tot = 0.0
for (grp, part, mat, spec, qty, w, note) in PARTS:
    L, W_, T, A = parse_spec(spec)
    if L is not None:
        piece = L / 1000 * W_ / 1000 * T / 1000 * 7850
    else:
        piece = w
    acc[T] += piece * qty
    tot += piece * qty
print('本地重算(按板厚):', {b: round(acc[b], 1) for b in BUCKETS})
print('本地重算小计:', round(tot, 1), ' 图纸小计:', DRAWING_SUBTOTAL,
      ' 差:', round(tot - DRAWING_SUBTOTAL, 2))
for b in BUCKETS:
    print('  t=%s 复核 %8.1f  图纸 %8.1f  差 %7.1f' % (b, acc[b], DRAWING_BY_T[b], acc[b] - DRAWING_BY_T[b]))
bg = wg['合计数量-节段01']
print('输出文件已生成:', OUT)
