# -*- coding: utf-8 -*-
"""按 _pmap.json 重算全部页内交叉引用（P<N> / 第 N 页）"""
import sys, re, json
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Pt
prs = openprs()
pm = {int(k): int(v) for k, v in json.load(open("_pmap.json")).items()}
def mp(n):
    return pm.get(n, n)

pat_P = re.compile(r"P(\d+)(\s*[–\-—]\s*P?(\d+))?")
pat_C = re.compile(r"第\s*(\d+)\s*页")

def fix_text(t):
    def rp(m):
        a = mp(int(m.group(1)))
        if m.group(3):
            b = mp(int(m.group(3)))
            sep = m.group(2)
            # 保留原分隔符但去掉中间 P
            sep2 = re.sub(r"\d", "", sep).replace("P", "")
            return f"P{a}{sep2}P{b}"
        return f"P{a}"
    t = pat_P.sub(rp, t)
    t = pat_C.sub(lambda m: f"第 {mp(int(m.group(1)))} 页", t)
    return t

n_changed = 0
log = []
for i, sl in enumerate(prs.slides, start=1):
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        tf = sh.text_frame
        for p in tf.paragraphs:
            for r in p.runs:
                new = fix_text(r.text)
                if new != r.text:
                    log.append((i, sh.name, r.text[:60], new[:60]))
                    r.text = new
                    n_changed += 1
print("changed runs", n_changed)
for l in log: print(l)
save(prs)
