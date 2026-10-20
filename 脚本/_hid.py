from pptx import Presentation
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
hidden = [(i+1) for i,s in enumerate(prs.slides) if s.element.get("show")=="0"]
print("total", len(prs.slides), "hidden", hidden)
