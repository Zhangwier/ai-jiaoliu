# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[88]
for sh in s.shapes:
    if sh.has_text_frame:
        t = sh.text_frame.text
        if t.startswith("共 16 页"):
            for p in sh.text_frame.paragraphs:
                if p.runs:
                    p.runs[0].text = "共 16 页，都可以直接拿去用：E1 降智测试（附两个自查工具） · E2 中转站与代理池 · E3 护栏与破限 · E4 开源平台对照"
                    for r in p.runs[1:]: r.text = ""
            print("ok")
prs.save(P)
