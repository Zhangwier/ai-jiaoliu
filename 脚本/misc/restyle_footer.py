# -*- coding: utf-8 -*-
"""弱化页脚溯源区：字号降到 8.5pt、颜色降为 #94A3B8、去掉加粗，
   分隔线由 C_LINE 改为更浅的 E4E7EB，让页面中心完全让给正文。
   保留“来源索引 ›”的整形状超链接。"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

GRAY = "94A3B8"          # 溯源文字
HAIRLINE = "E4E7EB"      # 分隔线
FOOTER_Y = 6.85          # 只处理这条线以下的元素


def is_footer(shape):
    return shape.top is not None and shape.top / 914400 >= FOOTER_Y


def restyle_run(r, size, bold=False):
    rPr = r._r.get_or_add_rPr()
    # 先清掉旧的颜色（主题色或上次写的）
    for el in rPr.findall(qn("a:solidFill")):
        rPr.remove(el)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(GRAY)


def main(src, dst, size=8.5):
    prs = Presentation(src)
    n_txt = n_link = n_line = 0

    for slide in prs.slides:
        for sh in slide.shapes:
            # 1) 分隔线变浅
            if sh.shape_type is not None and sh.name.startswith("Connector"):
                if sh.top is not None and abs(sh.top / 914400 - 7.00) < 0.06:
                    sh.line.color.rgb = RGBColor.from_string(HAIRLINE)
                    sh.line.width = Pt(0.5)
                    n_line += 1
                continue
            if not sh.has_text_frame or not is_footer(sh):
                continue

            text = sh.text_frame.text.strip()
            if not text:
                continue

            is_link = text.startswith("来源索引")
            is_src = text.startswith("来源") or text.startswith("说明：")

            if not (is_link or is_src):
                continue

            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    restyle_run(r, size, bold=False)

            if is_link:
                # 右对齐保持不变，只降灰
                n_link += 1
            else:
                n_txt += 1

    prs.save(dst)
    print(f"restyled: source {n_txt}, link {n_link}, hairlines {n_line} -> {dst}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
