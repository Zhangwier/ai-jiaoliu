# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Cm
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[23]
# 找错位的 moonshot 徽标（top≈12.16 且 left≈5.77）
kill=[]
for sh in list(s.shapes):
    if sh.name in ("LogoBadge","Logo"):
        t=sh.top/360000; l=sh.left/360000
        if abs(t-12.16)<0.1 and l>5.6:
            kill.append(sh)
for sh in kill:
    sh._element.getparent().remove(sh._element)
print("removed", len(kill))
prs.save(P)
