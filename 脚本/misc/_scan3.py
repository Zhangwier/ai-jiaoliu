# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
import re
found=[]
for i,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        t=sh.text_frame.text
        if "三项测试" in t or "降智测试" in t:
            found.append((i, sh.shape_id, t.replace("\n","\\n")[:120]))
for f in found: print(f)
