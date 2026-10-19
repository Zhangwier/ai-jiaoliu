# -*- coding: utf-8 -*-
"""列出附录页内所有 字母+数字 候选，人工确认哪些是附录编号。"""
import re
from pptx import Presentation

p = Presentation("AI交流分享-成品26.pptx")
CODE = re.compile(r"(?<![A-Za-z0-9])([A-E])(\d{1,2})(b?)(?![0-9A-Za-z])")

print("########## 形状内文本 ##########")
for i in range(58, 115):
    s = p.slides[i - 1]
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if not t.strip():
            continue
        hits = CODE.findall(t)
        if hits:
            ctx = []
            for m in CODE.finditer(t):
                ctx.append(t[max(0, m.start() - 12):m.end() + 10].replace("\n", "/"))
            print(f"[{i}][{sh.name}] " + " || ".join(ctx))

print("\n########## 备注 ##########")
for i in range(58, 115):
    s = p.slides[i - 1]
    if not s.has_notes_slide:
        continue
    t = s.notes_slide.notes_text_frame.text
    if not t.strip():
        continue
    hits = CODE.findall(t)
    if hits:
        ctx = []
        for m in CODE.finditer(t):
            ctx.append(t[max(0, m.start() - 12):m.end() + 10].replace("\n", "/"))
        print(f"[{i}][NOTES] " + " || ".join(ctx))
