# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[4]
for sh in s.shapes:
    if sh.shape_id==27:
        tf=sh.text_frame
        for k,p in enumerate(tf.paragraphs):
            print("p%d"%k, [(r.text, r.font.size.pt if r.font.size else None) for r in p.runs])
