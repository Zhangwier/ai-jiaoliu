# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Pt
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
def set_text(sh, new):
    tf=sh.text_frame; p=tf.paragraphs[0]; runs=list(p.runs)
    if not runs: tf.text=new; return
    runs[0].text=new
    for r in runs[1:]: r._r.getparent().remove(r._r)

# s58 副标题
for sh in prs.slides[57].shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("B1–B13"):
        set_text(sh, "B1–B13 每家一页：模型线、可用入口、强项、边界；OpenAI、Anthropic、DeepSeek、Qwen 各另有一页专题。")
        print("s58 ok")

# s90 E1 表格：标注哪两行是工具
s = prs.slides[89]; S = {sh.shape_id: sh for sh in s.shapes}
pairs = [(17, "鹈鹕动画"), (23, "糖果题"), (29, "多合一题\n（工具 · E1b）"), (35, "模型指纹\n（工具 · E1c）")]
for sid, txt in pairs:
    set_text(S[sid], txt)
# 名称列文字略缩，容纳两行
for sid in (17, 23, 29, 35):
    for p in S[sid].text_frame.paragraphs:
        for r in p.runs: r.font.size = Pt(13)
# 副标题改口径
set_text(S[6], "四种题各测一件事；前两种是判分用的测试题，后两种是可以直接拿来用的自查工具")
# 结论
set_text(S[45], "一道题只能测一件事；说清“推不出什么”，结论才立得住")
prs.save(P)
print("done")
