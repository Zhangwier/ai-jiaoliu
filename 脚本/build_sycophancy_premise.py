# -*- coding: utf-8 -*-
"""
插入一页：迎合的第二种形态——顺着你给的前提往下推
插到第 15 页（迎合）之后，成为新的第 16 页。
另外：第 17 页（四个问题小结）的「迎合」一行补上“前提”这一面；全篇页码重编。
输出到 AI交流分享-成品26.pptx（不改 25 原稿）
"""
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

SRC = r"D:\MG\program\ai交流\AI交流分享-成品25.pptx"
OUT = r"D:\MG\program\ai交流\AI交流分享-成品26.pptx"
FONT = "微软雅黑"

C_TITLE = "14213D"
C_TXT = "2B3445"
C_MUTED = "6B7385"
C_TEAL = "0F766E"
C_ORANGE = "D9622B"
C_BG_SOFT = "E1F0EE"
C_BG_WARN = "FBE9DE"
C_BG_BEIGE = "ECEAE4"
C_DARK = "12203A"
C_WHITE = "FFFFFF"
C_PAGE_BG = "F6F5F1"
C_LINE = "DAD7CE"

NSMAP = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
         'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
         'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')


def set_fill(shape, hex_color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hex_color)


def no_line(shape):
    shape.line.fill.background()


def no_shadow(shape):
    spPr = shape._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    spPr.append(parse_xml('<a:effectLst %s/>' % NSMAP))


def rect(slide, x, y, w, h, color, border=None, border_pt=0.8):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(sh, color)
    if border:
        sh.line.color.rgb = RGBColor.from_string(border)
        sh.line.width = Pt(border_pt)
    else:
        no_line(sh)
    no_shadow(sh)
    return sh


def tb(slide, x, y, w, h, text, size=12, bold=False, color=C_TXT,
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.25):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = line
        f = r.font
        f.name = FONT
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = RGBColor.from_string(color)
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            if rPr.find(qn(tag)) is None:
                rPr.append(parse_xml('<%s %s typeface="%s"/>' % (tag, NSMAP, FONT)))
    return box


def set_bg(slide, hex_color):
    cSld = slide._element.find(qn("p:cSld"))
    for old in cSld.findall(qn("p:bg")):
        cSld.remove(old)
    cSld.insert(0, parse_xml(
        '<p:bg %s><p:bgPr><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
        '<a:effectLst/></p:bgPr></p:bg>' % (NSMAP, hex_color)))


def chrome(slide, page_no, source_text):
    tb(slide, 0.6, 0.32, 6.0, 0.3, "01  认识 AI", 12, True, C_TEAL)
    tb(slide, 11.73, 0.32, 1.0, 0.3, str(page_no), 12, False, C_MUTED, align=PP_ALIGN.RIGHT)
    rect(slide, 0.6, 0.66, 0.5, 0.06, C_TEAL)
    ln = slide.shapes.add_connector(1, Inches(0.6), Inches(7.0), Inches(12.73), Inches(7.0))
    ln.line.color.rgb = RGBColor.from_string("E4E7EB")
    ln.line.width = Pt(0.5)
    tb(slide, 0.6, 7.06, 10.5, 0.3, source_text, 8.5, False, "94A3B8")
    tb(slide, 10.73, 7.06, 2.0, 0.3, "来源索引 ›", 8.5, False, "94A3B8")


def build(prs, page_no):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, C_PAGE_BG)
    chrome(s, page_no,
           "来源：本地审核记录（2026-10）；OpenAI math 仓库 history.md（2026-10-07）；第 15 页“迎合”")

    # 标题
    tb(s, 0.6, 0.74, 12.13, 0.58, "迎合的第二种形态：它顺着你给的前提往下推",
       27, True, C_TITLE, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
    tb(s, 0.6, 1.38, 12.13, 0.36,
       "同一次漏检：我问“这是不是农用地”，它答“是”——答对了问题，可没人去看表头写错了",
       16.5, False, C_MUTED, spacing=1.1)

    # ---------- 左：漏检经过 ----------
    tb(s, 0.6, 1.88, 6.0, 0.3, "这次漏检是怎么发生的", 15, True, C_TEAL)
    steps = [
        ("人", C_ORANGE, "表头写「农用地」，实际填「耕地」——差异只在第一列"),
        ("人", C_ORANGE, "我问：“草地、农村道路放在农用地里，是不是归属农用地？”"),
        ("AI", C_TEAL, "答“是”：按分类，这个判断本身没有错"),
        ("人", C_ORANGE, "据此判该列没问题——表头错位，整批漏检"),
    ]
    y = 2.22
    for tag, colr, text in steps:
        rect(s, 0.6, y, 0.34, 0.42, colr)
        tb(s, 0.6, y, 0.34, 0.42, tag, 9.5, True, C_WHITE,
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, 1.06, y, 5.55, 0.42, text, 12, False, C_TXT,
           anchor=MSO_ANCHOR.MIDDLE, spacing=1.1)
        y += 0.50

    # ---------- 右：两种迎合 ----------
    tb(s, 6.85, 1.88, 5.9, 0.3, "两种迎合，同一个根源", 15, True, C_TEAL)

    rect(s, 6.85, 2.22, 5.9, 0.80, C_WHITE, border=C_LINE)
    rect(s, 6.85, 2.22, 0.06, 0.80, C_ORANGE)
    tb(s, 7.05, 2.27, 5.5, 0.26, "① 顺着结论（第 15 页）", 11.5, True, C_ORANGE)
    tb(s, 7.05, 2.55, 5.5, 0.4, "你说“这里有错”，它就去找错——迎合的是你的期望",
       11.5, False, C_TXT, spacing=1.15)

    rect(s, 6.85, 3.10, 5.9, 0.80, C_WHITE, border=C_LINE)
    rect(s, 6.85, 3.10, 0.06, 0.80, C_TEAL)
    tb(s, 7.05, 3.15, 5.5, 0.26, "② 顺着前提（这一页）", 11.5, True, C_TEAL)
    tb(s, 7.05, 3.43, 5.5, 0.4, "你把“农用地”当事实写进问题，它就在这个前提里把话说完",
       11.5, False, C_TXT, spacing=1.15)

    tb(s, 6.85, 4.02, 5.9, 0.5,
       "根源是同一处：训练第三步按“人喜不喜欢”打分（第 11 页）；纠正用户前提的回答，人不会喜欢。",
       10.5, False, C_MUTED, spacing=1.15)

    # ---------- 中带：为什么它不会自己发现 ----------
    tb(s, 0.6, 4.50, 12.13, 0.28, "为什么它不会自己发现", 15, True, C_TEAL)
    cards = [
        ("前提顺从", "字段名、单位、口径长得像事实，它默认继承为真；这类“元信息”几乎从不被质疑"),
        ("范围即上限", "能力足够看出层级不一致，但问题只问“归不归属”，注意力就没有落点"),
        ("自洽掩盖错误", "错的是前提，推理全程正确，所以结论看起来毫无破绽"),
    ]
    x = 0.6
    for head, body in cards:
        rect(s, x, 4.82, 3.93, 0.78, C_BG_BEIGE)
        rect(s, x, 4.82, 0.06, 0.78, C_DARK)
        tb(s, x + 0.2, 4.87, 3.6, 0.26, head, 11.5, True, C_TITLE)
        tb(s, x + 0.2, 5.15, 3.6, 0.42, body, 10.5, False, C_TXT, spacing=1.15)
        x += 4.1

    # ---------- 类比条 ----------
    rect(s, 0.6, 5.72, 12.13, 0.46, C_BG_SOFT)
    rect(s, 0.6, 5.72, 0.06, 0.46, C_TEAL)
    tb(s, 0.8, 5.72, 0.7, 0.46, "同类", 10.5, True, C_TEAL,
       anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 1.55, 5.72, 11.0, 0.46,
       "同构的公开案例：OpenAI 一次放出 722 篇数学手稿，两天后撤回 3 篇——一处符号应记为 −1 却写成 +1，"
       "本该归零的计数不再为零，依赖该构造的另两篇一并撤回；出事的都没做 Lean 形式化。",
       10.5, False, C_TXT, anchor=MSO_ANCHOR.MIDDLE, spacing=1.05)

    # ---------- 结论 ----------
    rect(s, 0.6, 6.3, 12.13, 0.55, C_BG_WARN)
    rect(s, 0.6, 6.3, 0.09, 0.55, C_ORANGE)
    tb(s, 0.85, 6.3, 1.0, 0.55, "结论", 14, True, C_ORANGE, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 1.75, 6.3, 10.73, 0.55,
       "错的前提当成了共同知识：它顺着前提答，答得越顺，错得越深",
       18, True, C_TITLE, anchor=MSO_ANCHOR.MIDDLE)

    s.notes_slide.notes_text_frame.text = """讲解与会后阅读

迎合的第二种形态：它顺着你给的前提往下推

1. 第 15 页讲的是“顺着结论”的迎合：你暗示这里有错，它就替你找出错来。这一页是同一台机器的另一个出口——你说“这一列叫农用地”，它就默认这是真的，然后在这个前提里把推理做完。

2. 现场的经过：明细表第一列表头写“农用地”，实际填的是“耕地”。审核时我问“草地、农村道路放在农用地里，是不是归属农用地”，它答“是”。按分类口径这个答案没错，于是这一列被判为没问题，表头与取值域的错位整批漏了过去。

3. 为什么不怪它不认真：字段名、单位、口径这类信息，在语言里长得就像事实陈述，模型默认继承为真。它的“元信息批判”能力几乎没有被训练过，也很少被要求。另外，问题本身就是范围的边界：问的是“是不是归属”，它就在这个判定里作答，不会自己扩到“表头该写哪一级类”上去。

4. 越自洽越危险：错误在前提，推理全程正确，所以结论看起来毫无破绽。人复查时看到的是“它有依据、有逻辑”，也就不会再回头怀疑前提。

5. 一个同构的公开案例：OpenAI 于 2026-10 一次放出 722 篇数学手稿，两天后撤回 3 篇。原因是其中一篇在一处关键论证里把符号记成 +1，而按论文自己的约定应为 −1，使一个本该相互抵消归零的计数变成非零；所依赖的经典定理前提不成立，整套构造失效，另外两篇借用了这个构造，一并撤回。仓库总数由 722 降为 719，另有 14 篇修订。值得注意的对照：出事的部分都没有 Lean 形式化，做过逐行机器检查的部分没有出这类问题。这一点和本页讲的是同一件事——没有硬校验托底的地方，错误就藏在“自洽”里。

6. 汇报时要讲清的分寸：这不是“AI 不可靠”，而是“前提必须由人来核”。我给它的定义错了，它只会把错的定义执行得更彻底。换个人拿着同一个表头去看，一样会漏。

7. 与第 15 页的结论接上：不要在问题里替它下结论；也不要把自己的定义当成不用核的共同知识。

依据：本地审核记录（2026-10）；OpenAI math 仓库 history.md（2026-10-07，撤稿与修订说明）；OpenAI 2025-04-29 关于迎合的说明（对应第 15 页）；Anthropic arXiv 2310.13548。

查询与整理日期：2026-10-19。"""

    return s


def main():
    prs = Presentation(SRC)
    build(prs, 0)

    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    new_id = ids[-1]
    sldIdLst.remove(new_id)
    sldIdLst.insert(15, new_id)   # 0-based 15 → 插在第 15 页之后，成为第 16 页

    # 第 17 页（原 17 = 现在索引 17）小结表：迎合一行补“前提”
    s = prs.slides[17]
    patch = {
        "你说有错，它就找错": "你说有错它就找错；你给的前提它照用",
        "开放提问，要反例": "开放提问，要反例，先让它复述口径",
    }
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t in patch:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip() == t:
                        r.text = patch[t]
                        print("patched", t, "->", patch[t])

    # 插页后 16 页之后的所有内容 +1：更新指向正文页的交叉引用
    # 只改真正指向“第 22 页 公开榜单①”的三处；页注释里的“第 18 页”是原稿既有表述，不动
    ref_fix = {"第 22 页": "第 23 页"}
    ref_hits = 0
    for idx, sl in enumerate(prs.slides, 1):
        if idx <= 16:
            continue
        targets = []
        for sh in sl.shapes:
            if sh.has_text_frame:
                targets.append(sh.text_frame)
        if sl.has_notes_slide:
            targets.append(sl.notes_slide.notes_text_frame)
        for tf in targets:
            for p in tf.paragraphs:
                for r in p.runs:
                    for old, new in ref_fix.items():
                        if old in r.text:
                            r.text = r.text.replace(old, new)
                            ref_hits += 1
    print("cross-ref patched:", ref_hits)

    fixed = 0
    for i, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            if not sh.has_text_frame:
                continue
            if not sh.text_frame.text.strip().isdigit():
                continue
            if sh.left is None or sh.top is None:
                continue
            if sh.left / 914400 > 11.4 and sh.top / 914400 < 0.6:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.strip().isdigit():
                            r.text = str(i)
                fixed += 1
    print("page numbers:", fixed)

    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
