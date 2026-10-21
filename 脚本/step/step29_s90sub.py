# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[89]
S = {sh.shape_id: sh for sh in s.shapes}
sh = S[6]
for p in sh.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(15); r.font.bold = False
        r.font.color.rgb = RGBColor(0x5B,0x65,0x77)
        r.font.name = "微软雅黑"
prs.save(P)
print("ok", S[6].text_frame.text)
