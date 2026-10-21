# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[23]
kill = [sh for sh in s.shapes if sh.name in ("LogoBadge","Logo")]
for sh in kill: sh._element.getparent().remove(sh._element)
prs.save(P); print("removed", len(kill))
