# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
DIRS = [r"D:\MG\program\ai交流\成品\_work29\logos3", r"D:\MG\program\ai交流\成品\_work29\logos2", r"D:\MG\program\ai交流\成品\_work29\logos"]
def pick(n):
    for d in DIRS:
        p=os.path.join(d,n+".png")
        if os.path.exists(p): return p
    return None
prs = Presentation(P)
# 清掉上一轮只复制了圆角矩形、没带图的部分
for sn in (60,62,68,70):
    s = prs.slides[sn-1]
    for sh in list(s.shapes):
        if sh.shape_type != 4 and not sh.has_text_frame and sh.name.startswith("Rounded"):
            if abs(sh.top/360000-0.58)<0.15: sh._element.getparent().remove(sh._element); print("removed stray", sn, sh.name)
# 重新加
def badge(slide, path, left=27.23, top=0.58, size=1.88):
    sq = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(left), Cm(top), Cm(size), Cm(size))
    sq.fill.solid(); sq.fill.fore_color.rgb = RGBColor(0xFF,0xFF,0xFF)
    sq.line.color.rgb = RGBColor(0xE2,0xDF,0xD6); sq.line.width = Pt(0.75)
    sq.shadow.inherit = False; sq.name = "LogoBadge"
    try: sq.adjustments[0] = 0.22
    except Exception: pass
    pad = 0.155
    slide.shapes.add_picture(path, Cm(left+pad), Cm(top+pad), Cm(size-2*pad), Cm(size-2*pad)).name="Logo"
for sn, key in ((60,"openai"),(62,"anthropic"),(68,"deepseek"),(70,"qwen")):
    pth = pick(key)
    print(sn, key, bool(pth))
    if pth: badge(prs.slides[sn-1], pth)
prs.save(P)
