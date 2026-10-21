# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[24]
NEW = {
 "出字最快、最便宜；但先来一段“雷霆大思考”，自问自答完才给结果":
   "出字最快、最便宜；但先来一段“雷霆大思考”才给答案",
 "彻底顺着你：你说 A 对，它夸 A 对；你说 A 错，它立刻改口":
   "彻底顺着你：说 A 对就夸 A 对，说 A 错就立刻改口",
}
def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes); continue
        yield sh
for sh in walk(s.shapes):
    if not sh.has_text_frame: continue
    t = sh.text_frame.text
    if t in NEW:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.text = NEW[t]; break
            break
        print("ok:", NEW[t])
prs.save(P)
