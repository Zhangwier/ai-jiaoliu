# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

P  = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
LG = r"D:\MG\program\ai交流\成品\_work29\logos3"
LG4= r"D:\MG\program\ai交流\成品\_work29\logos4"
LG2= r"D:\MG\program\ai交流\成品\_work29\logos2"
import os
def pick(name):
    for d,f in ((LG,name+".png"),(LG2,name+".png"),(LG4,name+".png")):
        p=os.path.join(d,f)
        if os.path.exists(p): return p
    raise SystemExit("no logo "+name)

prs = Presentation(P)

def badge(slide, path, left, top, size, ring=RGBColor(0xE2,0xDF,0xD6)):
    ov = slide.shapes.add_shape(MSO_SHAPE.OVAL, Cm(left), Cm(top), Cm(size), Cm(size))
    ov.fill.solid(); ov.fill.fore_color.rgb = RGBColor(0xFF,0xFF,0xFF)
    ov.line.color.rgb = ring; ov.line.width = Pt(0.75)
    ov.shadow.inherit = False
    ov.name = "LogoBadge"
    pad = size*0.17
    pic = slide.shapes.add_picture(path, Cm(left+pad), Cm(top+pad), Cm(size-2*pad), Cm(size-2*pad))
    pic.name = "Logo"
    return ov, pic

def set_body(sh, txt, size=None, align=None):
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    runs = list(p0.runs)
    if runs:
        runs[0].text = txt
        for r in runs[1:]: r._r.getparent().remove(r._r)
    else:
        tf.text = txt
    if size:
        for r in tf.paragraphs[0].runs: r.font.size = Pt(size)
    if align is not None:
        for p in tf.paragraphs: p.alignment = align

# ============ SLIDE 25 ============
s = prs.slides[24]
S = {sh.shape_id: sh for sh in s.shapes}
ROWS = [
 (12,13,14,15, "GPT",      "gpt",
  "写代码偏“防御式”，兜底和异常处理写得很全；推理最硬",
  "Astra 是 FrontierMath Erdős 未解题唯一得分的模型"),
 (17,18,19,20, "Claude",   "claude",
  "最听得懂人话，一句话就能办对；会主动回去查证，也爱反过来纠正你",
  "报告案例 4 分钟出结果、金额全对、承建单位也查对"),
 (22,23,24,25, "Gemini",   "gemini",
  "知识面最广，中文表达最顺；但几乎不反对你，先夸再顺着说",
  "职业作家盲评榜前两名（2026-02，较旧）"),
 (27,28,29,30, "DeepSeek", "deepseek",
  "出字最快、最便宜；但先来一段“雷霆大思考”，自言自语完才给答案",
  "报告案例原因最具体，但有套错，耗时 34 分"),
 (32,33,34,35, "豆包",     "doubao",
  "彻底顺着你：你说 A 对就说 A 对，改口说 A 错，立刻也跟着改",
  "未进本次盲测，也没有同口径分数——只按日常使用记录"),
]
FONTS = {"gpt":"openai","openai":"openai"}
for (nid, rid, bid, eid, name, key, feel, ev) in ROWS:
    nm = S[nid]; nb = S[rid]
    top = nm.top; left = nm.left; w = nm.width; h = nm.height
    # 名称：字号 14，文本框下移给 logo 让位
    set_body(nm, name, size=14)
    nm.top = Cm(round(top/360000+0.52, 2)) if False else top + Cm(0.52)
    nm.height = h - Cm(0.52)
    # 说明文字
    set_body(S[bid], feel, size=14.5)
    set_body(S[eid], ev,   size=14.5)
    # logo
    sz = 0.62
    lb = left + (w - Cm(sz)) / 2
    badge(s, pick(FONTS.get(key, key)), lb/360000, (top + Cm(0.13))/360000, sz, ring=RGBColor(0xE2,0xDF,0xD6))

set_body(S[39], "同一个问题，五家交出来的东西不一样：脾气是真的，按任务挑，比按总分挑靠谱", size=17)
print("slide25 done")

# ============ SLIDE 20 ============
s = prs.slides[19]
S = {sh.shape_id: sh for sh in s.shapes}
ROWS20 = [(30,"ChatGPT","openai"),(36,"Claude","claude"),(41,"Gemini","gemini")]
for nid, label, key in ROWS20:
    nm = S[nid]
    top, left, w, h = nm.top, nm.left, nm.width, nm.height
    sz = 0.78
    badge(s, pick(key), (left+Cm(0.10))/360000, (top+(h-Cm(sz))/2)/360000, sz)
    nm.left = left + Cm(sz+0.16); nm.width = w - Cm(sz+0.20)
print("slide20 done")

# ============ SLIDE 26 ============
s = prs.slides[25]
S = {sh.shape_id: sh for sh in s.shapes}
ROWS26 = [(14,"WorkBuddy（腾讯）","tencent"),(20,"Trae 国内版（字节）","bytedance"),
          (25,"Qoder（阿里）","alibaba"),(31,"Cline","cline"),(36,"Codex（OpenAI）","codex"),
          (42,"Kiro（AWS）","aws"),(47,"Claude Code","claude")]
for nid, label, key in ROWS26:
    nm = S[nid]
    top, left, w, h = nm.top, nm.left, nm.width, nm.height
    sz = 0.72
    badge(s, pick(key), (left+Cm(0.02))/360000, (top+(h-Cm(sz))/2)/360000, sz)
    nm.left = left + Cm(sz+0.14); nm.width = w - Cm(sz+0.16)
    for p in nm.text_frame.paragraphs:
        for r in p.runs: r.font.size = Pt(14)
print("slide26 done")

prs.save(P)
