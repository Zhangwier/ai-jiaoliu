# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Pt
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
edits = [
 (90, 5, "E1b  工具一：多合一题，一张图看完“听不听话”"),
 (91, 5, "E1c  工具二：模型指纹，355 个数字里的“口音”"),
]
for idx, sid, txt in edits:
    for sh in prs.slides[idx].shapes:
        if sh.shape_id == sid:
            for p in sh.text_frame.paragraphs:
                if p.runs:
                    p.runs[0].text = txt
                    for r in p.runs[1:]: r.text = ""
            print(idx+1, sh.text_frame.text)
prs.save(P)
