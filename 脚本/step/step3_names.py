# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
def set_text(sh, new):
    tf = sh.text_frame; p = tf.paragraphs[0]; runs = list(p.runs)
    if not runs: tf.text = new; return
    runs[0].text = new
    for r in runs[1:]: r._r.getparent().remove(r._r)

B_NAME = "公司档案"
E_NAME = "降智测试与破限"
log=[]
for i,s in enumerate(prs.slides):
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        if "一页一家公司" in t:
            n = t.replace("一页一家公司", B_NAME)
            set_text(sh, n); log.append(f"s{i+1} B: {t[:40]!r} -> {n[:40]!r}")
        t2 = sh.text_frame.text
        if "案例、风险与破限" in t2:
            n = t2.replace("案例、风险与破限", E_NAME)
            set_text(sh, n); log.append(f"s{i+1} E: {t2[:40]!r} -> {n[:40]!r}")
prs.save(P)
print("\n".join(log))
