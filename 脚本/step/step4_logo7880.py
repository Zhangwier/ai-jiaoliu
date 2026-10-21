# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

P  = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
DIRS = [r"D:\MG\program\ai交流\成品\_work29\logos3",
        r"D:\MG\program\ai交流\成品\_work29\logos2",
        r"D:\MG\program\ai交流\成品\_work29\logos4"]
def pick(name):
    for d in DIRS:
        p=os.path.join(d,name+".png")
        if os.path.exists(p): return p
    return None
prs = Presentation(P)

def badge(slide, path, left, top, size, ring=RGBColor(0xE2,0xDF,0xD6)):
    ov = slide.shapes.add_shape(MSO_SHAPE.OVAL, Cm(left), Cm(top), Cm(size), Cm(size))
    ov.fill.solid(); ov.fill.fore_color.rgb = RGBColor(0xFF,0xFF,0xFF)
    ov.line.color.rgb = ring; ov.line.width = Pt(0.75); ov.shadow.inherit = False
    ov.name="LogoBadge"
    pad=size*0.17
    pic=slide.shapes.add_picture(path, Cm(left+pad), Cm(top+pad), Cm(size-2*pad), Cm(size-2*pad))
    pic.name="Logo"

CM = 360000
def add(sn, ids_keys, size=0.6, font=12.5, gap=0.10):
    s = prs.slides[sn-1]; S={sh.shape_id:sh for sh in s.shapes}; n=0
    for sid,key in ids_keys:
        sh = S.get(sid)
        if sh is None or not sh.has_text_frame: print("miss",sn,sid); continue
        pth = pick(key)
        if not pth: print("no logo", key); continue
        top,left,w,h = sh.top, sh.left, sh.width, sh.height
        badge(s, pth, (left+Cm(0.06))/CM, (top+(h-Cm(size))/2)/CM, size)
        sh.left = left + Cm(size+gap); sh.width = w - Cm(size+gap+0.02)
        for p in sh.text_frame.paragraphs:
            for r in p.runs: r.font.size = Pt(font)
        n+=1
    print(sn, "badges", n)

add(78, [(14,"openai"),(19,"claude"),(23,"gemini"),(28,"grok"),(32,"meta"),(37,"trae"),(41,"zhipu")], size=0.6, font=12.5)
add(79, [(18,"tencent"),(22,"moonshot"),(27,"claude"),(31,"openai"),(36,"gemini"),(40,"grok"),(45,"meta")], size=0.6, font=12.5)
add(80, [(14,"openai"),(19,"claude"),(23,"gemini"),(28,"moonshot"),(37,"grok"),(41,"meta"),
         (54,"tencent"),(59,"bytedance"),(63,"alibaba"),(68,"openai"),(72,"claude"),(77,"cline"),(81,"zhipu")], size=0.52, font=11.5, gap=0.08)
prs.save(P)
