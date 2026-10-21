# -*- coding: utf-8 -*-
import sys
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)

def set_text(sh, new):
    tf = sh.text_frame
    p = tf.paragraphs[0]
    runs = list(p.runs)
    if not runs:
        tf.text = new; return True
    runs[0].text = new
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return True

EDITS = {
 (5, 5):  "分享分四部分，附录收细节与出处",
 (5, 26): "附录 · 速查",
 (5, 32): "配套文件夹里有原始 HTML、案例材料和来源索引，随时可翻",
 (51, 18): "附录  常用术语 A1–A4 · 一页一家公司 B1–B13 · 订阅与省流 C1–C7 · 资讯渠道 D1–D4 · 降智测试与破限 E1–E4 · 原理与历史 F1–F9 · 来源 Z1–Z3",
 (52, 15): "附录 A–F（术语、一页一家公司、订阅与省流、资讯渠道、降智测试与破限、原理与历史），附录 Z 为来源索引；配套文件夹里有十份鹈鹕原始 HTML、图纸与报告的核对材料。",
 (90, 5):  "E1  降智测试：四道题，各测一件事",
 (91, 5):  "E1b  工具一 · 多合一题：一张图看完“听不听话”",
 (92, 5):  "E1c  工具二 · 模型指纹：355 个数字里的“口音”",
 (89, 5):  "共 16 页：E1 降智测试 · E1b、E1c 两个自查工具 · E2 中转站与代理池 · E3 护栏与破限 · E4 开源平台对照",
}
B = {
 59: "B1  OpenAI：以推理能力为核心的产品路线",
 60: "B1b  OpenAI 数学：从公开失误到可验证的推理进展",
 61: "B2  Anthropic：以智能体与工程任务为主战场",
 62: "B2b  Opus 生成视频：以代码逐帧渲染的实现路径",
 63: "B3  Google：知识覆盖与长文本输出最稳定",
 64: "B4  xAI：不以刷榜为目标，产品先做常驻型智能体",
 65: "B5  Meta：把长期记忆作为模型能力建设",
 66: "B6  Mistral：以可自部署为核心的欧洲路线",
 67: "B7  DeepSeek：以极致定价重构成本结构",
 68: "B7b  DeepSeek 的开源范围：开放了什么、保留了什么",
 69: "B8  阿里 Qwen：以工具调用与自主纠错见长",
 70: "B8b  Qwen 的尺寸矩阵：各参数档位全面铺开",
 71: "B9  智谱 GLM：开源模型中代码能力突出的一档",
 72: "B10  月之暗面 Kimi：超长上下文与长文档处理",
 73: "B11  MiniMax：语言与视频双线并行",
 74: "B12  腾讯混元 Hy：办公场景适配度较高的国产一档",
 75: "B13  百度文心：中文检索与生态整合最完整的一档",
}

log = []
for (sn, sid), new in EDITS.items():
    s = prs.slides[sn-1]
    hit = [sh for sh in s.shapes if sh.shape_id == sid]
    if not hit: log.append(f"MISS s{sn} id{sid}"); continue
    log.append(f"OK s{sn} id{sid}: {hit[0].text_frame.text[:36]!r} -> {new[:36]!r}")
    set_text(hit[0], new)

for sn, new in B.items():
    s = prs.slides[sn-1]
    done = False
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text[:2].rstrip("b") in ("B1","B2","B3","B4","B5","B6","B7","B8","B9") or (sh.has_text_frame and sh.text_frame.text.startswith(("B10","B11","B12","B13"))):
            log.append(f"OK s{sn}: {sh.text_frame.text[:40]!r} -> {new[:40]!r}")
            set_text(sh, new); done = True; break
    if not done: log.append(f"MISS B s{sn}")

# 全局「附录  会后速查」
n = 0
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == "附录  会后速查":
            set_text(sh, "附录  速查"); n += 1
log.append(f"footnote replaced: {n}")
prs.save(P)
print("\n".join(log))
