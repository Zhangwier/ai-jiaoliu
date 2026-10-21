# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Cm, Pt
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[79]
S = {sh.shape_id: sh for sh in s.shapes}
sh = S[32]
sh.text_frame.text = "豆包*/千问*"
for p in sh.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(10.5)
sh.left = Cm(2.72); sh.width = Cm(2.79)
# Hark 行与上面几行文字左端对齐（2.43）
S[46].left = Cm(2.43); S[46].width = Cm(3.06)
prs.save(P)
print("done")
