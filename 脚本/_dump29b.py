import sys, json
from pptx import Presentation
p = r"D:\MG\program\ai交流\成品\AI交流分享-成品29 .pptx"
prs = Presentation(p)
out=[]
for i, s in enumerate(prs.slides, 1):
    if i < 58: continue
    texts=[]
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.replace("\n"," | "))
    t = " || ".join(texts)
    out.append(f"[{i}] {t[:500]}")
print("\n".join(out))
