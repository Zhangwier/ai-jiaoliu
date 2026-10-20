from pptx import Presentation
from pptx.util import Emu
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
out=[]
for i, s in enumerate(prs.slides, 1):
    out.append(f"\n===== SLIDE {i} =====")
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                out.append(f"[{sh.name}] {t.replace(chr(10),' | ')}")
        elif sh.shape_type == 13:
            out.append(f"[{sh.name}] <PIC>")
open(r"D:\MG\program\ai交流\成品\_work29\alltext.txt","w",encoding="utf-8").write("\n".join(out))
print("ok", len(out))
