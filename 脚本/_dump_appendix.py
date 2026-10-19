# -*- coding: utf-8 -*-
"""Dump all runs of appendix pages (divider + content titles) for relabel planning."""
from pptx import Presentation

p = Presentation("AI交流分享-成品26.pptx")
for i in list(range(57, 115)):
    s = p.slides[i - 1]
    print("=" * 8, "SLIDE", i)
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text
        if not txt.strip():
            continue
        runs = []
        for pa in sh.text_frame.paragraphs:
            runs.append("".join(r.text for r in pa.runs))
        joined = " ⏎ ".join(runs)
        print(f"  [{sh.name}] @{round(sh.left/914400,2)},{round(sh.top/914400,2)} :: {joined[:160]}")
