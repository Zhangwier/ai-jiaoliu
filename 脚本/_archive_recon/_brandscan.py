from pptx import Presentation
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
BR = ["ChatGPT","GPT","OpenAI","Claude","Anthropic","Gemini","Google","DeepSeek","Qwen","千问","阿里","GLM","智谱","Kimi","月之暗面","MiniMax","混元","Hy","文心","百度","豆包","字节","Doubao","Grok","xAI","Meta","Mistral","Step","阶跃","MiMo","小米","Muse","Fable","Cursor","Cline","Trae","WorkBuddy","Qoder","Codex","Kiro","OpenCode","Devin","Jules","Hark","Gemma","Llama","Sora","Veo","Nano Banana","Seed","Hermes","OpenClaw","Copilot","Windsurf","Aider"]
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text.strip()
        if not t or len(t) > 16: continue
        if any(b.lower() in t.lower() for b in BR):
            print(i, "|", sh.name, "|", t)
