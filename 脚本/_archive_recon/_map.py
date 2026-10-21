# -*- coding: utf-8 -*-
import sys
from pptx import Presentation
from pptx.util import Emu
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
CM = 360000.0
def dump(idx):
    s = prs.slides[idx-1]
    print(f"\n===== slide {idx} =====")
    for sh in s.shapes:
        t = ""
        if sh.has_text_frame:
            t = sh.text_frame.text.replace("\n","/")
        sz = []
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size: sz.append(round(r.font.size.pt,1))
        f = ""
        try:
            if sh.fill.type == 1: f = "fill=%02X%02X%02X" % (sh.fill.fore_color.rgb[0],sh.fill.fore_color.rgb[1],sh.fill.fore_color.rgb[2]) if sh.fill.fore_color.type==1 else ""
        except Exception: f=""
        print(f"[{sh.shape_id}] {sh.name[:22]:22} l{sh.left/CM:6.2f} t{sh.top/CM:5.2f} w{sh.width/CM:6.2f} h{sh.height/CM:5.2f} {f} sz{sz[:3]} | {t[:90]}")
for i in [int(x) for x in sys.argv[1:]]:
    dump(i)
