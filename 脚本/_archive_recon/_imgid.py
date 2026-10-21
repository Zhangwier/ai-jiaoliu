# -*- coding: utf-8 -*-
import hashlib, glob, os
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
def h(b): return hashlib.sha1(b).hexdigest()[:10]
ai = {}
for f in glob.glob(r"D:\MG\program\ai交流\素材\AI娘\*.jpg"):
    ai[h(open(f,'rb').read())] = os.path.basename(f)
print("AI娘素材哈希:")
for k,v in ai.items(): print("  ",k,v)
print()
for idx in [59,61,63,64,65,66,67,69,71,72,73,74,75]:
    s = prs.slides[idx-1]
    for sh in s.shapes:
        if sh.shape_type == 13:
            try:
                b = sh.image.blob
            except Exception as e:
                print(idx, "ERR", e); continue
            if sh.width > 360000*5:
                print(f"s{idx} Pic{sh.shape_id} {sh.width/360000:.1f}x{sh.height/360000:.1f} {h(b)} -> {ai.get(h(b),'?')}")
