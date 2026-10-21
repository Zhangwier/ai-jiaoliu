import io, os, hashlib
from pptx import Presentation
from pptx.util import Emu
from PIL import Image
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
out = r"D:\MG\program\ai交流\成品\_work29\logos"
os.makedirs(out, exist_ok=True)
names = {58:"OpenAI",60:"Anthropic",62:"Google",63:"xAI",64:"Meta",65:"Mistral",66:"DeepSeek",
         68:"Qwen",70:"GLM",71:"Kimi",72:"MiniMax",73:"Hy",74:"Wenxin"}
for idx in sorted(names):
    s = prs.slides[idx]
    hit = None
    for sh in s.shapes:
        if sh.shape_type == 13 or sh.__class__.__name__ == "Picture":
            L = Emu(sh.left).cm; T = Emu(sh.top).cm
            if L > 26 and T < 2.0:
                hit = sh
    if not hit:
        print(names[idx], "NO LOGO PIC"); continue
    blob = hit.image.blob
    ext = hit.image.ext
    fn = os.path.join(out, f"{names[idx]}.{ext}")
    open(fn, "wb").write(blob)
    try:
        im = Image.open(io.BytesIO(blob))
        print(f"slide{idx+1} {names[idx]:<10} {im.size} {im.mode} {len(blob)}B -> {fn}")
    except Exception as e:
        print(idx, e)
