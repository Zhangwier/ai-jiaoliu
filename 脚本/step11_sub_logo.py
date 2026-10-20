# -*- coding: utf-8 -*-
import copy
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
PAIRS = {60: 59, 62: 61, 68: 67, 70: 69}
for child, parent in PAIRS.items():
    sp = prs.slides[parent-1]
    geom = [sh for sh in sp.shapes if abs(sh.top/360000-0.58)<0.12 and abs(sh.left/360000-27.23)<0.2 and sh.shape_type!=MSO_SHAPE_TYPE.PICTURE]
    pic  = [sh for sh in sp.shapes if sh.shape_type==MSO_SHAPE_TYPE.PICTURE and abs(sh.top/360000-0.71)<0.25 and sh.left/360000>26]
    if len(geom)!=1 or len(pic)!=1:
        print("s%d parent mismatch" % child, len(geom), len(pic)); continue
    cs = prs.slides[child-1]
    for sh in (geom[0], pic[0]):
        cs.shapes._spTree.append(copy.deepcopy(sh._element))
    print("s%d <- logo from s%d ok (%s)" % (child, parent, pic[0].name))
prs.save(P)
