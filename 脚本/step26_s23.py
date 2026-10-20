# -*- coding: utf-8 -*-
import os, re
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
DIRS = [r"D:\MG\program\ai交流\成品\_work29\logos3", r"D:\MG\program\ai交流\成品\_work29\logos4", r"D:\MG\program\ai交流\成品\_work29\logos2", r"D:\MG\program\ai交流\成品\_work29\logos"]
def pick(n):
    for d in DIRS:
        p=os.path.join(d,n+".png")
        if os.path.exists(p): return p
prs = Presentation(P)
s = prs.slides[22]
S = {sh.shape_id: sh for sh in s.shapes}
SIZE=0.50
def badge(sl, pth, x, top):
    ov = sl.shapes.add_shape(MSO_SHAPE.OVAL, Cm(x), Cm(top), Cm(SIZE), Cm(SIZE))
    ov.fill.solid(); ov.fill.fore_color.rgb = RGBColor(0xFF,0xFF,0xFF)
    ov.line.color.rgb = RGBColor(0xE2,0xDF,0xD6); ov.line.width = Pt(0.75)
    ov.shadow.inherit=False; ov.name="LogoBadge"
    pad=SIZE*0.17
    sl.shapes.add_picture(pth, Cm(x+pad), Cm(top+pad), Cm(SIZE-2*pad), Cm(SIZE-2*pad)).name="Logo"
MAP = {16:"claude", 22:"openai", 27:"openai", 33:"openai", 38:"claude", 44:"gemini", 49:"claude"}
for sid, key in MAP.items():
    sh = S[sid]
    top = (sh.top + (sh.height - Cm(SIZE))/2)/360000
    x = 15.62
    pth = pick(key)
    if not pth:
        print("missing", key); continue
    badge(s, pth, x, top)
    old_l = sh.left/360000; old_w = sh.width/360000
    sh.left = Cm(x+SIZE+0.12); sh.width = Cm(old_w - (x+SIZE+0.12-old_l))
prs.save(P)
print("ok")
