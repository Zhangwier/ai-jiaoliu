# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
for idx in (78,):
    s = prs.slides[idx]
    for sh in s.shapes:
        t = sh.text_frame.text.replace("\n","\\n")[:40] if sh.has_text_frame else ""
        if t: print(f"{sh.shape_id:<3} l={sh.left/360000:6.2f} t={sh.top/360000:6.2f} w={sh.width/360000:6.2f} {t!r}")
