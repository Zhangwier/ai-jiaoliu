# -*- coding: utf-8 -*-
"""第三轮：按“结论条不要每页都用强断言、标题更短”的原则做收尾润色。"""
import sys
from pptx import Presentation
from pptx.util import Pt

EDITS = [
    # (页序号0基, 形状名, 新文本)
    (9,  "TextBox 36", "按“人是否偏好”打分来训练，与后面提到的“迎合”有关"),
    (21, "TextBox 49", "本次榜单上：推理与数学靠 Astra，写代码与干活靠 Claude，聊天体感 Gemini、Muse 靠前"),
    (37, "TextBox 4",  "合同条文：改写容易删掉条件"),
    (45, "TextBox 51", "本稿经验：工作流首先要让人能验收，而不是先让模型听话"),
    (81, "TextBox 41", "自己搭网关，是目前比较能把“谁用了多少、数据去了哪”说清楚的做法"),
    (19, "TextBox 4",  "实际效果取决于模型与外壳两件事"),
    (46, "TextBox 4",  "多个 Agent 协作：什么时候值，什么时候更贵"),
    (30, "TextBox 4",  "看不见图就找不到：只靠 OCR 的配置零命中"),
]


def main(src, dst):
    prs = Presentation(src)
    n = 0
    for idx, name, text in EDITS:
        for sh in prs.slides[idx].shapes:
            if sh.name == name and sh.has_text_frame:
                p = sh.text_frame.paragraphs[0]
                p.runs[0].text = text
                for r in p.runs[1:]:
                    r._r.getparent().remove(r._r)
                print("ok", idx + 1, name)
                n += 1
                break
        else:
            print("! MISS", idx + 1, name)
    prs.save(dst)
    print("applied", n, "->", dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
