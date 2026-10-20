from pptx import Presentation
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
s = prs.slides[23]
for sh in s.shapes:
    if sh.shape_id in (9,13,17,21):
        print("cell", sh.shape_id, round(sh.left/360000,2), round(sh.top/360000,2), round(sh.width/360000,2), round(sh.height/360000,2), sh.fill.fore_color.rgb if sh.fill.type==1 else sh.fill.type)
    if sh.shape_id in (10,14,18,22):
        print("  txt", sh.shape_id, round(sh.left/360000,2), round(sh.width/360000,2), repr(sh.text_frame.text))
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                print("     ", repr(r.text), r.font.size.pt if r.font.size else None, r.font.color.rgb if r.font.color and r.font.color.type is not None else None)
print("=== s21 rep column fonts")
s = prs.slides[20]
for sh in s.shapes:
    if sh.shape_id in (38,43,48,52,56,60):
        print(sh.shape_id, round(sh.left/360000,2), round(sh.width/360000,2), round(sh.height/360000,2), repr(sh.text_frame.text))
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                print("   ", repr(r.text), r.font.size.pt if r.font.size else None, r.font.name)
