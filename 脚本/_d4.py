from pptx import Presentation
from pptx.util import Emu
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
for i in [59,61,67,69]:
    s = prs.slides[i]
    print("### SLIDE", i+1)
    for sh in s.shapes:
        if sh.shape_type is not None and ("Picture" in sh.name or "Rounded" in sh.name or sh.shape_type==1):
            print("  ", sh.shape_id, sh.name, round(Emu(sh.left).cm,2), round(Emu(sh.top).cm,2), round(Emu(sh.width).cm,2), round(Emu(sh.height).cm,2))
