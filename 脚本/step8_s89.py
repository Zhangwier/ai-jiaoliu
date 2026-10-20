# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Pt
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[88]
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("附录"):
        ps = sh.text_frame.paragraphs
        # p0 -> '附录', p1 -> '分享与工具'
        p0, p1 = ps[0], ps[1]
        r0 = p0.runs[0]; r0.text = "附录"
        for r in p0.runs[1:]: r._r.getparent().remove(r._r)
        r1 = p1.runs[0]; r1.text = "分享与工具"
        for r in p1.runs[1:]: r._r.getparent().remove(r._r)
        print("title fixed:", [p.text for p in ps])
    if sh.has_text_frame and sh.text_frame.text.startswith("E  "):
        r = sh.text_frame.paragraphs[0].runs[0]
        r.text = "E  分享与工具（十六页）"
        print("E label:", r.text)
    if sh.has_text_frame and sh.text_frame.text.startswith("共 16 页"):
        r = sh.text_frame.paragraphs[0].runs[0]
        r.text = "共 16 页，都可以直接拿去用：E1 三项测试（一项降智测试、两项自查工具）· E2 中转站与代理池 · E3 护栏与破限 · E4 开源平台对照"
        print("sub:", r.text[:30])
prs.save(P)
