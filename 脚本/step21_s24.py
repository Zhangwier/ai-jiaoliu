# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
DIRS = [r"D:\MG\program\ai交流\成品\_work29\logos3", r"D:\MG\program\ai交流\成品\_work29\logos4", r"D:\MG\program\ai交流\成品\_work29\logos2", r"D:\MG\program\ai交流\成品\_work29\logos"]
def pick(n):
    for d in DIRS:
        p=os.path.join(d,n+".png")
        if os.path.exists(p): return p
prs = Presentation(P)
s = prs.slides[23]
SIZE=0.52; GAP=0.08
def badge(sl, pth, x, top):
    ov = sl.shapes.add_shape(MSO_SHAPE.OVAL, Cm(x), Cm(top), Cm(SIZE), Cm(SIZE))
    ov.fill.solid(); ov.fill.fore_color.rgb = RGBColor(0xFF,0xFF,0xFF)
    ov.line.color.rgb = RGBColor(0xE2,0xDF,0xD6); ov.line.width = Pt(0.75)
    ov.shadow.inherit=False; ov.name="LogoBadge"
    pad=SIZE*0.17
    sl.shapes.add_picture(pth, Cm(x+pad), Cm(top+pad), Cm(SIZE-2*pad), Cm(SIZE-2*pad)).name="Logo"
blocks = [(9,["openai","anthropic"]),(13,["google","meta"]),(17,["grok","alibaba","zhipu","moonshot"]),(21,["tencent","minimax","doubao"])]
S = {sh.shape_id: sh for sh in s.shapes}
for rid, keys in blocks:
    rect = S[rid]
    L = rect.left/360000; T = rect.top/360000; W = rect.width/360000
    n = len(keys); total = n*SIZE + (n-1)*GAP
    x = L + W - 0.25 - total; y = T + 0.14
    for k in keys:
        pth = pick(k)
        if pth: badge(s, pth, x, y)
        else: print("missing", k)
        x += SIZE+GAP
prs.save(P)
print("ok")
