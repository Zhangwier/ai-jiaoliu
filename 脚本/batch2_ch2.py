# -*- coding: utf-8 -*-
"""batch2: 第二章挂钩主题
 1) s18 分隔页副标 + 右栏目录
 2) 新增「十种结果的三个来源」页（插在 s19 之后）
 3) s21 结论条呼应
"""
import sys, copy
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

prs = openprs()
def S(n): return prs.slides[n-1]
def g(sl, n): return find(sl, name=n)

# ---- 1) s18 分隔页 ----
s = S(18)
settext(g(s,"TextBox 4"), "这一章回答一个问题：同一个问题，为什么别人用着好用、你用了就想放弃。差别只有三处：模型的档次、模型的专精、外面的壳。", size=14)
items = ["鹈鹕盲测：十个配置，十种结果","差别一：能力有档次","差别二：各有专精",
         "差别三：外壳决定能干什么","同一牌子，免费版和 API 不是一个模型","榜单与全球格局",
         "各家脾气 · 免费额度"]
boxes = [sh for sh in s.shapes if sh.has_text_frame and abs(sh.left/CM-23.11)<0.15 and 4.9 < sh.top/CM < 15.2]
boxes.sort(key=lambda x: x.top)
print("s18 boxes", len(boxes))
while len(boxes) < len(items):
    el = copy.deepcopy(boxes[-1]._element); s.shapes._spTree.append(el)
    boxes = [sh for sh in s.shapes if sh.has_text_frame and abs(sh.left/CM-23.11)<0.15 and 4.9 < sh.top/CM < 15.2]
    boxes.sort(key=lambda x: x.top)
y = 5.05
for i, txt in enumerate(items):
    b = boxes[i]; b.top = Cm(y); b.height = Cm(1.14); b.left = Cm(23.11); b.width = Cm(9.92)
    settext(b, txt, size=15.5); y += 1.24
for extra in boxes[len(items):]:
    extra._element.getparent().remove(extra._element)

# ---- 2) 新增页 ----
NEW_AFTER = 19
lay = prs.slides[NEW_AFTER-1].slide_layout
s = prs.slides.add_slide(lay)
for sh in list(s.shapes): sh._element.getparent().remove(sh._element)

def tb(sl,x,y,w,h,text,size,color=NAVY,bold=False,align=None,anchor=None,space=1.18,name=None):
    b = sl.shapes.add_textbox(Cm(x),Cm(y),Cm(w),Cm(h))
    tf = b.text_frame; tf.word_wrap=True
    if anchor is not None: tf.vertical_anchor = anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    for i,ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if align is not None: p.alignment = align
        p.line_spacing = space
        r = p.add_run(); r.text = ln
        r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name="微软雅黑"
    if name: b.name=name
    return b

def rect(sl,x,y,w,h,fill=None,line=None,name=None):
    r = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x),Cm(y),Cm(w),Cm(h))
    if fill is None: r.fill.background()
    else: r.fill.solid(); r.fill.fore_color.rgb=fill
    if line is None: r.line.fill.background()
    else: r.line.color.rgb=line; r.line.width=Pt(0.75)
    r.shadow.inherit=False
    if name: r.name=name
    return r

tb(s,1.52,0.81,15.24,0.76,"02  模型与工具",12,GREY)
tb(s,29.80,0.81,2.54,0.76,"20",12,RGBColor(0x6B,0x73,0x85),align=PP_ALIGN.RIGHT)
rect(s,1.52,1.68,1.27,0.15,fill=TEAL)
tb(s,1.52,1.98,30.82,1.55,"十种结果的三个来源：档次、专精、外壳",27,NAVY,True)
tb(s,1.52,3.55,30.82,0.95,"鹈鹕不是“谁画得好看”，是同一份任务下，三种差别叠出来的结果；先认清是哪一种，再决定要不要换",16.5,GREY,space=1.1)
conn = s.shapes.add_connector(1,Cm(1.52),Cm(17.78),Cm(32.34),Cm(17.78)); conn.line.color.rgb=RGBColor(0xE2,0xDF,0xE6); conn.line.width=Pt(0.75)
tb(s,1.52,17.93,26.67,0.76,"来源：第 19 页十份鹈鹕实测；AA 智能指数（2026-10-18 查询）；本次分享的客户端实测记录",8.5,GREY)
tb(s,27.26,17.93,5.08,0.76,"来源索引 ›",8.5,GREY)

cols = [
 ("01  档次差别", TEAL, "同一条产品线上，分数真不一样",
  "GPT-6.1 Sol 与 GPT-6 Luna 都是 GPT，AA 智能指数差 20 分以上；Opus 5.5 与 Haiku 5.5 也是同理。\n\n"
  "表现：能做的事一样，做对的概率不一样。\n\n"
  "结论：别用最便宜的那一档去判断“AI 行不行”。", "换档位，是最直接的提升"),
 ("02  专精差别", ORANGE, "同一档里，各有各的长处",
  "Claude 强在长任务和写代码；Gemini 强在长文与中文；DeepSeek 强在便宜快；豆包强在顺着你。\n\n"
  "表现：同一个任务，有人 4 分钟做完，有人 34 分钟还套错。\n\n"
  "结论：别拿 A 的短板去否定 B 的长处。", "按任务挑，比按总分挑划算"),
 ("03  外壳差别", NAVY, "同一个模型，换个壳就换个人",
  "能不能读图、能不能改文件、能不能跑脚本、能连续干多久——这些不写在模型里，写在外壳里。\n\n"
  "表现：图纸案例里，只靠 OCR 的配置零命中；能自己把 PDF 渲染成图的配置全部命中。\n\n"
  "结论：模型一样时，先换壳。", "壳决定它能干什么，模型决定干得好不好"),
]
x = 1.52
W = 9.86; GAP = 0.62
for i,(tag,color,head,body,foot) in enumerate(cols):
    xx = x + i*(W+GAP)
    rect(s,xx,4.72,W,10.20,fill=WHITE)
    rect(s,xx,4.72,W,1.10,fill=color)
    tb(s,xx+0.28,4.72,W-0.56,1.10,tag,14,WHITE,True,anchor=MSO_ANCHOR.MIDDLE)
    tb(s,xx+0.28,6.02,W-0.56,0.95,head,16.5,NAVY,True)
    tb(s,xx+0.28,7.05,W-0.56,5.30,body,12.5,NAVY,space=1.22)
    rect(s,xx+0.28,14.10,W-0.56,0.78,fill=CREAM)
    tb(s,xx+0.45,14.10,W-0.90,0.78,foot,12,GREY,anchor=MSO_ANCHOR.MIDDLE)

rect(s,1.52,16.05,30.82,1.30,fill=MINT)
rect(s,1.52,16.05,0.23,1.30,fill=TEAL)
tb(s,2.16,16.05,2.54,1.30,"结论",14,NAVY,anchor=MSO_ANCHOR.MIDDLE)
tb(s,4.45,16.05,27.26,1.30,"觉得“AI 不行”之前，先按这个顺序查一遍：档位太低、用错了专精、壳不匹配",17,NAVY,anchor=MSO_ANCHOR.MIDDLE)
move_to(prs, s, NEW_AFTER)

# ---- 3) s21 结论 ----
s = S(22)
for sh in s.shapes:
    if sh.has_text_frame and "模型与运行环境都会影响结果" in sh.text_frame.text:
        settext(sh, "模型决定上限，外壳决定它此刻能干什么；比较两个结果时，这两样都要写清楚", size=17)

save(prs); print("ok")
