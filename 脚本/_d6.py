from pptx import Presentation
from pptx.util import Emu, Pt
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
print("=== E page titles")
for i in range(92,105):
    for sh in prs.slides[i].shapes:
        if sh.has_text_frame and sh.text_frame.text[:2] in ("E1","E2","E3","E4"):
            print(f"s{i+1}: {sh.text_frame.text!r}")
print("=== s58 separator")
for sh in prs.slides[57].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print(" ", sh.shape_id, sh.name, repr(sh.text_frame.text[:120]))
def fmt(sh):
    out=[]
    for para in sh.text_frame.paragraphs:
        for r in para.runs:
            out.append(f"sz={r.font.size.pt if r.font.size else None} b={r.font.bold} c={r.font.color.rgb if r.font.color and r.font.color.type is not None else None} name={r.font.name}")
    return out
print("=== s20 name cells")
for sid in (30,36,41):
    for sh in prs.slides[19].shapes:
        if sh.shape_id==sid: print(" ", sh.text_frame.text, fmt(sh)[:2])
print("=== s26 name cells")
for sid in (14,20,25,31,36,42,47):
    for sh in prs.slides[25].shapes:
        if sh.shape_id==sid: print(" ", sh.text_frame.text, fmt(sh)[:2])
print("=== s25 name cells")
for sid in (12,17,22,27,32):
    for sh in prs.slides[24].shapes:
        if sh.shape_id==sid:
            print(" ", sh.text_frame.text, fmt(sh)[:2], "anchor", sh.text_frame.vertical_anchor, "wrap", sh.text_frame.word_wrap, "ml", sh.text_frame.margin_left)
