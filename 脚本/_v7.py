from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
for sn in (53,58,76,84,89,106,116):
    s = prs.slides[sn-1]
    print("## s",sn)
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            sz = sh.text_frame.paragraphs[0].runs
            print("   id%-3d sz%-5s %s" % (sh.shape_id, sz[0].font.size.pt if sz and sz[0].font.size else "-", repr(sh.text_frame.text[:80])))
