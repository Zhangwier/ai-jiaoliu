# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Cm
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[82]
S = {sh.shape_id: sh for sh in s.shapes}
# 删除 s83 三栏标题的 badge/logo
kill=[sh for sh in s.shapes if sh.name in ("LogoBadge","Logo")]
for sh in kill: sh._element.getparent().remove(sh._element)
print("removed", len(kill))
# 恢复文字 left/width 原值
for sid, l, w in ((12,2.16,8.76),(22,12.55,8.76),(32,22.94,8.76)):
    S[sid].left = Cm(l); S[sid].width = Cm(w)
prs.save(P)
print("reverted s83")
