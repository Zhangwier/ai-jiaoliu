# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
bad = ["会后","散会","一页一家公司","案例、风险与破限","降智测试与破限","讲完","结束后","回头再","以后再查","稍后","待会"]
hits = {}
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        for b in bad:
            if b in t:
                hits.setdefault(b, []).append((i, t.replace("\n","\\n")[:70]))
for b, v in hits.items():
    print("##", b, len(v))
    for i, t in v[:8]: print("  ", i, t)
print("total flagged", sum(len(v) for v in hits.values()))
