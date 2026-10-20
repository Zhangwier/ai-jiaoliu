# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
for i,s in enumerate(prs.slides):
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text
            for k in ("降智测试与破限","案例、风险与破限","分享与工具"):
                if k in t:
                    print(f"s{i+1} id{sh.shape_id} [{k}] {t[:80]!r}")
