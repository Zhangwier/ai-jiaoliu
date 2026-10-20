# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[23]
for sh in s.shapes:
    if sh.name in ("LogoBadge","Logo"):
        print(f"{sh.name:<10} l={sh.left/360000:6.2f} t={sh.top/360000:6.2f} w={sh.width/360000:5.2f}")
