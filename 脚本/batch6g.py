import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm, Pt
prs = openprs(); s = prs.slides[85]
def g(n): return find(s, name=n)
def rw(nm, lines, size=12.5, ls=1.40):
    sh = g(nm); tf = sh.text_frame
    for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs)[1:]: r._r.getparent().remove(r._r)
    r0 = p0.runs[0]; r0.text = lines[0]
    r0.font.size=Pt(size); r0.font.bold=False; r0.font.color.rgb=RGBColor(0x2B,0x34,0x45); r0.font.name="微软雅黑"
    p0.line_spacing = ls
    for ln in lines[1:]:
        p=tf.add_paragraph(); p.line_spacing=ls
        r=p.add_run(); r.text=ln
        r.font.size=Pt(size); r.font.bold=False; r.font.color.rgb=RGBColor(0x2B,0x34,0x45); r.font.name="微软雅黑"

rw("TextBox 13", ["Cline：首月 $1（之后 $9.99）", "Codex：Go 档 $8/月", "GLM Coding：Lite ¥118/月", "首月半价：先看续费价"])
rw("TextBox 20", ["WorkBuddy：500 积分＋签到", "Trae：每天 100 次（会员 200）", "Qoder：2 周试用＋300 次", "Codex 免费号：只有小模型"])
rw("TextBox 27", ["开源中转、公益站（L 站常见）", "闲鱼共享号、代充：几元/月", "提示词与资料过别人的服务器", "维护靠人，随时可能停"])
for nm in ("TextBox 13","TextBox 20","TextBox 27"):
    g(nm).top = Cm(6.45); g(nm).height = Cm(2.75)
save(prs); print("ok")
