# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE

W,H = 13.3333, 7.5
PAPER = RGBColor(0xF6,0xF4,0xEF)
INK   = RGBColor(0x1A,0x22,0x33)
BODY  = RGBColor(0x3D,0x47,0x57)
GREY  = RGBColor(0x8A,0x90,0x9A)
RED   = RGBColor(0xB3,0x40,0x2C)
GREEN = RGBColor(0x1F,0x5C,0x55)
LINE  = RGBColor(0xD5,0xD1,0xC7)

prs = Presentation()
prs.slide_width = Inches(W); prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]

def sld():
    s = prs.slides.add_slide(blank)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0,0,prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = PAPER; r.line.fill.background(); r.shadow.inherit=False
    return s

def txt(s,x,y,w,h,t,size=14,color=BODY,bold=False,font='Microsoft YaHei',align=PP_ALIGN.LEFT,ls=1.0,sp=None):
    tb = s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf = tb.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    lines = t.split('\n')
    for i,l in enumerate(lines):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = ls
        if sp is not None: p.space_after = Pt(sp)
        r = p.add_run(); r.text = l
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = font
        r.font.color.rgb = color
    return tb

def rule(s,x,y,w,th=0.75,c=LINE):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Pt(th))
    r.fill.solid(); r.fill.fore_color.rgb = c; r.line.fill.background(); r.shadow.inherit=False

# slide 1 font probe
s = sld()
txt(s,0.75,0.6,11.8,0.3,'Noto Serif SC 40 · 交付物越像样，越要知道它会在哪里错',40,INK,False,'Noto Serif SC')
txt(s,0.75,1.5,11.8,0.3,'Noto Serif SC bold 40 · 数字 41,638.9336 万元',40,INK,True,'Noto Serif SC')
txt(s,0.75,2.4,11.8,0.3,'Microsoft YaHei Light 28 · 交付物越像样，越要知道它会在哪里错',28,INK,False,'Microsoft YaHei Light')
txt(s,0.75,3.05,11.8,0.3,'Microsoft YaHei 28 · 交付物越像样，越要知道它会在哪里错',28,INK,False,'Microsoft YaHei')
txt(s,0.75,3.7,11.8,0.3,'Microsoft YaHei bold 28 · 交付物越像样，越要知道在哪里错',28,INK,True,'Microsoft YaHei')
txt(s,0.75,4.35,11.8,0.3,'Consolas 26 · 41,638.9336 / 41,465.1967 / -173.7369 / 55% 81%',26,RED,False,'Consolas')
txt(s,0.75,5.0,11.8,0.3,'Segoe UI 22 · AI Assist 2026 Aa Bb 0123456789',22,BODY,False,'Segoe UI')
txt(s,0.75,5.5,11.8,0.3,'DengXian 24 · 交付物越像样，越要知道它会在哪里错',24,BODY,False,'DengXian')
txt(s,0.75,6.0,11.8,0.3,'SimHei 20 · 交付物越像样，越要知道它会在哪里错',20,BODY,False,'SimHei')
rule(s,0.75,6.5,11.8,0.75)
rule(s,0.75,6.7,11.8,1.5,RED)
txt(s,0.75,6.85,11.8,0.3,'hairline 0.75pt vs 1.5pt red — 排版网格参考线',10,GREY,False,'Microsoft YaHei')

# slide 2: dense paragraph capacity test
s = sld()
txt(s,0.75,0.5,11.8,0.3,'A 正文容量测试 12pt × 11.8in',12,GREY,True)
body = ('这是一个用于验证中文排版行宽与行距的测试段落。工程图纸审查、工程量清单复算、修编报告编写，'
        '这三类任务对准确性的要求远高于对文采的要求；而模型的问题恰恰出现在它最擅长的地方——'
        '它能把格式写得比人还整齐，却也可能在数字上悄悄偏一位。')
txt(s,0.75,0.9,11.8,3.0,body,12,BODY,False,'Microsoft YaHei',PP_ALIGN.LEFT,1.45,6)
txt(s,0.75,4.0,5.6,0.3,'B 半宽 5.6in 12pt',12,GREY,True)
txt(s,0.75,4.4,5.6,3.0,body,12,BODY,False,'Microsoft YaHei',PP_ALIGN.LEFT,1.45,6)
txt(s,7.0,4.0,5.6,0.3,'C 半宽 5.6in 15pt 行距1.6',15,GREY,True)
txt(s,7.0,4.4,5.6,3.0,body,15,BODY,False,'Microsoft YaHei',PP_ALIGN.LEFT,1.6,6)

prs.save(r'_v35\probe.pptx')
print('ok')
