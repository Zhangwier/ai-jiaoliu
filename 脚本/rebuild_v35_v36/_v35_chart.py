# -*- coding: utf-8 -*-
"""矢量图表基元：用 freeform 画真实折线/曲线，避免用截图当图表。"""
import math
from pptx.util import Emu, Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import ds
from ds import SANS, MONO, INK, MUT, RULE, FAINT, BODY


def _ff(s, pts):
    x0, y0 = pts[0]
    ff = s.shapes.build_freeform(Emu(int(I(x0))), Emu(int(I(y0))), scale=1.0)
    ff.add_line_segments([(Emu(int(I(x)), Emu(int(I(y))))) for x, y in pts[1:]], close=False)
    return ff.convert_to_shape()


def poly(s, pts, color=INK, w=1.5, dash=None, close=False):
    if close:
        x0, y0 = pts[0]
        ff = s.shapes.build_freeform(Emu(int(I(x0))), Emu(int(I(y0))), scale=1.0)
        ff.add_line_segments([(Emu(int(I(x))), Emu(int(I(y)))) for x, y in pts[1:]], close=True)
        sh = ff.convert_to_shape()
    else:
        sh = _ff(s, pts)
    sh.fill.background()
    sh.line.color.rgb = color
    sh.line.width = Pt(w)
    if dash:
        sh.line.dash_style = dash
    sh.shadow.inherit = False
    return sh


def fillpoly(s, pts, color, alpha=None):
    x0, y0 = pts[0]
    ff = s.shapes.build_freeform(Emu(int(I(x0))), Emu(int(I(y0))), scale=1.0)
    ff.add_line_segments([(Emu(int(I(x))), Emu(int(I(y)))) for x, y in pts[1:]], close=True)
    sh = ff.convert_to_shape()
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    if alpha is not None:
        from pptx.oxml.ns import qn
        import lxml.etree as etree
        srgb = sh.fill.fore_color._xFill.find(qn('a:srgbClr'))
        a = etree.SubElement(srgb, qn('a:alpha'))
        a.set('val', str(int(alpha * 100000)))
    return sh


def bar_h(s, x, y, w, h, frac, color, bg=ds.RULE, r=0.0):
    if bg is not None:
        ds.box(s, x, y, w, h, fill=bg)
    if frac > 0:
        ds.box(s, x, y, max(0.012, w * frac), h, fill=color)


def dot(s, cx, cy, d, color, line=None):
    return ds.circ(s, cx, cy, d, color, line)


def label(s, x, y, w, t, size=8.5, color=MUT, align=PP_ALIGN.LEFT, mono=False, bold=False):
    ds.text(s, x, y, w, 0.18, t, size, color, MONO if mono else SANS, bold, align, 1.0, check=False)
