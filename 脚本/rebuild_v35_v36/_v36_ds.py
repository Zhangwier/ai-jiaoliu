# -*- coding: utf-8 -*-
"""设计系统 v36 —— 沿用原稿视觉语言（藏青 / 青绿 / 橙 / 米白），重建排版秩序
画布 13.3326 x 7.5 (16:9)"""
import os
from pptx import Presentation
from pptx.util import Inches as I, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

# ---------------- 画布与栅格 ----------------
W, H = 13.33264, 7.5
ML, MR = 0.60, 0.60
CW = W - ML - MR                 # 12.1326
X2 = W - MR                      # 12.7326
MID = W / 2
TOP = 1.95                       # 内容区上缘
BOT = 6.82                       # 内容区下缘
CH = BOT - TOP                   # 4.87

# ---------------- 色板（原稿） ----------------
def C(r, g, b): return RGBColor(r, g, b)

NAVY = C(0x14, 0x21, 0x3D)
INK = C(0x12, 0x20, 0x3A)
DEEP = C(0x1C, 0x2E, 0x4F)
TEAL = C(0x0F, 0x76, 0x6E)
TEALL = C(0x9C, 0xCF, 0xC8)
TEALXL = C(0xE7, 0xF2, 0xF0)
MINT = C(0xE1, 0xF0, 0xEE)
ORANGE = C(0xD9, 0x62, 0x2B)
ORANGEL = C(0xFB, 0xE9, 0xDE)
CREAM = C(0xEC, 0xEA, 0xE4)
PAPER = C(0xF6, 0xF5, 0xF1)
CARD = C(0xFF, 0xFF, 0xFF)
GREY = C(0x5B, 0x65, 0x77)
GREY2 = C(0x6B, 0x73, 0x85)
FAINT = C(0x94, 0xA3, 0xB8)
LINE = C(0xE3, 0xDF, 0xD4)
LINE2 = C(0xCF, 0xC9, 0xBA)
WHITE = C(0xFF, 0xFF, 0xFF)
DARKBG = INK
DARKLINE = C(0x2C, 0x3A, 0x55)
DARKGREY = C(0xA8, 0xB2, 0xC4)

SANS = "微软雅黑"
MONO = "Consolas"
SERIF = "思源宋体"     # 仅用于极少数点睛句
LINES = 1.32           # 微软雅黑行高系数

MEDIA = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "_v34", "media")
WARN = []


def new_deck():
    prs = Presentation()
    prs.slide_width = Emu(12191695)
    prs.slide_height = Emu(6858000)
    return prs


def _nostyle(shp):
    sp = shp._element
    st = sp.find(qn('p:style'))
    if st is not None:
        sp.remove(st)


def bgrect(s, color, x=0, y=0, w=W, h=H):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    _nostyle(r)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background()
    r.shadow.inherit = False
    if r.has_text_frame:
        r.text_frame.text = ""
    return r


def slide(prs, paper=PAPER):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bgr = bgrect(s, paper)
    bgr._element.getparent().remove(bgr._element)
    s.shapes._spTree.insert(2, bgr._element)
    return s


# ---------------- 文字宽度估算 ----------------
def est_w(t, size, bold=False, cs=0.0):
    """返回宽度（pt）"""
    if not t: return 0.0
    w = 0.0
    for ch in str(t):
        o = ord(ch)
        if o < 128:
            f = 0.58
            if ch in " .,:;'|!il": f = 0.30
            elif ch.isupper(): f = 0.72
        else:
            f = 1.0
        w += f * size
    if bold: w *= 1.035
    w += cs * max(len(str(t)) - 1, 0)
    return w


def nlines(t, size, w, bold=False, cs=0.0):
    if not t: return 1
    lim = w * 72
    n = 0
    for para in str(t).split("\n"):
        cw = 0.0; k = 1
        for ch in para:
            o = ord(ch)
            if o < 128:
                f = 0.58
                if ch in " .,:;'|!il": f = 0.30
                elif ch.isupper(): f = 0.72
            else:
                f = 1.0
            cw += f * size + cs
            if cw > lim:
                k += 1; cw = f * size
        n += k
    return n


def text(s, x, y, w, h, t, size=15.5, color=GREY, font=SANS, bold=False,
         align=PP_ALIGN.LEFT, ls=1.35, anchor=MSO_ANCHOR.TOP, check=True,
         italic=False, cs=None):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0;  tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(str(t).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = ls
        r = p.add_run(); r.text = ln
        f = r.font
        f.size = Pt(size); f.bold = bold; f.name = font; f.italic = italic
        f.color.rgb = color
        if cs is not None:
            f._rPr.set('spc', str(int(cs * 100)))
    if check:
        need = nlines(t, size, w, bold, cs or 0) * size * LINES * ls / 72.0
        if need > h * 1.03 + 0.03:
            WARN.append(f"OVF {need:.2f}>{h:.2f} @{x:.2f},{y:.2f} {str(t)[:28]!r}")
    return tb


def rich(s, x, y, w, h, paras, size=15.5, ls=1.35, anchor=MSO_ANCHOR.TOP, wrap=True):
    """paras: [ {'runs':[(txt,{size,color,bold,font,cs})...], 'align':..,'ls':..,'sp':..} ]"""
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0;  tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, pa in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = pa.get("align", PP_ALIGN.LEFT)
        p.line_spacing = pa.get("ls", ls)
        if pa.get("sp") is not None:
            p.space_before = Pt(pa["sp"])
        for txt, st in pa["runs"]:
            r = p.add_run(); r.text = txt
            f = r.font
            f.size = Pt(st.get("size", size))
            f.bold = st.get("bold", False)
            f.name = st.get("font", SANS)
            f.italic = st.get("italic", False)
            f.color.rgb = st.get("color", GREY)
            if st.get("cs") is not None:
                f._rPr.set('spc', str(int(st["cs"] * 100)))
    return tb


# ---------------- 基础图元 ----------------
def hr(s, x, y, w, th=1.0, color=LINE):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), Emu(int(th * 12700)))
    _nostyle(r); r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background(); r.shadow.inherit = False
    return r


def vr(s, x, y, h, th=1.0, color=LINE):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), Emu(int(th * 12700)), I(h))
    _nostyle(r); r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background(); r.shadow.inherit = False
    return r


def box(s, x, y, w, h, fill=None, line=None, lw=1.0, dash=None, radius=None):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    r = s.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    _nostyle(r)
    if radius:
        try: r.adjustments[0] = radius
        except Exception: pass
    if fill is None:
        r.fill.background()
    else:
        r.fill.solid(); r.fill.fore_color.rgb = fill
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = line; r.line.width = Pt(lw)
        if dash: r.line.dash_style = dash
    r.shadow.inherit = False
    if r.has_text_frame:
        r.text_frame.text = ""
    return r


def card(s, x, y, w, h, accent=None, fill=CARD, line=LINE, accent_h=0.055):
    """原稿式卡片：白底细边 + 顶部色条"""
    r = box(s, x, y, w, h, fill=fill, line=line, lw=1.0)
    if accent:
        box(s, x, y, w, accent_h, fill=accent)
    return r


def circ(s, cx, cy, d, fill, line=None, lw=1.0):
    r = s.shapes.add_shape(MSO_SHAPE.OVAL, I(cx - d / 2), I(cy - d / 2), I(d), I(d))
    _nostyle(r)
    if fill is None: r.fill.background()
    else:
        r.fill.solid(); r.fill.fore_color.rgb = fill
    if line is None: r.line.fill.background()
    else:
        r.line.color.rgb = line; r.line.width = Pt(lw)
    r.shadow.inherit = False
    return r


def line(s, x1, y1, x2, y2, color=LINE2, lw=1.2, head=False, dash=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    _nostyle(c)
    c.line.color.rgb = color; c.line.width = Pt(lw)
    if dash: c.line.dash_style = dash
    if head:
        ln = c.line._get_or_add_ln()
        t = etree.SubElement(ln, qn('a:tailEnd'))
        t.set('type', 'triangle'); t.set('w', 'med'); t.set('len', 'med')
    return c


def frame(s, x, y, w, h, color=LINE, lw=1.0, fill=None):
    return box(s, x, y, w, h, fill=fill, line=color, lw=lw)


def pic(s, fn, x, y, w=None, h=None, mode="fit", path=None, border=None):
    from PIL import Image
    p = path or os.path.join(MEDIA, fn)
    iw, ih = Image.open(p).size
    ar = iw / ih
    if w is None: w = h * ar
    if h is None: h = w / ar
    if mode == "fit":
        if w / h > ar:
            w2 = h * ar
            sh = s.shapes.add_picture(p, I(x + (w - w2) / 2), I(y), I(w2), I(h))
        else:
            h2 = w / ar
            sh = s.shapes.add_picture(p, I(x), I(y + (h - h2) / 2), I(w), I(h2))
    else:
        sh = s.shapes.add_picture(p, I(x), I(y), I(w), I(h))
        if w / h > ar:
            f = 1 - (w / (h * ar))
            if f > 0: sh.crop_left = f / 2; sh.crop_right = f / 2
        else:
            f = 1 - (h * ar / w)
            if f > 0: sh.crop_top = f / 2; sh.crop_bottom = f / 2
    if border:
        sh.line.color.rgb = border; sh.line.width = Pt(1.0)
    return sh


# ---------------- 页面构件 ----------------
def header(s, chapter, pno, dark=False, rule=False):
    """顶部：章节标签（左，青绿粗）+ 页码（右，灰）"""
    cc = TEALL if dark else TEAL
    pc = DARKGREY if dark else GREY2
    if chapter:
        text(s, ML, 0.30, 8.6, 0.28, chapter, 12, cc, SANS, True, PP_ALIGN.LEFT, 1.0, cs=1.2, check=False)
    if pno is not None:
        text(s, X2 - 1.0, 0.30, 1.0, 0.28, f"{pno:02d}", 12, pc, SANS, False, PP_ALIGN.RIGHT, 1.0,
             check=False)
    if rule:
        hr(s, ML, 0.70, CW, 1.0, DARKLINE if dark else LINE)


def head(s, chapter, pno, ttl, sub=None, size=28, dark=False, w=None, sub_w=None,
         sub_size=16.5, bar=True):
    """页头：章节标 + 标题（带前置色条）+ 导语；返回导语下沿 y"""
    header(s, chapter, pno, dark)
    ww = w if w else CW
    sz = size
    while sz > 18 and est_w(ttl, sz, True) > ww * 72 * 0.995:
        sz -= 0.5
    if bar:
        box(s, ML, 0.66, 0.50, 0.055, fill=TEALL if dark else TEAL)
    tcol = WHITE if dark else NAVY
    nl = nlines(ttl, sz, ww, True)
    th = nl * sz * LINES * 1.16 / 72.0
    text(s, ML, 0.80, ww, th + 0.06, ttl, sz, tcol, SANS, True, PP_ALIGN.LEFT, 1.16, check=False)
    y = 0.80 + th + 0.14
    if sub:
        sw = sub_w if sub_w else ww
        sn = nlines(sub, sub_size, sw)
        sh = sn * sub_size * LINES * 1.42 / 72.0
        text(s, ML, y, sw, sh + 0.05, sub, sub_size, DARKGREY if dark else GREY, SANS, False,
             PP_ALIGN.LEFT, 1.42, check=False)
        y += sh + 0.05
    return y


def footer(s, src, dark=False, y=7.00):
    hr(s, ML, y, CW, 1.0, DARKLINE if dark else CREAM)
    text(s, ML, y + 0.09, CW - 0.1, 0.24, src, 8.5,
         DARKGREY if dark else FAINT, SANS, False, PP_ALIGN.LEFT, 1.0, check=False)


def chip(s, x, y, t, color=TEAL, fill=MINT, size=11, padx=0.13, h=0.30, w=None, bold=True, cs=0.4):
    ww = w if w else est_w(t, size, bold, cs) / 72.0 + padx * 2
    box(s, x, y, ww, h, fill=fill, radius=0.18)
    text(s, x + padx, y + (h - size * 1.32 / 72) / 2, ww - padx * 1.4, size * 1.4 / 72, t, size, color,
         SANS, bold, PP_ALIGN.LEFT, 1.0, cs=cs, check=False)
    return ww


def stat(s, x, y, big, unit, cap, color=TEAL, big_size=44, w=3.2, cap_size=12,
         cap_color=GREY, cap_ls=1.35, mono=False, bold=True):
    text(s, x, y, w, big_size * 1.24 / 72.0, big, big_size, color, MONO if mono else SANS,
         bold, PP_ALIGN.LEFT, 1.0, check=False)
    if unit:
        bw = est_w(big, big_size, bold) / 72.0
        text(s, x + bw + 0.06, y + big_size * 0.54 / 72.0, 1.8, 0.24, unit, cap_size, cap_color,
             SANS, False, PP_ALIGN.LEFT, 1.0, check=False)
    cy = y + big_size * 1.06 / 72.0
    cn = nlines(cap, cap_size, w)
    text(s, x, cy, w, cn * cap_size * LINES * cap_ls / 72.0 + 0.05, cap, cap_size, cap_color,
         SANS, False, PP_ALIGN.LEFT, cap_ls, check=False)


def bullets(s, x, y, w, items, size=15, gap=None, color=GREY, dot=ORANGE, d=0.075, ls=1.4, pad=0.24):
    yy = y
    for it in items:
        circ(s, x + d / 2, yy + size * 1.30 * 0.5 / 72.0, d, dot)
        nl = nlines(it, size, w - pad)
        hh = nl * size * LINES * ls / 72.0
        text(s, x + pad, yy, w - pad, hh + 0.05, it, size, color, SANS, False, PP_ALIGN.LEFT, ls,
             check=False)
        yy += hh + (gap if gap is not None else 0.16)
    return yy


def takeaway(s, t, y=6.05, color=NAVY, dark=False, w=None, tag="结论", fill=None):
    """底部结论条：色块 + 标签 + 一句判断（原稿有类似横条，这里做成实心块）"""
    ww = w if w else CW
    nl = nlines(t, 16, ww - 1.42, True)
    hh = max(0.62, nl * 16 * LINES * 1.36 / 72.0 + 0.30)
    bg = fill if fill else (DARKLINE if dark else TEALXL)
    box(s, ML, y, ww, hh, fill=bg)
    text(s, ML + 0.26, y + (hh - 15 * 1.1 / 72) / 2 - 0.01, 1.10, 0.26, tag, 12,
         TEALL if dark else TEAL, SANS, True, PP_ALIGN.LEFT, 1.0, cs=1.6, check=False)
    vr(s, ML + 1.30, y + 0.13, hh - 0.26, 1.0, TEALL if dark else C(0xB4, 0xD4, 0xCE))
    text(s, ML + 1.50, y + (hh - nl * 16 * LINES * 1.36 / 72.0) / 2 - 0.01, ww - 1.72, hh - 0.10, t, 16,
         WHITE if dark else NAVY, SANS, True, PP_ALIGN.LEFT, 1.36, check=False)
    return y + hh


def _no_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        ln = tcPr.find(qn(tag))
        if ln is None:
            ln = etree.SubElement(tcPr, qn(tag))
        ln.set('w', '0')


def table(s, x, y, w, rows, colw, size=14, head_size=None, head=True, headfill=TEAL,
          aligns=None, rowlines=True, first_bold=False, colors=None, line_color=LINE,
          valign=MSO_ANCHOR.MIDDLE, pad=0.13, minrow=0.42, head_color=WHITE,
          body_color=C(0x2B, 0x34, 0x45), zebra=None, rowh=None, ls=1.28,
          head_h=0.46, head_ls=1.1):
    nr, nc = len(rows), len(colw)
    tot = sum(colw)
    ws = [w * c / tot for c in colw]
    gt = s.shapes.add_table(nr, nc, I(x), I(y), I(w), I(minrow * nr))
    tbl = gt.table
    tbl._tbl.tblPr.set('firstRow', '1' if head else '0')
    tbl._tbl.tblPr.set('bandRow', '0')
    for i, ww in enumerate(ws):
        tbl.columns[i].width = Emu(int(I(ww)))
    hs = head_size if head_size else size
    for ci in range(nc):
        c = tbl.cell(0, ci)
        _no_border(c)
        c.fill.solid()
        c.fill.fore_color.rgb = headfill if head else CARD
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_left = I(pad); c.margin_right = I(pad * 0.6)
        c.margin_top = 0; c.margin_bottom = 0
        c.text_frame.word_wrap = True
        p = c.text_frame.paragraphs[0]
        p.alignment = (aligns[ci] if aligns else PP_ALIGN.LEFT)
        p.line_spacing = head_ls
        r = p.add_run(); r.text = str(rows[0][ci])
        r.font.size = Pt(hs); r.font.bold = True; r.font.name = SANS
        r.font.color.rgb = head_color if head else GREY
    tbl.rows[0].height = I(head_h if head else minrow)

    for ri in range(1, nr):
        mx = 1
        for ci in range(nc):
            mx = max(mx, nlines(rows[ri][ci], size, ws[ci] - pad * 1.7, first_bold and ci == 0, 0))
        rh = rowh if rowh else max(minrow, mx * size * LINES * ls / 72.0 + 0.20)
        tbl.rows[ri].height = I(rh)
        for ci in range(nc):
            c = tbl.cell(ri, ci)
            _no_border(c)
            c.fill.solid()
            c.fill.fore_color.rgb = (zebra if (zebra and ri % 2 == 0) else CARD)
            c.vertical_anchor = valign
            c.margin_left = I(pad); c.margin_right = I(pad * 0.6)
            c.margin_top = I(0.03); c.margin_bottom = I(0.03)
            c.text_frame.word_wrap = True
            p = c.text_frame.paragraphs[0]
            p.alignment = (aligns[ci] if aligns else PP_ALIGN.LEFT)
            p.line_spacing = ls
            r = p.add_run(); r.text = str(rows[ri][ci])
            r.font.size = Pt(size); r.font.name = SANS
            r.font.bold = bool(first_bold and ci == 0)
            col = body_color
            if colors:
                col = colors(ri, ci) or body_color
            r.font.color.rgb = col
    if rowlines:
        top = y + (head_h if head else minrow)
        yy = top
        for ri in range(1, nr):
            hr(s, x, yy, w, 1.0, line_color)
            yy += tbl.rows[ri].height / 914400.0
    return y + sum(r.height for r in tbl.rows) / 914400.0
