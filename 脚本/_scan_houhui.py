from pptx import Presentation
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
import re
for i, s in enumerate(prs.slides, 1):
    hits=[]
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        if re.search(r"会后", t):
            hits.append((sh.name, t.replace("\n"," | ")[:300]))
    if hits:
        print(f"### slide {i}")
        for n,t in hits: print("   ", n, "=>", t)
