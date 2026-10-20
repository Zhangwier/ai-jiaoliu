# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
log=[]
# A. s5 目录：p0 残留串行 -> 单行；p4 改名
s = prs.slides[4]
for sh in s.shapes:
    if sh.shape_id==27:
        ps = sh.text_frame.paragraphs
        r0 = ps[0].runs[0]; r0.text = "A 常用术语"
        r4 = ps[4].runs[0]; r4.text = "E 分享与工具"
        log.append("s5 list: "+" | ".join(p.text for p in ps))
# B. 全局改名
for i,s in enumerate(prs.slides):
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        if "降智测试与破限" in t or "案例、风险与破限" in t:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    nt = r.text.replace("降智测试与破限","分享与工具").replace("案例、风险与破限","分享与工具")
                    if nt != r.text:
                        log.append(f"s{i+1} id{sh.shape_id}: {r.text[:30]!r} -> {nt[:30]!r}")
                        r.text = nt
prs.save(P)
print("\n".join(log))
