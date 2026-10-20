# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[88]
for sh in s.shapes:
    print(sh.shape_id, sh.name, sh.has_text_frame and repr(sh.text_frame.text))
    if sh.has_text_frame:
        for k,p in enumerate(sh.text_frame.paragraphs):
            print("    p%d" % k, [ (r.text, r.font.size.pt if r.font.size else None) for r in p.runs ])
