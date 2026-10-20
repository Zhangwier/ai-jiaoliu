# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
DIRS = [r"D:\MG\program\ai交流\成品\_work29\logos3", r"D:\MG\program\ai交流\成品\_work29\logos4", r"D:\MG\program\ai交流\成品\_work29\logos2", r"D:\MG\program\ai交流\成品\_work29\logos"]
def pick(n):
    for d in DIRS:
        p=os.path.join(d,n+".png")
        if os.path.exists(p): return p
    return None
prs = Presentation(P)
# s80 左侧 豆包/千问、Hark 两行缺 logo；拆成两行 logo
s = prs.slides[79]
S = {sh.shape_id: sh for sh in s.shapes}
SIZE=0.52
def badge(sl, pth, x, top):
    ov = sl.shapes.add_shape(MSO_SHAPE.OVAL, Cm(x), Cm(top), Cm(SIZE), Cm(SIZE))
    ov.fill.solid(); ov.fill.fore_color.rgb = RGBColor(0xFF,0xFF,0xFF)
    ov.line.color.rgb = RGBColor(0xE2,0xDF,0xD6); ov.line.width = Pt(0.75)
    ov.shadow.inherit=False; ov.name="LogoBadge"
    pad=SIZE*0.17
    sl.shapes.add_picture(pth, Cm(x+pad), Cm(top+pad), Cm(SIZE-2*pad), Cm(SIZE-2*pad)).name="Logo"

# 豆包* / 千问* 行：两个 logo，文字右移
sh = S[32]; h = sh.height
top = (sh.top + (h - Cm(SIZE))/2)/360000
x0 = 1.89
for k in ("doubao","qwen"):
    pth = pick(k)
    if pth: badge(s, pth, x0, top); x0 += SIZE+0.08
sh.left = Cm(x0+0.02); sh.width = Cm(3.68 - (x0+0.02-1.83))
print("s80 豆包/千问 ->", sh.text_frame.text, round(sh.width/360000,2))
prs.save(P)
