# -*- coding: utf-8 -*-
"""1) 把 P25（OpenAI 数学）与 P26（Opus 出视频）从 02 正文移到附录 E，作为 E14/E15；
   2) 新增 E0b「模型架构：稠密、MoE、稀疏、注意力与线性化」，放在 E0 之后。
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

from build_v3 import (put, rect, hline, vline, dot, chrome, section_bar,
                      callout, rowband, L, W, C_TITLE, C_BODY, C_MUTED,
                      C_TEAL, C_ORANGE, C_HEAD, C_LINE, C_WHITE, C_WARM,
                      C_COOL, C_ZEBRA)


def move_slide(prs, from_idx, to_idx):
    lst = prs.slides._sldIdLst
    items = list(lst)
    el = items[from_idx]
    lst.remove(el)
    items2 = list(lst)
    if to_idx >= len(items2):
        lst.append(el)
    else:
        items2[to_idx].addprevious(el)


def page_arch(slide, page):
    chrome(
        slide, page,
        "E0b  模型架构：稠密、MoE、稀疏，差在哪",
        "同样是「大模型」，内部怎么分配算力不一样；这决定了它跑得多贵、长文多稳",
        "来源：Shazeer 等（2017，MoE 原始论文）；DeepSeek-V3 与 Kimi K3 技术报告；本稿整理",
        notes=(
            "讲解与会后阅读\n\n"
            "这一页补一个前面没讲的东西：模型的内部结构。听众常听到 MoE、稀疏、稠密、激活参数这些词，"
            "但很少有人说清它们到底差在哪、对使用者有什么影响。\n\n"
            "一，稠密（Dense）。最直白的结构：每生成一个 token，整个网络的所有参数都要算一遍。"
            "好处是行为可预测、同一输入结果稳定；坏处是参数越多越贵、越慢——"
            "想更聪明就得付出线性增长的成本。多数“小模型”和早期模型都是稠密结构。\n\n"
            "二，MoE（Mixture of Experts，混合专家）。把前馈层拆成很多个“专家”，"
            "每个 token 只被路由到其中少数几个（常见是 2 个），其余不参与计算。"
            "于是出现两个参数口径：总参数（所有专家加起来）和激活参数（这次真正算到的）。"
            "比如 770B 总参数 / 49B 激活，意思是“知识容量按 770B 存，但每次只动 49B”。"
            "这就是“又快又便宜又有大模型知识”的来源。\n\n"
            "三，稀疏（Sparse）。严格说 MoE 就是一种条件计算/稀疏激活；工程上还会用稀疏注意力"
            "（只让每个 token 关注部分位置）等手段降低长文的计算量。所以说“稀疏模型”时，"
            "要问清是“参数稀疏（MoE）”还是“注意力稀疏”。\n\n"
            "四，注意力的两个演进方向：MLA（多头潜在注意力）把 KV 缓存压小，长文更省显存；"
            "线性注意力/混合线性注意力把注意力的计算从随长度平方增长压到接近线性，"
            "代价是可能损失一部分精确回忆能力。Kimi 的 Kimi Delta Attention 属于这一类。\n\n"
            "五，对使用者意味着什么：\n"
            "· 总参数大 ≠ 一定更聪明：很多“大”是存在硬盘里的容量，不是每次都用上；\n"
            "· 激活参数小 = 更快更便宜，但极难任务上可能不如稠密大模型稳；\n"
            "· 稠密模型行为更可预测，适合需要复现的实验；\n"
            "· MoE 的专家路由可能带来“同一问题两次答案差异更大”的现象，做基准测试要注意。\n\n"
            "六，对部署的影响：MoE 的总参数决定显存占用（要装得下所有专家），"
            "激活参数决定算力开销。所以“770B 总参数能在单机上跑”这种说法，要看显存是否够装全部专家。\n\n"
            "一句话收尾：架构决定的是“成本曲线”，不是“聪明程度”；聪明程度还是看评测和你自己的任务。"
        ))

    # 三种结构并排
    cards = [
        ("稠密 Dense", "每个 token 都过全部参数", [
            "全部参数都参与计算",
            "行为稳定、可复现",
            "参数越多越贵越慢",
            "例：多数中小模型",
        ], C_TEAL),
        ("MoE 混合专家", "每个 token 只走少数专家", [
            "总参数 ≫ 激活参数",
            "如 770B 总 / 49B 激活",
            "知识容量大、单次算力小",
            "例：Mistral、Qwen、混元、DeepSeek",
        ], C_ORANGE),
        ("稀疏 Sparse", "只在需要时算需要的那部分", [
            "MoE 属于参数稀疏",
            "还有稀疏注意力（少看一些位置）",
            "问“稀疏”要先问哪一种",
            "长文场景常用",
        ], C_HEAD),
    ]
    x = L
    cw = 3.98
    for title, sub, items, col in cards:
        rect(slide, x, 1.90, cw, 2.16, C_WHITE)
        rect(slide, x, 1.90, cw, 0.06, col)
        put(slide, x + 0.18, 2.02, cw - 0.36, 0.26, title, size=13, bold=True,
            color=C_TITLE)
        put(slide, x + 0.18, 2.30, cw - 0.36, 0.22, sub, size=9, color=C_MUTED)
        yy = 2.56
        for it in items:
            put(slide, x + 0.18, yy, cw - 0.36, 0.24, "· " + it, size=9,
                color=C_BODY, anchor=MSO_ANCHOR.TOP)
            yy += 0.30
        x += cw + 0.19

    # 下半：两个参数口径 + 注意力
    section_bar(slide, L, 4.24, 6.05, "两个参数口径，别看错", h=0.34, size=12.5)
    rows = [
        ("总参数", "所有专家加起来的总量；决定显存要装多大", C_ORANGE),
        ("激活参数", "一次前向真正算到的量；决定跑得多快多贵", C_TEAL),
    ]
    y = 4.68
    for k, v, col in rows:
        rect(slide, L, y, 6.05, 0.44, C_ZEBRA)
        put(slide, 0.76, y, 1.30, 0.44, k, size=10, bold=True, color=col,
            anchor=MSO_ANCHOR.MIDDLE)
        put(slide, 2.10, y, 4.30, 0.44, v, size=9.5, color=C_BODY,
            anchor=MSO_ANCHOR.MIDDLE)
        y += 0.50

    put(slide, L, 5.86, 6.05, 0.46,
        "所以“770B 能在单机跑”要看显存装不装得下全部专家；"
        "“才 49B 激活”说的是算力，不是它只学了 49B 的知识。",
        size=9, color=C_MUTED, spacing=1.15, anchor=MSO_ANCHOR.TOP)

    section_bar(slide, 6.95, 4.24, 5.78, "注意力也在省算力", color=C_ORANGE, h=0.34,
                size=12.5)
    att = [
        ("MLA 多头潜在", "把 KV 缓存压小，长文更省显存"),
        ("线性 / 混合线性", "把平方增长压到接近线性，代价是可能丢细节"),
        ("滑动窗口", "只看附近一段，远处靠层数传导"),
    ]
    y = 4.68
    for k, v in att:
        put(slide, 7.12, y, 1.75, 0.22, k, size=9.5, bold=True, color=C_TEAL,
            anchor=MSO_ANCHOR.TOP)
        put(slide, 8.95, y - 0.01, 3.72, 0.44, v, size=9, color=C_BODY,
            spacing=1.12, anchor=MSO_ANCHOR.TOP)
        y += 0.44

    rect(slide, 6.95, 5.94, 5.78, 0.38, C_COOL)
    put(slide, 7.12, 5.94, 5.46, 0.38,
        "架构决定成本曲线，不决定聪明程度；聪明程度看评测与任务",
        size=9, color=C_BODY, spacing=1.1, anchor=MSO_ANCHOR.MIDDLE)

    callout(slide, "一句话", "总参数看显存，激活参数看钱；两者别混着比", kind="cool")


def main(src, dst):
    prs = Presentation(src)

    # ---- 定位
    def title(i):
        for sh in prs.slides[i].shapes:
            if sh.has_text_frame and sh.name == "TextBox 4":
                return sh.text_frame.text.split("\n")[0]
        return ""

    src_math = next(i for i in range(len(prs.slides)) if title(i).startswith("AI 开始解数学"))
    src_video = next(i for i in range(len(prs.slides)) if title(i).startswith("Opus 5.5 一句话出视频"))
    e0 = next(i for i in range(len(prs.slides)) if title(i).startswith("E0  按任务挑公司"))
    print("math", src_math + 1, "video", src_video + 1, "E0", e0 + 1)

    # ---- 先改标题与标签：这两页现在是 E14/E15
    t_math = [sh for sh in prs.slides[src_math].shapes
              if sh.has_text_frame and sh.name == "TextBox 4"][0]
    t_math.text_frame.paragraphs[0].runs[0].text = "E14  OpenAI 数学：从翻车到可验证的进展"
    s_math = [sh for sh in prs.slides[src_math].shapes
              if sh.has_text_frame and sh.name == "TextBox 1"][0]
    s_math.text_frame.paragraphs[0].runs[0].text = "附录  会后速查"

    t_vid = [sh for sh in prs.slides[src_video].shapes
             if sh.has_text_frame and sh.name == "TextBox 4"][0]
    t_vid.text_frame.paragraphs[0].runs[0].text = "E15  Opus 出视频：一句话，靠写代码画帧"
    s_vid = [sh for sh in prs.slides[src_video].shapes
             if sh.has_text_frame and sh.name == "TextBox 1"][0]
    s_vid.text_frame.paragraphs[0].runs[0].text = "附录  会后速查"

    # ---- 搬页：先把 math 移到末尾，再重新定位 video 并移到末尾
    move_slide(prs, src_math, len(prs.slides))
    src_video = next(i for i in range(len(prs.slides))
                     if title(i).startswith("E15  Opus 出视频"))
    move_slide(prs, src_video, len(prs.slides))

    # 按标题定位 E14 / E15，插到 E13 之后（E14 在前）
    e13 = next(i for i in range(len(prs.slides)) if title(i).startswith("E13"))
    m14 = next(i for i in range(len(prs.slides)) if title(i).startswith("E14"))
    move_slide(prs, m14, e13 + 1)
    m15 = next(i for i in range(len(prs.slides)) if title(i).startswith("E15"))
    move_slide(prs, m15, e13 + 2)

    # ---- 插入 E0b（架构页）：放在 E0 之后
    e0 = next(i for i in range(len(prs.slides)) if title(i).startswith("E0  按任务挑公司"))
    blank = next(l for l in prs.slide_layouts if l.name == "Blank")
    new = prs.slides.add_slide(blank)
    page_arch(new, 0)
    move_slide(prs, len(prs.slides) - 1, e0 + 1)

    # ---- 页码重排
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if not sh.has_text_frame or sh.left is None:
                continue
            if abs(sh.left / 914400 - 11.73) > 0.05 or sh.width / 914400 > 1.3:
                continue
            p = sh.text_frame.paragraphs[0]
            if p.runs and p.runs[0].text.strip().isdigit():
                p.runs[0].text = str(i)

    # ---- E 扉页副标题：E0 / E0b / E1–E15
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            if "E1–E13" in sh.text_frame.text:
                sh.text_frame.paragraphs[0].runs[0].text = (
                    "E0 是按任务挑公司的对照表；E0b 讲模型架构；E1–E15 每家一页，"
                    "最后两页是今年的两个进展。")

    prs.save(dst)
    print("saved", dst, len(prs.slides))
    for i, s in enumerate(prs.slides, 1):
        if 86 <= i <= 112:
            print(f"  {i:3d} {title(i-1)[:52]}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
