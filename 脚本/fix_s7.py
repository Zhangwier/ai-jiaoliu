# -*- coding: utf-8 -*-
"""s7 第五行缩短为单行"""
import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Pt
prs = openprs(); s = prs.slides[6]
row5 = None
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "奖励黑客":
        row5 = sh; break
print("row5", row5.left/360000, row5.top/360000, row5.height/360000)
targets = []
for sh in s.shapes:
    if not sh.has_text_frame: continue
    if abs(sh.top/360000 - row5.top/360000) < 0.05 and sh.left/360000 > row5.left/360000 + 1:
        targets.append(sh)
targets.sort(key=lambda x: x.left)
for t in targets: print("  ", t.left/360000, t.width/360000, t.text_frame.text[:40])
texts = ["钻空子、走捷径、多要权限", "奖励只看结果指标", "写清红线，过程也要查"]
for t, nt in zip(targets, texts):
    settext(t, nt, size=17)
save(prs); print("ok")
