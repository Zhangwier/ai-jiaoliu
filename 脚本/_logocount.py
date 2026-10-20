# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes); continue
        yield sh
rows=[]
for i,s in enumerate(prs.slides,1):
    n=sum(1 for sh in walk(s.shapes) if sh.name in ("Logo","LogoBadge"))
    if n: rows.append((i, n//2 if n%2==0 else n))
print("pages with logos:", len(rows))
print(rows)
