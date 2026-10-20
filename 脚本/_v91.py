# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
for idx in (90,91):
    print("=== slide", idx+1)
    for sh in prs.slides[idx].shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t=sh.text_frame.text.replace("\n"," | ")
            if sh.top/360000 < 5.2: print("  ", sh.shape_id, t[:110])
