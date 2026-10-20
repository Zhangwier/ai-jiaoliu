from pptx import Presentation
from pptx.util import Emu
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
def dump(i):
    s = prs.slides[i]
    print(f"########## SLIDE {i+1}")
    for sh in s.shapes:
        L=lambda v: round(Emu(v).cm,2) if v is not None else None
        t = ""
        if sh.has_text_frame:
            t = sh.text_frame.text.replace("\n"," / ")[:100]
        elif sh.has_table:
            t = "TABLE " + str(len(sh.table.rows)) + "x" + str(len(sh.table.columns))
        print(f"{sh.shape_id:>3} {sh.name:<20} l={L(sh.left)} t={L(sh.top)} w={L(sh.width)} h={L(sh.height)} | {t}")
for i in [4,22,23,50,51,88,89]: dump(i)
print("===== B TITLES")
for i in range(58,75):
    for sh in prs.slides[i].shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith("B"):
            print(f"s{i+1}: {sh.text_frame.text!r}")
