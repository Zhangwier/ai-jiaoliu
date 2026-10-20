from pptx import Presentation
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
s = prs.slides[23]
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() and sh.shape_id in (9,13,17,21):
        print("### id", sh.shape_id, round(sh.left/360000,2), round(sh.top/360000,2), round(sh.width/360000,2), round(sh.height/360000,2))
        tf=sh.text_frame
        print("  anchor", tf.vertical_anchor, "wrap", tf.word_wrap)
        for j,para in enumerate(tf.paragraphs):
            print(f"   p{j} align={para.alignment} lspc={para.line_spacing} space={para.space_before}/{para.space_after}")
            for r in para.runs:
                f=r.font
                print(f"     run {r.text!r} sz={f.size.pt if f.size else None} b={f.bold} color={f.color.rgb if f.color and f.color.type is not None else None} name={f.name}")
    if sh.shape_id in (8,10,12,14,16,18,20,22):
        try:
            print("box", sh.shape_id, sh.name, round(sh.top/360000,2), sh.fill.fore_color.rgb if sh.fill.type==1 else sh.fill.type)
        except Exception as e: print("box", sh.shape_id, e)
