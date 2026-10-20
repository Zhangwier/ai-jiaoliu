# -*- coding: utf-8 -*-
"""公用工具库"""
import copy
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
CM = 360000

NAVY  = RGBColor(0x14,0x21,0x3D)
INK   = RGBColor(0x12,0x20,0x3A)
DEEP  = RGBColor(0x1C,0x2E,0x4F)
TEAL  = RGBColor(0x0F,0x76,0x6E)
ORANGE= RGBColor(0xD9,0x62,0x2B)
GREY  = RGBColor(0x5B,0x65,0x77)
CREAM = RGBColor(0xEC,0xEA,0xE4)
MINT  = RGBColor(0xE1,0xF0,0xEE)
PEACH = RGBColor(0xFB,0xE9,0xDE)
WHITE = RGBColor(0xFF,0xFF,0xFF)

def openprs():
    return Presentation(P)

def save(prs):
    prs.save(P)

def find(slide, name=None, sid=None):
    for sh in slide.shapes:
        if name is not None and sh.name == name: return sh
        if sid  is not None and sh.shape_id == sid: return sh
    return None

def _runs(p):
    return list(p.runs)

def settext(shape, text, size=None, bold=None, color=None):
    """改一段文字，保留原格式；text 可含 \n 表示多段"""
    tf = shape.text_frame
    lines = str(text).split("\n")
    # 保证段落数
    while len(tf.paragraphs) < len(lines):
        p0 = tf.paragraphs[-1]._p
        newp = copy.deepcopy(p0)
        p0.addnext(newp)
    while len(tf.paragraphs) > len(lines):
        p = tf.paragraphs[-1]._p
        p.getparent().remove(p)
    for i, line in enumerate(lines):
        p = tf.paragraphs[i]
        rs = _runs(p)
        if not rs:
            r = p.add_run(); r.text = line; rs=[r]
        else:
            rs[0].text = line
            for extra in rs[1:]:
                extra._r.getparent().remove(extra._r)
        for r in p.runs:
            if size:  r.font.size = Pt(size)
            if bold is not None: r.font.bold = bold
            if color is not None: r.font.color.rgb = color
    return shape

def setpara(shape, texts, size=None, bold=None, color=None):
    return settext(shape, "\n".join(texts), size=size, bold=bold, color=color)

def clone_slide(prs, src):
    """深拷贝整页形状（无图片依赖的页用）"""
    new = prs.slides.add_slide(src.slide_layout)
    spTree = new.shapes._spTree
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    for el in src.shapes._spTree:
        tag = el.tag.split('}')[-1]
        if tag in ("nvGrpSpPr", "grpSpPr"): continue
        spTree.append(copy.deepcopy(el))
    # 重排 shape id，避免冲突
    n = 1
    for el in spTree:
        tag = el.tag.split('}')[-1]
        if tag in ("nvGrpSpPr", "grpSpPr"): continue
        c = el.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}cNvPr')
        if c is not None:
            n += 1
            c.set('id', str(n + 1000))
    return new

def sld_ids(prs):
    return list(prs.slides._sldIdLst)

def move_to(prs, slide, idx):
    lst = prs.slides._sldIdLst
    el = None
    for i, s in enumerate(prs.slides):
        if s.slide_id == slide.slide_id:
            el = list(lst)[i]; break
    lst.remove(el)
    lst.insert(idx, el)

def title_of(slide):
    best=None
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text.strip()
        if not t: continue
        sz=0
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size: sz=max(sz,r.font.size.pt)
        if sz>=22 and 1.5 < sh.top/CM < 4.5:
            if best is None or sz>best[0]: best=(sz,t)
    return best[1] if best else ""

def index_of(prs, key):
    for i,s in enumerate(prs.slides):
        if key in title_of(s): return i
    return -1
