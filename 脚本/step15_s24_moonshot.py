# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[23]
pth = r"D:\MG\program\ai交流\成品\_work29\logos4\moonshot.png"
SIZE = 0.62
x = 5.77; top = 12.16
ov = s.shapes.add_shape(MSO_SHAPE.OVAL, Cm(x), Cm(top), Cm(SIZE), Cm(SIZE))
ov.fill.solid(); ov.fill.fore_color.rgb = RGBColor(0xFF,0xFF,0xFF)
ov.line.color.rgb = RGBColor(0x8F,0xC9,0xC3); ov.line.width = Pt(0.75)
ov.shadow.inherit = False; ov.name="LogoBadge"
pad = SIZE*0.17
s.shapes.add_picture(pth, Cm(x+pad), Cm(top+pad), Cm(SIZE-2*pad), Cm(SIZE-2*pad)).name="Logo"
prs.save(P)
print("moonshot added")
