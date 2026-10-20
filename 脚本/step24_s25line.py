# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[24]
NEW = {
 "最听得懂人话，一句话就能办对；会主动查证，也爱反过来纠正你":
   "最听得懂人话，一句话就办对；会主动查证，也爱纠正你",
 "知识面最广，中文表达最顺；但几乎不反对你，先夸再顺着说":
   "知识面最广，中文最顺；但几乎不反对你，先夸再顺着说",
 "彻底顺着你：你说 A 对它就夸 A 对，你说 A 错它立刻改口":
   "彻底顺着你：你说 A 对，它夸 A 对；你说 A 错，它立刻改口",
}
def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
            continue
        yield sh
n=0
for sh in walk(s.shapes):
    if not sh.has_text_frame: continue
    t = sh.text_frame.text
    if t in NEW:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.text = NEW[t]
                break
            break
        print("updated:", NEW[t])
        n+=1
print("count", n)
prs.save(P)
