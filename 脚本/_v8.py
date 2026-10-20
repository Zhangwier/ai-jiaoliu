from pptx import Presentation
from lxml import etree
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
for sh in prs.slides[57].shapes:
    if sh.shape_id==4:
        print(etree.tostring(sh.text_frame._txBody, pretty_print=True).decode())
