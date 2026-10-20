# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
# s5 p1
for sh in prs.slides[4].shapes:
    if sh.shape_id==27:
        ps=sh.text_frame.paragraphs
        ps[1].runs[0].text = "B 公司档案"
        print("s5:", [p.text for p in ps])
# s58 副标题 p1
for sh in prs.slides[57].shapes:
    if sh.shape_id==4:
        ps=sh.text_frame.paragraphs
        print("s58 before:", [p.text for p in ps])
        ps[1].runs[0].text = "公司档案"
        print("s58 after:", [p.text for p in ps])
prs.save(P)
