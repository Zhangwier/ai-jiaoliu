import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm
EMU=360000
prs = openprs(); s = prs.slides[39]
for sh in s.shapes:
    if sh.name in ("Rectangle 9","Rectangle 20","Rectangle 31"):
        sh.height = Cm(9.98)
    if sh.has_text_frame and abs(sh.top/EMU-14.72)<0.08:
        sh.top = Cm(14.62); sh.height = Cm(0.55)
        for p in sh.text_frame.paragraphs: p.line_spacing = 1.0
save(prs); print("ok")
