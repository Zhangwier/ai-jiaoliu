# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[79]
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print(f"id{sh.shape_id:<3} l={sh.left/360000:6.2f} t={sh.top/360000:6.2f} w={sh.width/360000:6.2f} h={sh.height/360000:5.2f} {sh.text_frame.text[:26]!r}")
