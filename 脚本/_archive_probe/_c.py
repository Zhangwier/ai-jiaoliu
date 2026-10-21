import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.enum.shapes import MSO_SHAPE_TYPE
EMU=360000
prs = openprs(); s = prs.slides[39]
for sh in s.shapes:
    t=sh.top/EMU; h=sh.height/EMU; l=sh.left/EMU; w=sh.width/EMU
    if t>4.0:
        print(f"{sh.shape_type} {sh.name:18} l{l:6.2f} t{t:6.2f} w{w:5.2f} h{h:6.2f} b{t+h:6.2f} | {sh.text_frame.text[:10] if sh.has_text_frame else ''}")
