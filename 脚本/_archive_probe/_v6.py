# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
for i,s in enumerate(prs.slides):
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        for k in ("一页一家公司","公司档案"):
            if k in t:
                print(f"s{i+1} id{sh.shape_id} [{k}] {t[:110]!r}")
