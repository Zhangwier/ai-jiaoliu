from pptx import Presentation
from pptx.util import Emu
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
def info(sn, ids):
    s=prs.slides[sn-1]
    for sh in s.shapes:
        if sh.shape_id in ids and sh.has_text_frame:
            r = sh.text_frame.paragraphs[0].runs
            sz = r[0].font.size.pt if r and r[0].font.size else None
            print(f"s{sn} id{sh.shape_id} l={round(Emu(sh.left).cm,2)} w={round(Emu(sh.width).cm,2)} h={round(Emu(sh.height).cm,2)} sz={sz} {sh.text_frame.text[:24]!r}")
info(78,[14,19,23,27,32,36,41])
info(79,[13,17,22,26,30,35,39,44,48])
info(80,[14,18,22,27,31,36,40,45,53,58,62,67,71,76,80])
