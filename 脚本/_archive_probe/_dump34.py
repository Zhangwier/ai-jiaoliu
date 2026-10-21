import sys, io, json
from pptx import Presentation
from pptx.util import Emu

sys.stdout.reconfigure(encoding='utf-8')
p = r"D:\MG\program\ai交流\成品\AI交流分享-成品34.pptx"
prs = Presentation(p)
print("SLIDE_SIZE", prs.slide_width, prs.slide_height, prs.slide_width/914400, prs.slide_height/914400)
print("COUNT", len(prs.slides))

def walk(shapes, depth=0):
    out = []
    for sh in shapes:
        info = {"name": sh.name, "type": str(sh.shape_type), "id": sh.shape_id}
        try:
            info["box"] = [round(sh.left/914400,2), round(sh.top/914400,2), round(sh.width/914400,2), round(sh.height/914400,2)]
        except Exception:
            info["box"] = None
        if sh.shape_type == 6 or sh.__class__.__name__ == 'GroupShape':
            info["children"] = walk(sh.shapes, depth+1)
        if sh.has_text_frame:
            txt = sh.text_frame.text
            if txt.strip():
                info["text"] = txt
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = []
            for r in sh.table.rows:
                rows.append([c.text for c in r.cells])
            info["table"] = rows
        if getattr(sh, "has_chart", False) and sh.has_chart:
            ch = sh.chart
            info["chart"] = {"type": str(ch.chart_type)}
        if sh.shape_type == 13:
            try:
                info["image"] = sh.image.filename or sh.image.content_type
                info["imgsize"] = sh.image.size
            except Exception:
                pass
        out.append(info)
    return out

for i, s in enumerate(prs.slides, 1):
    print("="*20, "SLIDE", i, "layout:", s.slide_layout.name, "hidden:", s._element.get('show'))
    for sh in walk(s.shapes):
        print(json.dumps(sh, ensure_ascii=False))
