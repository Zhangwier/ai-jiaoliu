from pptx import Presentation
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
for sn, ids in ((78,[27,32,36,41]),(79,[13,26,30,35,39,44]),(80,[14,27,36,45,58,67])):
    s=prs.slides[sn-1]
    print("## s",sn)
    for sh in s.shapes:
        if sh.shape_id in ids:
            print("  ", sh.shape_id, sh.name, repr(sh.text_frame.text[:30]))
