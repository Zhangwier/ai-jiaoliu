# -*- coding: utf-8 -*-
"""紧凑提取：每页按位置输出文字+字号+粗体，忽略装饰形状。"""
import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
p = r"D:\MG\program\ai交流\成品\AI交流分享-成品34.pptx"
prs = Presentation(p)

def sizes(sh):
    out = []
    for para in sh.text_frame.paragraphs:
        for r in para.runs:
            s = r.font.size.pt if r.font.size else 0
            out.append((s, bool(r.font.bold)))
    if not out:
        return 0, False
    return max(x[0] for x in out), any(x[1] for x in out)

def collect(shapes, acc):
    for sh in shapes:
        if sh.__class__.__name__ == 'GroupShape':
            collect(sh.shapes, acc)
            continue
        try:
            top, left = sh.top / 914400, sh.left / 914400
            w, h = sh.width / 914400, sh.height / 914400
        except Exception:
            continue
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                s, b = sizes(sh)
                acc.append((round(top, 2), round(left, 2), w, h, s, b, t.replace("\n", " ⏎ ")))
        elif getattr(sh, "has_table", False) and sh.has_table:
            rows = ["|".join(c.text.replace("\n", " ") for c in r.cells) for r in sh.table.rows]
            acc.append((round(top, 2), round(left, 2), w, h, -1, False, "TBL: " + " ;; ".join(rows)))
        elif sh.shape_type == 13:
            acc.append((round(top, 2), round(left, 2), w, h, -2, False, "[IMG %.1fx%.1f]" % (w, h)))
        elif sh.shape_type == 16:
            acc.append((round(top, 2), round(left, 2), w, h, -3, False, "[VIDEO]"))

start = int(sys.argv[1]) if len(sys.argv) > 1 else 1
end = int(sys.argv[2]) if len(sys.argv) > 2 else len(prs.slides)
for i, s in enumerate(prs.slides, 1):
    if i < start or i > end:
        continue
    acc = []
    collect(s.shapes, acc)
    acc.sort(key=lambda x: (x[0], x[1]))
    print("\n=== P%d (%d shapes)" % (i, len(s.shapes)))
    for top, left, w, h, sz, b, t in acc:
        print("  y%.2f x%.2f | %s%s | %s" % (top, left, ("%.1fpt" % sz) if sz > 0 else "", "B" if b else "", t[:500]))
