from pptx import Presentation
from pptx.util import Emu
from lxml import etree
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
s = prs.slides[24]
for sh in s.shapes:
    if sh.name in ("TextBox 11","TextBox 13","TextBox 38","TextBox 4"):
        print("#####", sh.name)
        print(etree.tostring(sh._element, pretty_print=True).decode()[:2500])
