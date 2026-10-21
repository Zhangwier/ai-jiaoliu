from pptx import Presentation
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
import os
for idx in [58,60,62,63,66,68,69,70,71,72,73,74]:
    s = prs.slides[idx]
    print("slide", idx+1)
    for sh in s.shapes:
        if sh.shape_type == 13:
            img = sh.image
            print("   ", sh.name, sh.width, sh.height, img.filename, img.blob and len(img.blob), img.sha1[:8])
