from pptx import Presentation
p = Presentation("成品/AI交流分享-成品45.pptx")
for i, s in enumerate(p.slides, 1):
    if s._element.get("show") == "0": continue
    best = None
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text.strip()
        if not t: continue
        top = sh.top/360000 if sh.top is not None else 99
        sz = 0
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if r.font.size: sz = max(sz, r.font.size.pt)
        if 1.4 < top < 4.6 and sz >= 20:
            if best is None or sz > best[0]: best = (sz, t.replace("\n"," / "))
    print(f"{i}\t{best[1] if best else ''}")
