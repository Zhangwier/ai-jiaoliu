from pptx import Presentation
from pptx.util import Emu
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
print("slide size cm:", prs.slide_width/360000, prs.slide_height/360000)
s = prs.slides[24]
for sh in s.shapes:
    if sh.shape_id in (11,16,36):
        print(sh.shape_id, sh.name, sh.left, sh.top, sh.width, sh.height)
