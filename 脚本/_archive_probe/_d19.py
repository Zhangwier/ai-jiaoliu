from pptx import Presentation
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
for sn in (78,79,80):
    s=prs.slides[sn-1]
    print("## s",sn)
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            print("  ", sh.shape_id, repr(sh.text_frame.text[:34]))
