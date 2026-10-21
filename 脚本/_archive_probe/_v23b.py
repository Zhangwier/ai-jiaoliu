# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[22]
S = {sh.shape_id: sh for sh in s.shapes}
for sid in (16,20,25,37):
    sh = S[sid]
    p = sh.text_frame.paragraphs[0]
    r = p.runs[0] if p.runs else None
    print(sid, sh.text_frame.text[:22], "algn", p.alignment, "sz", r.font.size.pt if r and r.font.size else None,
          "margins", sh.text_frame.margin_left, sh.text_frame.margin_top)
