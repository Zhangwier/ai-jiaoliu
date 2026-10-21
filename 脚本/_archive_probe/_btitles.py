# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
for i in range(58, 76):
    s = prs.slides[i]
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("B"):
            t = sh.text_frame.text.replace("\n"," | ")
            if len(t) > 8:
                print(i+1, t[:80])
                break
