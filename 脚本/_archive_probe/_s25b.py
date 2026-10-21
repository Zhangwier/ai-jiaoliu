from pptx import Presentation
from pptx.util import Emu
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
s = prs.slides[24]
for sh in s.shapes:
    L = lambda v: round(Emu(v).cm,2) if v is not None else None
    fill = ""
    try:
        if sh.fill.type is not None and sh.fill.type == 1:
            fill = "fill=" + str(sh.fill.fore_color.rgb)
    except Exception as e:
        pass
    print(f"{sh.shape_id:>3} {sh.name:<14} {str(sh.shape_type):<12} l={L(sh.left)} t={L(sh.top)} w={L(sh.width)} h={L(sh.height)} {fill}")
