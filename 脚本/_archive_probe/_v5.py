from pptx import Presentation
from lxml import etree
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[4]
for sh in s.shapes:
    if sh.shape_id==27:
        print(etree.tostring(sh.text_frame._txBody, pretty_print=True).decode()[:3000])
