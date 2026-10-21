import sys, json
from pptx import Presentation
from pptx.util import Emu
p = r"D:\MG\program\ai交流\成品\AI交流分享-成品29 .pptx"
prs = Presentation(p)
print("size", prs.slide_width, prs.slide_height)
print("slides", len(prs.slides))
out=[]
for i, s in enumerate(prs.slides, 1):
    texts=[]
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.replace("\n"," | "))
    t = " || ".join(texts)
    out.append(f"[{i}] {t[:400]}")
print("\n".join(out))
