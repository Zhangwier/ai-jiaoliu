from pptx import Presentation
from pptx.util import Cm
import sys
for f in ["成品/AI交流分享-成品45.pptx"]:
    p = Presentation(f)
    print("===", f, "slides:", len(p.slides))
    hid = 0
    for i, s in enumerate(p.slides, 1):
        try:
            if s._element.get("show") == "0": hid += 1
        except Exception: pass
    print("hidden:", hid)
