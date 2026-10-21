from pptx import Presentation
from pptx.util import Emu
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
def dump(i):
    s = prs.slides[i]
    print(f"########## SLIDE {i+1}")
    for sh in s.shapes:
        L=lambda v: round(Emu(v).cm,2) if v is not None else None
        t = ""
        if sh.has_text_frame: t = sh.text_frame.text.replace("\n"," / ")[:80]
        elif sh.has_table:
            tb=sh.table; t=f"TABLE {len(tb.rows)}x{len(tb.columns)}"
        print(f"{sh.shape_id:>3} {sh.name:<18} l={L(sh.left)} t={L(sh.top)} w={L(sh.width)} h={L(sh.height)} | {t}")
for i in [18,77,78,79]: dump(i)
