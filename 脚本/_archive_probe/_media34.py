import sys, os
from pptx import Presentation
sys.stdout.reconfigure(encoding='utf-8')
p = r"D:\MG\program\ai交流\成品\AI交流分享-成品34.pptx"
out = r"D:\MG\program\ai交流\_v34\media"
os.makedirs(out, exist_ok=True)
prs = Presentation(p)
rows = []
for i, s in enumerate(prs.slides, 1):
    k = 0
    for sh in s.shapes:
        if sh.shape_type == 13:
            k += 1
            try:
                img = sh.image
                blob = img.blob
                ext = img.ext
            except Exception as e:
                print("SKIP", i, k, e)
                continue
            name = "s%03d_%d.%s" % (i, k, ext)
            with open(os.path.join(out, name), "wb") as f:
                f.write(blob)
            rows.append((i, k, name, sh.width/914400, sh.height/914400, len(blob)))
for r in rows:
    print("%d\t%d\t%s\t%.2fx%.2f\t%d" % r)
print("TOTAL", len(rows))
