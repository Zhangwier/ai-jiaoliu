# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[89]
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print(sh.shape_id, repr(sh.text_frame.text[:120]))
