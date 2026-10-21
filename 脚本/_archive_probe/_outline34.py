import sys, json
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
p = r"D:\MG\program\ai交流\成品\AI交流分享-成品34.pptx"
prs = Presentation(p)

def collect(shapes, acc, depth=0):
    for sh in shapes:
        if sh.__class__.__name__ == 'GroupShape':
            collect(sh.shapes, acc, depth+1)
            continue
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                acc.append((round(sh.top/914400,2), round(sh.left/914400,2), t.replace("\n"," / ")))
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = ["|".join(c.text.replace("\n"," ") for c in r.cells) for r in sh.table.rows]
            acc.append((round(sh.top/914400,2), round(sh.left/914400,2), "TABLE: " + " ; ".join(rows)))
        if sh.shape_type == 13:
            acc.append((round(sh.top/914400,2), round(sh.left/914400,2), "[PIC %.1fx%.1f]" % (sh.width/914400, sh.height/914400)))
        if sh.shape_type == 16:
            acc.append((round(sh.top/914400,2), round(sh.left/914400,2), "[MEDIA]"))

for i, s in enumerate(prs.slides, 1):
    acc = []
    collect(s.shapes, acc)
    acc.sort(key=lambda x: (x[0], x[1]))
    print("### %d  (%d shapes) hidden=%s" % (i, len(s.shapes), s._element.get('show')))
    for top, left, t in acc:
        print("   y%.2f x%.2f | %s" % (top, left, t[:400]))
