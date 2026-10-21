# -*- coding: utf-8 -*-
from pptx import Presentation
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
for sn in (83,):
    s = prs.slides[sn-1]
    print("## s", sn)
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            r = sh.text_frame.paragraphs[0].runs
            print(f"  id{sh.shape_id} {sh.name[:16]:<16} l={sh.left/360000:5.2f} t={sh.top/360000:5.2f} w={sh.width/360000:5.2f} h={sh.height/360000:5.2f} sz={r[0].font.size.pt if r and r[0].font.size else None} {sh.text_frame.text[:36]!r}")
