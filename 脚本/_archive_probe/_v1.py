# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
# 1) any 会后 left?
hits=[]
for i,s in enumerate(prs.slides):
    for sh in s.shapes:
        if sh.has_text_frame and "会后" in sh.text_frame.text:
            hits.append((i+1, sh.shape_id, sh.text_frame.text[:50]))
print("会后 hits:", len(hits)); [print("  ",h) for h in hits]
# 2) s25 final
print("\n--- s25")
for sh in prs.slides[24].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print(" ", sh.shape_id, repr(sh.text_frame.text[:60]))
# 3) s89 subtitle + s90 rows
print("\n--- s89")
for sh in prs.slides[88].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print(" ", sh.shape_id, repr(sh.text_frame.text[:90]))
print("\n--- s90 titles")
for sh in prs.slides[89].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() and sh.shape_id in (5,6,17,23,29,35,45):
        print(" ", sh.shape_id, repr(sh.text_frame.text[:70]))
