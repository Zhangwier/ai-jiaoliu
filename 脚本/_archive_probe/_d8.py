from pptx import Presentation
from pptx.util import Emu
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
print("### s92 subtitle/title")
for sh in prs.slides[91].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print("  ", sh.shape_id, repr(sh.text_frame.text[:110]))
print("### fills")
def fill(sh):
    try:
        f=sh.fill
        if f.type is None: return None
        if f.type==1: return str(f.fore_color.rgb)
        return str(f.type)
    except Exception as e: return "ERR"
for i in (19,25):
    for sh in prs.slides[i].shapes:
        if sh.name.startswith("Rectangle") or sh.name.startswith("Rounded"):
            print(i+1, sh.shape_id, sh.name, round(Emu(sh.top).cm,2), fill(sh))
