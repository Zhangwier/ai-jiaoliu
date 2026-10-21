from pptx import Presentation
from pptx.util import Emu
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
s = prs.slides[58]
for sh in s.shapes:
    if sh.name.startswith("Rounded"):
        f = sh.fill
        print(sh.name, sh.shape_id, f.type, f.fore_color.rgb if f.type==1 else None)
        ln = sh.line
        print("  line:", ln.fill.type, getattr(ln.color,'rgb',None) if ln.fill.type==1 else None, ln.width)
