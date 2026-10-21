# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
P  = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
DIRS = [r"D:\MG\program\ai交流\成品\_work29\logos3", r"D:\MG\program\ai交流\成品\_work29\logos2", r"D:\MG\program\ai交流\成品\_work29\logos4"]
def pick(name):
    for d in DIRS:
        p=os.path.join(d,name+".png")
        if os.path.exists(p): return p
prs = Presentation(P)
def badge(slide, path, left, top, size, fill=RGBColor(0xFF,0xFF,0xFF), ring=RGBColor(0xE2,0xDF,0xD6)):
    ov = slide.shapes.add_shape(MSO_SHAPE.OVAL, Cm(left), Cm(top), Cm(size), Cm(size))
    ov.fill.solid(); ov.fill.fore_color.rgb = fill
    ov.line.color.rgb = ring; ov.line.width = Pt(0.75); ov.shadow.inherit=False; ov.name="LogoBadge"
    pad=size*0.17
    pic=slide.shapes.add_picture(path, Cm(left+pad), Cm(top+pad), Cm(size-2*pad), Cm(size-2*pad))
    pic.name="Logo"
CM=360000
s = prs.slides[25-1]; S={sh.shape_id:sh for sh in s.shapes}
for sid,key in [(12,"openai"),(17,"claude"),(22,"gemini"),(27,"deepseek"),(32,"doubao")]:
    print(sid, key, pick(key))
# vertical alignment tuning on slide 25 name cells
for sid in (12,17,22,27,32):
    sh=S[sid]
    print(sid, round(sh.top/CM,2), round(sh.height/CM,2), round(sh.left/CM,2), round(sh.width/CM,2))
