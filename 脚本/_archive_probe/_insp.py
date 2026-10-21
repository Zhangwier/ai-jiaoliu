from pptx import Presentation
from pptx.util import Emu
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
for idx in (58, 19, 20, 23, 88):
    s = prs.slides[idx]
    print("="*20, "slide", idx+1)
    for sh in s.shapes:
        print(f"-- {sh.shape_type} name={sh.name!r} l={round(Emu(sh.left).cm,2) if sh.left is not None else None} t={round(Emu(sh.top).cm,2) if sh.top is not None else None} w={round(Emu(sh.width).cm,2) if sh.width is not None else None} h={round(Emu(sh.height).cm,2) if sh.height is not None else None}")
        if sh.has_text_frame and sh.text_frame.text.strip():
            print("   T:", sh.text_frame.text.replace("\n"," | ")[:260])
