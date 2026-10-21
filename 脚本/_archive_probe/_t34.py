# -*- coding: utf-8 -*-
"""紧凑正文提取：只输出每页文字（标题优先），不含坐标。"""
import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
p = r"D:\MG\program\ai交流\成品\AI交流分享-成品34.pptx"
prs = Presentation(p)
print("SIZE", prs.slide_width / 914400, prs.slide_height / 914400, "COUNT", len(prs.slides))

def maxsize(sh):
    mx = 0
    for para in sh.text_frame.paragraphs:
        for r in para.runs:
            if r.font.size:
                mx = max(mx, r.font.size.pt)
    return mx

def collect(shapes, acc):
    for sh in shapes:
        if sh.__class__.__name__ == 'GroupShape':
            collect(sh.shapes, acc)
            continue
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                acc.append((round(sh.top / 914400, 2), maxsize(sh), t.replace("\n", " ⏎ ")))
        elif getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                acc.append((round(sh.top / 914400, 2), 0, "  TBL| " + " | ".join(c.text.replace("\n", " ") for c in r.cells)))

start = int(sys.argv[1]) if len(sys.argv) > 1 else 1
end = int(sys.argv[2]) if len(sys.argv) > 2 else len(prs.slides)
for i, s in enumerate(prs.slides, 1):
    if i < start or i > end:
        continue
    acc = []
    collect(s.shapes, acc)
    acc.sort(key=lambda x: (-x[1], x[0]))
    print("\n[P%d]" % i)
    for top, sz, t in acc:
        print("  %s" % t[:600])
