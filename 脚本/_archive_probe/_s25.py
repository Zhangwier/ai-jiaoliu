from pptx import Presentation
from pptx.util import Emu
p = r"D:\MG\program\ai交流\成品\AI交流分享-成品29 .pptx"
prs = Presentation(p)
s = prs.slides[24]
print("=== slide 25 ===")
for sh in s.shapes:
    print(f"-- {sh.shape_type} name={sh.name!r} l={Emu(sh.left).cm if sh.left is not None else None} t={Emu(sh.top).cm if sh.top is not None else None} w={Emu(sh.width).cm if sh.width is not None else None} h={Emu(sh.height).cm if sh.height is not None else None}")
    if sh.has_text_frame:
        for pi,para in enumerate(sh.text_frame.paragraphs):
            print(f"   p{pi}: {para.text!r}")
    if sh.has_table:
        tb = sh.table
        for ri,row in enumerate(tb.rows):
            print("   ROW", ri, [c.text for c in row.cells])
