# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
BRANDS = ["ChatGPT","Claude","Gemini","Grok","DeepSeek","千问","Qwen","GLM","Kimi","MiniMax","豆包","文心","混元","Hy","Muse","Hark","Cursor","Cline","Codex","Trae","Qoder","WorkBuddy","Kiro","Copilot","OpenCode","Devin","Jules","Mistral","Meta","OpenAI","Anthropic","Google","xAI","Seed","Step","MiMo"]
for i,s in enumerate(prs.slides):
    npic = sum(1 for x in s.shapes if x.shape_type==13)
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text.strip()
        if not t or len(t) > 24: continue
        hits = [b for b in BRANDS if b in t]
        if len(hits) == 1 and len(t) <= len(hits[0]) + 8:
            print(f"s{i+1:>3} id{sh.shape_id:<3} l={sh.left/360000:5.2f} t={sh.top/360000:5.2f} w={sh.width/360000:5.2f} h={sh.height/360000:4.2f} pics={npic} {t!r}")
