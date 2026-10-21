from pptx import Presentation
from pptx.util import Emu
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
def dump(i):
    s = prs.slides[i]
    print(f"########## SLIDE {i+1}")
    for sh in s.shapes:
        L=lambda v: round(Emu(v).cm,2) if v is not None else None
        t = ""
        if sh.has_text_frame:
            t = sh.text_frame.text.replace("\n"," / ")[:70]
        print(f"{sh.shape_id:>3} {sh.name:<22} l={L(sh.left)} t={L(sh.top)} w={L(sh.width)} h={L(sh.height)} | {t}")
for i in [19,20,25,90]: dump(i)
