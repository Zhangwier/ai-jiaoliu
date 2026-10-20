from pptx import Presentation
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
s = prs.slides[24]
for sh in s.shapes:
    if sh.has_text_frame:
        print(sh.shape_id, sh.name, "|", repr(sh.text_frame.text))
