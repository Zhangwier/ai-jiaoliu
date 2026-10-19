# -*- coding: utf-8 -*-
"""在 E7（DeepSeek）后插 E7b 开源账；在 E8（Qwen）后插 E8b 尺寸阶梯。"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

from build_v3 import (put, rect, hline, vline, dot, chrome, section_bar,
                      callout, rowband, L, W, C_TITLE, C_BODY, C_MUTED,
                      C_TEAL, C_ORANGE, C_HEAD, C_LINE, C_WHITE, C_WARM,
                      C_COOL, C_ZEBRA)


def clear(slide):
    from pptx.oxml.ns import qn
    spTree = slide.shapes._spTree
    for child in list(spTree):
        if child.tag in (qn("p:sp"), qn("p:cxnSp"), qn("p:pic"),
                         qn("p:graphicFrame"), qn("p:grpSp")):
            spTree.remove(child)


# ==================================================== E7b  DeepSeek 开源账
def page_ds_open(slide, page):
    chrome(
        slide, page,
        "E7b  DeepSeek 的开源账：开的到底是什么",
        "不是“把权重放出来”就算开源：代码、权重、训练细节、底层算子，是四种不同的“开”",
        "来源：deepseeklicense.github.io（官方许可 FAQ）；DeepSeek-V3 技术报告 arXiv:2412.19437；"
        "DeepSeek 开源周（2025-02-24 至 02-28）官方仓库",
        notes=(
            "讲解与会后阅读\n\n"
            "这一页回答“DeepSeek 的开源到底彻底在哪”，也是听众最容易被“开源”两个字糊弄的地方。\n\n"
            "先拆成四层，从浅到深：\n"
            "一，权重开放。把模型参数放出来，能下载、能自己跑。多数“开源模型”只做到这一层。\n"
            "二，许可宽松。DeepSeek 的做法是代码用 MIT、模型用自建许可（基于 OpenRAIL 修改）。"
            "官方 FAQ 明确：可商用、免费、无需申请登记；衍生模型不必开源（不是 GPL 那种传染式），"
            "但要保留版权声明、显著声明修改、遵守用途限制。作为对比，Meta 的 Llama 用的是自家社区许可，"
            "限制更多（比如月活到一定规模要额外申请）。\n"
            "三，训练细节公开。V3 技术报告公开发表（671B 总参数 / 37B 激活），"
            "包含架构、训练策略和成本说明。这是“能被复现研究”的关键一层。\n"
            "四，底层算子与工具开源。2025 年 2 月 24 至 28 日的“开源周”，连续放了五个库："
            "FlashMLA（Hopper 上的 MLA 解码内核）、DeepEP（MoE 训练推理的 EP 通信库）、"
            "DeepGEMM（矩阵乘算子库）、EPLB 与 DualPipe（负载均衡与流水并行）、"
            "3FS（Fire-Flyer 文件系统）。这些是“怎么把硬件榨干”的东西，通常厂商不外传。\n\n"
            "为什么第四层最能说明态度：前三层给你“结果”，第四层给你“方法”。"
            "愿意公开方法，等于承认别人可以照着做得更便宜。\n\n"
            "关于“国产追上”的时间点要说准：DeepSeek 第一代是 2023 年的 DeepSeek-LLM 与 DeepSeek-Coder；"
            "真正让国产大模型被全球认真对待的是 2024 年底的 V3 与 2025 年初的 R1。"
            "所以讲“转折点”要落在 V3/R1 这一波，不要说成 V1。\n\n"
            "最后一句要提醒：开源不等于没有成本。权重能下载，算力和运维还得自己扛；"
            "许可宽松也不等于可以随便用，用途限制仍在。"
        ))

    # 四层阶梯
    section_bar(slide, L, 1.88, W, "四种“开”，一层比一层彻底", h=0.36, size=13)
    layers = [
        ("① 权重开放", "参数能下载、能自己跑", "多数“开源模型”只到这一层", C_MUTED),
        ("② 许可宽松", "代码 MIT、模型自建许可；可商用、免登记、衍生可闭源",
         "对比：Llama 社区许可限制更多（规模大要另行申请）", C_TEAL),
        ("③ 细节公开", "V3 技术报告：671B 总 / 37B 激活，架构、策略、成本都写",
         "能被复现研究，而不只是能用", C_TEAL),
        ("④ 底层开源", "开源周五个库：FlashMLA、DeepEP、DeepGEMM、EPLB/DualPipe、3FS",
         "给的是“怎么榨干硬件”的方法，通常不外传", C_ORANGE),
    ]
    y = 2.32
    for i, (k, v, note, col) in enumerate(layers):
        rowband(slide, L, y, W, 0.86, i)
        rect(slide, L, y, 0.06, 0.86, col)
        put(slide, 0.76, y, 1.45, 0.86, k, size=12, bold=True, color=C_TITLE)
        put(slide, 2.30, y + 0.06, 6.25, 0.74, v, size=10, color=C_BODY,
            spacing=1.18, anchor=MSO_ANCHOR.TOP)
        put(slide, 8.70, y + 0.06, 3.55, 0.74, note, size=9, color=C_MUTED,
            spacing=1.18, anchor=MSO_ANCHOR.TOP)
        y += 0.86

    # 底部：时间点提醒
    rect(slide, L, 5.86, 6.05, 0.42, C_WARM)
    put(slide, 0.76, 5.86, 5.75, 0.42,
        "时间点要说准：转折在 V3 / R1（2024 年底—2025 年初），不是 2023 年的第一代",
        size=9, color=C_BODY, spacing=1.1, anchor=MSO_ANCHOR.MIDDLE)

    rect(slide, 6.95, 5.86, 5.78, 0.42, C_ZEBRA)
    put(slide, 7.12, 5.86, 5.46, 0.42,
        "开源 ≠ 零成本：权重能下，算力运维自己扛；许可宽松也有用途限制",
        size=9, color=C_BODY, spacing=1.1, anchor=MSO_ANCHOR.MIDDLE)

    callout(slide, "一句", "前三层给“结果”，第四层给“方法”——愿意给方法，才是真彻底", kind="cool")


# ==================================================== E8b  Qwen 尺寸阶梯
def page_qwen_sizes(slide, page):
    chrome(
        slide, page,
        "E8b  Qwen 的另一个战场：每一档尺寸都要最强",
        "旗舰之外还有一条从 0.6B 到 32B 的完整梯队；单位里真正跑得动的，往往是这一头的",
        "来源：Qwen3 发布的官方说明与媒体报道（2026）；联发科天玑平台 Qwen3 适配技术页；"
        "Hugging Face 开放模型报告；本稿整理",
        notes=(
            "讲解与会后阅读\n\n"
            "前面讲总参数与激活参数、讲私有化部署时，都指向同一个现实：多数单位最后真能用的，"
            "不是旗舰，而是塞得进自己机器的中小模型。这一页就是 Qwen 在这条线上的布局。\n\n"
            "Qwen3 一次放出 8 款：2 款 MoE（30B、235B），6 款稠密（0.6B / 1.7B / 4B / 8B / 14B / 32B）。"
            "官方与报道的说法是“每款都是同尺寸开源模型的 SOTA”——注意这句话的分量："
            "不是“和旗舰比”，是“和自己这个尺寸的对手比”。\n\n"
            "为什么这件事对单位有意义，三个场景：\n"
            "一，端侧。0.6B/1.7B/4B 已经能在手机上跑（联发科天玑平台做过实机验证，性能与功耗平衡）。"
            "现场勘验、外业记录这类没网也要用的场景，靠的就是这一档。\n"
            "二，内网单机。8B 到 32B 是“一张消费级显卡或一台工作站能跑”的区间，"
            "够做分类、抽取、改写这类高频窄任务，成本可控。\n"
            "三，蒸馏与教学。小模型常被用来给大模型做蒸馏目标，或者做本地实验。\n\n"
            "还要说清一个容易混淆的点：同尺寸最强 ≠ 能替代旗舰。"
            "32B 再强，复杂长链推理也难追上 2.4T 的旗舰；选型要看任务难度，不是看榜单名次。\n\n"
            "另外 Qwen 的生态规模也值得一句：累计开源 460+ 模型、衍生模型 30 万+、"
            "近半年下载破 30 亿次，是全球下载量最高的开放模型家族。这意味着“遇到问题更容易搜到答案、"
            "更容易找到别人改好的版本”。\n\n"
            "一句话收尾：旗舰决定它能拿多少分，尺寸梯队决定它能不能真的装进你的机房。"
        ))

    # 尺寸阶梯（横向）
    section_bar(slide, L, 1.88, W, "一条完整的尺寸梯队：0.6B → 32B → 旗舰", h=0.36, size=13)

    steps = [
        ("0.6B", "手机端侧", "够做分类、抽取、短问答"),
        ("1.7B / 4B", "端侧与轻量服务", "天玑平台实机验证；功耗与性能平衡"),
        ("8B / 14B", "单卡工作站", "高频窄任务：改写、结构化、初筛"),
        ("32B", "单机 / 小集群", "接近可用的通用能力，仍装得下"),
        ("30B / 235B MoE", "服务器", "MoE 激活小，单次算力省"),
        ("2.4T 旗舰", "集群", "最难的任务，最贵的价格"),
    ]
    x = L
    for i, (size, scene, what) in enumerate(steps):
        last = (i == len(steps) - 1)
        w = 1.92 if not last else 2.10
        col = C_ORANGE if last else (C_TEAL if i < 4 else C_HEAD)
        box = rect(slide, x, 2.34, w, 1.32, C_WHITE, MSO_SHAPE.ROUNDED_RECTANGLE,
                   radius=0.08)
        box.line.color.rgb = RGBColor.from_string(col)
        box.line.width = Pt(1.0)
        put(slide, x + 0.08, 2.40, w - 0.16, 0.26, size, size=11.5, bold=True,
            color=col, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
        put(slide, x + 0.08, 2.70, w - 0.16, 0.24, scene, size=8.5, color=C_TITLE,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
        put(slide, x + 0.08, 2.96, w - 0.16, 0.64, what, size=8, color=C_MUTED,
            align=PP_ALIGN.CENTER, spacing=1.14, anchor=MSO_ANCHOR.TOP)
        x += w + 0.14

    put(slide, L, 3.76, W, 0.24,
        "读法：越往左越装得下，越往右越聪明；选型看任务难度，不看榜单名次",
        size=9, color=C_MUTED, anchor=MSO_ANCHOR.TOP)

    # 三个场景
    section_bar(slide, L, 4.18, 6.05, "为什么尺寸梯队对单位重要", color=C_ORANGE,
                h=0.34, size=12.5)
    cases = [
        ("端侧", "现场勘验、外业记录：没网也要用"),
        ("内网单机", "8B–32B 一张卡能跑，做高频窄任务"),
        ("蒸馏与实验", "给小模型做教师，或在本地做实验"),
    ]
    y = 4.62
    for k, v in cases:
        rect(slide, L, y, 6.05, 0.44, C_ZEBRA)
        put(slide, 0.76, y, 1.20, 0.44, k, size=9.5, bold=True, color=C_TEAL,
            anchor=MSO_ANCHOR.MIDDLE)
        put(slide, 2.02, y, 4.40, 0.44, v, size=9, color=C_BODY,
            anchor=MSO_ANCHOR.MIDDLE)
        y += 0.50

    section_bar(slide, 6.95, 4.18, 5.78, "别忘了这两句", color=C_HEAD, h=0.34, size=12.5)
    put(slide, 7.12, 4.62, 5.46, 1.28,
        "· 同尺寸最强 ≠ 能替代旗舰：32B 再强，复杂长链也难追 2.4T\n"
        "· 生态规模是隐性优势：开源 460+ 模型、衍生 30 万+、下载破 30 亿次，"
        "遇到问题更容易搜到答案、找到别人改好的版本",
        size=9, color=C_BODY, spacing=1.3, anchor=MSO_ANCHOR.TOP)

    callout(slide, "一句", "旗舰决定它能拿多少分，尺寸梯队决定它能不能装进你的机房", kind="cool")


def main(src, dst):
    prs = Presentation(src)
    blank = next(l for l in prs.slide_layouts if l.name == "Blank")

    def title(i):
        for sh in prs.slides[i].shapes:
            if sh.has_text_frame and sh.name == "TextBox 4":
                return sh.text_frame.text.split("\n")[0]
        return ""

    # 先做 DeepSeek（在 E7 之后插）
    e7 = next(i for i in range(len(prs.slides)) if title(i).startswith("E7  DeepSeek"))
    s1 = prs.slides.add_slide(blank)
    page_ds_open(s1, 0)
    lst = prs.slides._sldIdLst
    items = list(lst)
    el = items[-1]
    lst.remove(el)
    items[e7 + 1].addprevious(el)
    print("inserted E7b after", title(e7)[:20])

    # 再做 Qwen（E8 之后）
    e8 = next(i for i in range(len(prs.slides)) if title(i).startswith("E8  阿里 Qwen"))
    s2 = prs.slides.add_slide(blank)
    page_qwen_sizes(s2, 0)
    items = list(prs.slides._sldIdLst)
    el = items[-1]
    prs.slides._sldIdLst.remove(el)
    items = list(prs.slides._sldIdLst)
    items[e8 + 1].addprevious(el)
    print("inserted E8b after", title(e8)[:20])

    # 页码重排（隐藏页不参与）
    page = 0
    for s in prs.slides:
        if s._element.get("show") != "0":
            page += 1
        for sh in s.shapes:
            if not sh.has_text_frame or sh.left is None:
                continue
            if abs(sh.left / 914400 - 11.73) > 0.05 or sh.width / 914400 > 1.3:
                continue
            p = sh.text_frame.paragraphs[0]
            if p.runs and p.runs[0].text.strip().isdigit():
                p.runs[0].text = str(page)

    # E 扉页文案
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and "E1–E13" in sh.text_frame.text:
                sh.text_frame.paragraphs[0].runs[0].text = (
                    "E1–E13 每家一页：模型线、外壳、强项、边界；"
                    "OpenAI、Anthropic、DeepSeek、Qwen 各多一页值得单独讲的。")

    prs.save(dst)
    print("saved", dst, len(prs.slides))
    for i, t in enumerate([title(i) for i in range(len(prs.slides))], 1):
        if 96 <= i <= 110:
            print(f"  {i:3d} {t[:50]}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
