from pptx import Presentation
from pptx.util import Emu
import os, hashlib
p = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(p)
os.makedirs(r"D:\MG\program\ai交流\成品\_work29\media_out", exist_ok=True)
seen={}
for i, s in enumerate(prs.slides,1):
    for sh in s.shapes:
        if sh.shape_type == 13:
            b = sh.image.blob
            h = hashlib.sha1(b).hexdigest()
            if h in seen: continue
            seen[h]=(i, sh.name, sh.image.ext, len(b), sh.image.size)
            fn = r"D:\MG\program\ai交流\成品\_work29\media_out\%s.%s" % (h[:10], sh.image.ext)
            open(fn,"wb").write(b)
for h,(i,n,e,l,sz) in sorted(seen.items(), key=lambda x:x[1][0]):
    print(i, n, e, l, sz, h[:10])
print("count", len(seen))
