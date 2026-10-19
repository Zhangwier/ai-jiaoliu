# -*- coding: utf-8 -*-
"""把 成品16+附录E 里的附录 E（E 分隔页 + E0–E13）合并进 成品20，去掉 E14（小红书）。
   - 插入位置：D4 之后、Z 分隔页之前
   - 图片用原始 blob 重建（保证在新文件里有效）
   - 全篇页码按物理顺序重排
   - 五个分节扉页与封面目录补上 E 行
"""
import sys
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

GRAY = "94A3B8"


def title_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name == "TextBox 4":
            return sh.text_frame.text.split("\n")[0]
    return ""


def find_slide(prs, pred):
    for i, s in enumerate(prs.slides):
        if pred(title_of(s), s):
            return i
    return None


def build_picture(slide, src_pic):
    """用源图片的 blob 在目标页重建图片，保留位置尺寸与边框。"""
    pic = slide.shapes.add_picture(
        __import__("io").BytesIO(src_pic.image.blob),
        src_pic.left, src_pic.top, src_pic.width, src_pic.height)
    if src_pic.crop_left or src_pic.crop_right or src_pic.crop_top or src_pic.crop_bottom:
        pic.crop_left = src_pic.crop_left
        pic.crop_right = src_pic.crop_right
        pic.crop_top = src_pic.crop_top
        pic.crop_bottom = src_pic.crop_bottom
    try:
        if src_pic.line.color and src_pic.line.color.type is not None:
            pic.line.color.rgb = src_pic.line.color.rgb
            pic.line.width = src_pic.line.width
    except Exception:
        pass
    return pic


def clone_slide(src_prs, dst_prs, idx):
    s = src_prs.slides[idx]
    name = s.slide_layout.name
    layout = next((l for l in dst_prs.slide_layouts if l.name == name),
                  dst_prs.slide_layouts[0])
    new = dst_prs.slides.add_slide(layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)

    for sh in s.shapes:
        if sh.shape_type == 13:            # PICTURE
            build_picture(new, sh)
        elif sh.has_text_frame:
            el = copy.deepcopy(sh._element)
            new.shapes._spTree.append(el)
        elif sh.shape_type is not None and "Connector" in (sh.name or ""):
            new.shapes._spTree.append(copy.deepcopy(sh._element))
        else:
            new.shapes._spTree.append(copy.deepcopy(sh._element))

    # 背景（若有）
    s_csld = s._element.find(qn("p:cSld"))
    bg = s_csld.find(qn("p:bg")) if s_csld is not None else None
    if bg is not None:
        n_csld = new._element.find(qn("p:cSld"))
        old = n_csld.find(qn("p:bg"))
        if old is not None:
            n_csld.remove(old)
        n_csld.insert(0, copy.deepcopy(bg))

    if s.has_notes_slide:
        new.notes_slide.notes_text_frame.text = s.notes_slide.notes_text_frame.text
    return new


def add_rail_e(slide):
    """在分节扉页右侧索引里插一行 E，并把 Z 往下推。"""
    shapes = {}
    for sh in slide.shapes:
        if sh.has_text_frame:
            shapes[sh.name] = sh
    z = shapes.get("TextBox 11")
    if z is None or not z.text_frame.text.strip().startswith("Z"):
        print("  ! rail Z not found")
        return False
    # 新行：复制 Z 的形状
    new_el = copy.deepcopy(z._element)
    z._element.addnext(new_el)
    new_sh = [sh for sh in slide.shapes if sh._element is new_el][0]
    new_sh.name = "TextBox E"
    new_sh.top = z.top
    new_sh.text_frame.paragraphs[0].runs[0].text = "E  一页一家公司（十四页）"
    z.top = z.top + int(0.5 * 914400)
    return True


def main(src20, src16, dst):
    a = Presentation(src20)
    b = Presentation(src16)

    def is_divider(s):
        big = False
        rail = False
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if t == "E" and sh.width is not None and sh.width / 914400 > 3:
                big = True
            if t.startswith("这一部分"):
                rail = True
        return big and rail

    e_div = find_slide(b, lambda t, s: is_divider(s))
    def is_content(s, prefix):
        t = title_of(s)
        hdr = any(sh.has_text_frame and sh.name == "TextBox 1" and
                  "附录" in sh.text_frame.text for sh in s.shapes)
        return t.startswith(prefix) and hdr

    e0 = find_slide(b, lambda t, s: is_content(s, "E0"))
    e13 = find_slide(b, lambda t, s: is_content(s, "E13"))
    print("E divider", e_div + 1, "E0", e0 + 1, "E13", e13 + 1)
    idxs = [e_div] + list(range(e0, e13 + 1))     # 不含 E14

    z_idx = find_slide(a, lambda t, s: any(
        sh.has_text_frame and sh.text_frame.text.strip() == "Z" for sh in s.shapes))
    print("Z divider at", z_idx + 1, "| E pages:", len(idxs))

    made = [clone_slide(b, a, i) for i in idxs]

    lst = a.slides._sldIdLst
    items = list(lst)
    new_ids = items[-len(made):]
    for el in items:
        lst.remove(el)
    keep = [el for el in items if el not in new_ids]
    for el in keep[:z_idx] + new_ids + keep[z_idx:]:
        lst.append(el)

    # 页码：全篇按物理序号重写（有页脚框的才写）
    for i, s in enumerate(a.slides, 1):
        for sh in s.shapes:
            if not sh.has_text_frame or sh.left is None:
                continue
            if abs(sh.left / 914400 - 11.73) > 0.05 or sh.width / 914400 > 1.3:
                continue
            p = sh.text_frame.paragraphs[0]
            if p.runs and p.runs[0].text.strip().isdigit():
                p.runs[0].text = str(i)

    # 分节扉页 / 封面目录补 E 行
    targets = []
    for s in a.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.name == "TextBox 11":
                if sh.text_frame.text.strip().startswith("Z  参考资料"):
                    targets.append(s)
                    break
            if sh.has_text_frame and sh.name == "TextBox 26" and \
                    "B 常用术语" in sh.text_frame.text:
                targets.append(s)
                break
    print("rail targets:", len(targets))
    for s in targets:
        add_rail_e(s)

    # 封面左侧目录里那行文本
    for s in a.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.name == "TextBox 26":
                txt = sh.text_frame.text
                if "D  案例、风险与破限" in txt and "E 一页" not in txt:
                    sh.text_frame.paragraphs[0].runs[0].text = \
                        "A 订阅与省流版\nB 常用术语\nC 资讯渠道\nD  案例、风险与破限\nE 一页一家公司\nZ 参考资料来源"
                    print("  cover index updated")

    a.save(dst)
    print("saved", dst, len(a.slides))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
