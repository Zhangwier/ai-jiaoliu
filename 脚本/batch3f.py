import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm
EMU=360000
prs = openprs(); s = prs.slides[39]
for sh in s.shapes:
    if not sh.has_text_frame and abs(sh.top/EMU-4.50)<0.03 and sh.height/EMU>9.5:
        sh.height = Cm(10.05)
    if sh.has_text_frame and abs(sh.top/EMU-15.02)<0.06:
        sh.top = Cm(14.72)
save(prs); print("ok")
