# -*- coding: utf-8 -*-
"""第 3 页目录：04 那一行的说明补上"我日常在做什么"。"""
from pptx import Presentation

F = r"D:\MG\program\ai交流\AI交流分享-成品26.pptx"
prs = Presentation(F)
s = prs.slides[2]
hit = 0
for sh in s.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text.strip()
    if t == "任务书、上下文、资料库、工作流、落地":
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.text.strip() == t:
                    r.text = "我日常在做什么、任务书、上下文、资料库、工作流、落地"
                    hit += 1
print("patched:", hit)
prs.save(F)
