# -*- coding: utf-8 -*-
"""修正：打码账号标签 A01…A11 被误当作附录编号改写成了 C0x，改回 A0x。"""
import re
from pptx import Presentation

F = r"D:\MG\program\ai交流\AI交流分享-成品26.pptx"
prs = Presentation(F)

PAT = re.compile(r"(?<![A-Za-z0-9])C(0\d)(?![0-9])")
n = 0
for idx in (98, 101, 102):
    s = prs.slides[idx - 1]
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                nt = PAT.sub(lambda m: "A" + m.group(1), r.text)
                if nt != r.text:
                    r.text = nt
                    n += 1
    if s.has_notes_slide:
        tf = s.notes_slide.notes_text_frame
        for p in tf.paragraphs:
            for r in p.runs:
                nt = PAT.sub(lambda m: "A" + m.group(1), r.text)
                if nt != r.text:
                    r.text = nt
                    n += 1

prs.save(F)
print("account labels restored:", n)
