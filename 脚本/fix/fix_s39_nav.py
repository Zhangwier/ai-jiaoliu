# -*- coding: utf-8 -*-
"""第 40 页改为"判据"后，同步第 39 页（04 分节页）的目录项与导语，以及第 3 页总目录描述。"""
from pptx import Presentation

F = r"D:\MG\program\ai交流\AI交流分享-成品26.pptx"
prs = Presentation(F)


def set_single(shape, text):
    p = shape.text_frame.paragraphs[0]
    runs = p.runs
    if not runs:
        return False
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return True


# ---- 第 39 页 ----
s = prs.slides[38]
n = 0
for sh in s.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text.strip()
    if t == "我日常在做什么：三件活":
        set_single(sh, "什么活适合交给 AI")
        n += 1
    elif t.startswith("先说我日常拿它干的三件活"):
        set_single(sh, "先给一套判断标准：什么活适合交出去、什么活自己干；"
                       "再看为什么要用 RAG、怎么整理成工作流、多 Agent 何时值、企业怎么从试点走到交付。")
        n += 1
print("divider40 patched:", n)

# ---- 第 3 页总目录 ----
s3 = prs.slides[2]
m = 0
for sh in s3.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text.strip()
    if t.startswith("我日常在做什么、任务书"):
        set_single(sh, "什么活适合交给 AI、任务书、上下文、资料库、工作流、落地")
        m += 1
print("toc patched:", m)

prs.save(F)
