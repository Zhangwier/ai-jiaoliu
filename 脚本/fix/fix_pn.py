# -*- coding: utf-8 -*-
"""全局重编页码：每个有页码框的页 -> 实际序号"""
import sys, re
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Pt
prs = openprs()
changed=[]
for i, sl in enumerate(prs.slides, start=1):
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        if sh.left/360000 > 29 and sh.top/360000 < 1.6:
            t = sh.text_frame.text.strip()
            if re.fullmatch(r"\d+", t) and int(t) != i:
                changed.append((int(t), i))
                p = sh.text_frame.paragraphs[0]
                r = p.runs[0] if p.runs else p.add_run()
                r.text = str(i)
                for extra in list(p.runs)[1:]: extra._r.getparent().remove(extra._r)
print(len(changed), changed[:12])
save(prs)
