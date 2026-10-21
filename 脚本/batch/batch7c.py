import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm
from pptx.enum.text import MSO_ANCHOR
prs = openprs(); s = prs.slides[39]
fix = {
 "Rectangle 18": 13.10, "Rectangle 29": 13.10, "Rectangle 40": 13.10,
}
for sh in s.shapes:
    if sh.name in fix:
        sh.top = Cm(13.10); sh.height = Cm(1.34)
txts = {
 "TextBox 19": ("怎么做：先问“有哪些路子”，再让它把路子列成一张表", 12),
 "TextBox 30": ("怎么做：一次跑一百个，先抽 3 个验格式和条数", 12),
 "TextBox 41": ("怎么做：拆成小步，每步留一个能查的结果", 12),
}
for sh in s.shapes:
    if sh.name in txts:
        t, sz = txts[sh.name]
        sh.top = Cm(13.12); sh.height = Cm(1.30); sh.width = Cm(9.02 if sh.left/360000 < 11 else 8.94)
        settext(sh, t, size=sz)
        sh.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
save(prs); print("ok")
