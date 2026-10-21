# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
EMU=360000
prs = openprs(); s = prs.slides[108]

def tb(l,t,w,h, lines, size=11.5, bold=False, color=GREY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.22):
    box = s.shapes.add_textbox(Cm(l),Cm(t),Cm(w),Cm(h))
    tf = box.text_frame; tf.word_wrap=True
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    tf.vertical_anchor=anchor
    for i,ln in enumerate(lines):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.line_spacing=ls
        r=p.add_run(); r.text=ln
        r.font.size=Pt(size); r.font.bold=bold; r.font.name="微软雅黑"
        r.font.color.rgb=color
    return box

def rect(l,t,w,h,fill,line=None):
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Cm(l),Cm(t),Cm(w),Cm(h))
    r.fill.solid(); r.fill.fore_color.rgb=fill
    if line: r.line.color.rgb=line; r.line.width=Pt(0.75)
    else: r.line.fill.background()
    r.shadow.inherit=False
    return r

# 清掉正文区
for sh in list(s.shapes):
    if 4.4 < sh.top/EMU < 17.6:
        sh._element.getparent().remove(sh._element)

# 标题与副标
for sh in s.shapes:
    if not sh.has_text_frame: continue
    if sh.top/EMU < 2.2 and sh.left/EMU < 3:
        settext(sh, "E4  开源落地：从权重到能用的四层", size=26)
    elif 3.4 < sh.top/EMU < 4.3:
        settext(sh, "模型权重只是第一层；从权重到“同事能上手”，中间还有三层要拼——每层都有成熟的开源项目，难的不是装，是运维和权限", size=15)

COLS = [
    ("01  跑模型 · 推理引擎", TEAL, [
        ("vLLM", "并发与吞吐最高，服务器上跑开源权重的默认选择"),
        ("SGLang", "结构化输出与多轮对话快，社区更新勤"),
        ("Ollama / llama.cpp", "量化之后，笔记本或单台工作站就能试"),
    ], "这一层决定一张卡能同时服务几个人"),
    ("02  统一接口 · 网关", ORANGE, [
        ("LiteLLM", "几十家模型收成一个 OpenAI 兼容地址，带成本记账"),
        ("One-API", "国内常用，账号池、令牌、额度都在一张表里"),
        ("自建网关", "加脱敏、限流、审计三个中间件就够用"),
    ], "换模型不改代码，全靠这一层"),
    ("03  接资料与界面 · 应用层", DEEP, [
        ("Open WebUI", "像 ChatGPT 的界面，可接本地模型与知识库"),
        ("Dify", "拖拽式工作流，把检索、模型、输出串成一条"),
        ("RAGFlow", "表格与扫描件解析强，检索质量靠它兜底"),
    ], "同事真正上手操作的，是这一层"),
    ("04  看住它 · 可观测与权限", NAVY, [
        ("Langfuse", "每次调用的输入输出、耗时、成本都留痕"),
        ("向量库", "pgvector / Milvus / Weaviate，按数据量选"),
        ("审批与审计", "敏感动作人工确认，日志能按人查"),
    ], "没有这一层，出了事查不出来"),
]
W=7.29; TOP=4.72; BAR=1.02; CARD=8.28
for i,(head,color,items,tail) in enumerate(COLS):
    x = 1.52 + i*(W+0.556)
    bar = rect(x,TOP,W,BAR,color)
    tb(x+0.30, TOP+0.22, W-0.55, 0.62, [head], size=13, bold=True, color=WHITE)
    rect(x, TOP+BAR, W, CARD, WHITE, line=RGBColor(0xE2,0xDF,0xD6))
    y = TOP+BAR+0.18
    for name,desc in items:
        tb(x+0.30, y, W-0.60, 0.55, [name], size=12.5, bold=True, color=color)
        tb(x+0.30, y+0.62, W-0.60, 1.30, [desc], size=11, color=RGBColor(0x2B,0x34,0x45))
        y += 2.18
    rect(x+0.24, TOP+BAR+6.55, W-0.48, 1.02, MINT)
    tb(x+0.42, TOP+BAR+6.72, W-0.80, 0.72, [tail], size=10.5, bold=True, color=TEAL)

# 落地顺序
rect(1.52, 14.22, 30.83, 0.92, MINT)
rect(1.52, 14.22, 0.22, 0.92, TEAL)
tb(1.92, 14.36, 1.60, 0.62, ["顺序"], size=11.5, bold=True, color=TEAL)
tb(3.90, 14.36, 28.20, 0.66,
   ["本机先用 Ollama 跑通一条 → 换 vLLM 上服务器 → 前面加 LiteLLM 统一入口 → 再接 Open WebUI 或 Dify → 最后补日志与审批。别一上来就上集群。"],
   size=11, bold=True, color=NAVY)

# 开源权重映射
tb(1.52, 15.32, 30.83, 0.62,
   ["能直接塞进这条链路的开源权重：DeepSeek V4.1 · Qwen3.8 · GLM-5.3 · Kimi K3 · Hy4 · MiniMax M3（各家的边界见附录 B）"],
   size=11.5, bold=True, color=TEAL)

# 结论条
rect(1.52, 16.13, 30.83, 1.27, PEACH)
rect(1.52, 16.13, 0.23, 1.27, ORANGE)
tb(2.16, 16.13, 2.54, 1.27, ["提醒"], size=14, bold=True, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)
tb(4.52, 16.13, 27.40, 1.27, ["开源能落地，但落地的是项目不是产品：先跑通一条链路，再谈替换"], size=17, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)

# 来源
for sh in s.shapes:
    if sh.has_text_frame and sh.top/EMU > 17.6 and sh.left/EMU < 27:
        settext(sh, "来源：各项目 GitHub 仓库 README（vLLM、SGLang、Ollama、LiteLLM、One-API、Open WebUI、Dify、RAGFlow、Langfuse）；2026-10-20 查询；本稿为选型清单，未做部署实测", size=9.5)

save(prs); print("done")
