# -*- coding: utf-8 -*-
"""修 Muse 标签 + 新增 C4b 省钱页（插在 C4 之后）"""
import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
EMU=360000
prs = openprs()

# --- 修 Muse 标签 ---
s = prs.slides[68]
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "什么时候值得试":
        settext(sh, "何时值得试", size=11.5)

# --- 定位 C4 ---
idx = None
for i, sl in enumerate(prs.slides):
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("C4  订阅与客户端价格速查"):
            idx = i; break
    if idx is not None: break
print("C4 at slide", idx+1)

NEW_AFTER = idx + 2   # 1-based = idx+1，插到它后面
lay = prs.slides[idx].slide_layout
s = prs.slides.add_slide(lay)
for sh in list(s.shapes): sh._element.getparent().remove(sh._element)

def tb(x,y,w,h,text,size,color=NAVY,bold=False,align=None,anchor=None,space=1.28):
    b=s.shapes.add_textbox(Cm(x),Cm(y),Cm(w),Cm(h))
    tf=b.text_frame; tf.word_wrap=True
    if anchor is not None: tf.vertical_anchor=anchor
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    for i,ln in enumerate(text.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if align is not None: p.alignment=align
        p.line_spacing=space
        r=p.add_run(); r.text=ln
        r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name="微软雅黑"
    return b

def rect(x,y,w,h,fill=None,line=None,name=None):
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Cm(x),Cm(y),Cm(w),Cm(h))
    if fill is None: r.fill.background()
    else: r.fill.solid(); r.fill.fore_color.rgb=fill
    if line is None: r.line.fill.background()
    else: r.line.color.rgb=line; r.line.width=Pt(0.75)
    r.shadow.inherit=False
    if name: r.name=name
    return r

tb(1.52,0.81,15.24,0.76,"附录  速查",12,GREY)
tb(29.80,0.81,2.54,0.76,"81",12,RGBColor(0x6B,0x73,0x85),align=PP_ALIGN.RIGHT)
rect(1.52,1.68,1.27,0.15,fill=TEAL)
tb(1.52,1.98,30.81,1.55,"C4b  省钱的路子：三种，各自的代价不一样",26,NAVY,True)
tb(1.52,3.55,30.81,0.95,"便宜不是捡来的：省下的钱，通常换成时间、稳定性或数据风险——先认代价，再决定省不省",16.5,GREY)
conn=s.shapes.add_connector(1,Cm(1.52),Cm(17.78),Cm(32.33),Cm(17.78))
conn.line.color.rgb=RGBColor(0xE2,0xDF,0xD6); conn.line.width=Pt(0.75)
tb(1.52,17.93,26.67,0.76,"来源：各产品定价与活动页；L 站（linux.do）公益与共享服务的社区反馈；闲鱼同类商品观察（2026-10-18）",8.5,GREY)
tb(27.26,17.93,5.08,0.76,"来源索引 ›",8.5,GREY)

cards = [
 ("01  官方活动与首月价", TEAL,
  "代价最低，但窗口短",
  "Cline：ClinePass 首月 $1 的活动（之后 $9.99/月）。\n"
  "Codex：Go 档 $8/月，够轻度用。\n"
  "GLM Coding：Lite ¥118/月，只能在支持的工具里用。\n"
  "各类“新用户首月半价”：记住续费价再决定。",
  "适合：想先试一个月，不打算长期绑定。\n关键是设个日历提醒，别被自动续费扣住。"),
 ("02  国内签到与免费积分", NAVY,
  "代价是时间，不是钱",
  "WorkBuddy：每月 500 积分＋每日签到。\n"
  "Trae：每天 100 次（会员 200）。\n"
  "Qoder：2 周试用＋300 次。\n"
  "Codex 免费账号：可用小模型，额度少。",
  "适合：任务不重、能接受每天手动领一次。\n代价是额度有上限，赶 deadline 时不一定够。"),
 ("03  共享、代充与公益站", ORANGE,
  "代价最高：数据、稳定、合规",
  "开源中转与公益站（L 站社区里常见的那几个，如 anyrouter 一类）：免费或极低价，靠人维护。\n"
  "闲鱼上的共享号、代充、拼车：几块钱一个月。\n"
  "共同点：你的提示词和资料要经过别人的服务器。",
  "适合：只跑公开材料、不涉及单位数据。\n不建议：合同、图纸、报告、客户信息。"),
]
x=1.52; W=9.86; GAP=0.62
for i,(tag,color,sub,body,note) in enumerate(cards):
    xx=x+i*(W+GAP)
    rect(xx,4.55,W,10.30,fill=WHITE)
    rect(xx,4.55,W,1.02,fill=color)
    b=s.shapes.add_textbox(Cm(xx+0.30),Cm(4.55),Cm(W-0.60),Cm(1.02))
    tf=b.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    p=tf.paragraphs[0]; r=p.add_run(); r.text=tag
    r.font.size=Pt(15); r.font.bold=True; r.font.color.rgb=WHITE; r.font.name="微软雅黑"
    tb(xx+0.30,5.72,W-0.60,0.60,sub,14,color if color!=WHITE else TEAL,True)
    tb(xx+0.30,6.40,W-0.60,3.55,body,12,RGBColor(0x2B,0x34,0x45),space=1.34)
    rect(xx+0.30,10.15,W-0.60,0.04,fill=RGBColor(0xE2,0xDF,0xD6))
    tb(xx+0.30,10.35,W-0.60,2.60,note,12,NAVY,space=1.34)

# 中间一条判断提示
rect(1.52,15.05,30.81,0.82,fill=CREAM)
tb(1.85,15.05,30.20,0.82,"判断标准只有一条：这段材料能不能给陌生人看。能，三种随便选；不能，只走前两种。",13.5,NAVY,True,anchor=MSO_ANCHOR.MIDDLE)

rect(1.52,16.10,30.81,1.22,fill=MINT)
rect(1.52,16.10,0.23,1.22,fill=TEAL)
tb(2.16,16.10,2.54,1.22,"提醒",14,NAVY,anchor=MSO_ANCHOR.MIDDLE)
tb(4.45,16.10,27.26,1.22,"共享账号随时会被封、会换模型、会跑路；别把交付压在它上面",17,NAVY,anchor=MSO_ANCHOR.MIDDLE)

move_to(prs, s, NEW_AFTER)
save(prs); print("ok -> slide", NEW_AFTER)
