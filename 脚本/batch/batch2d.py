# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm, Pt
EMU=360000
prs = openprs(); s = prs.slides[19]
new = [
"GPT-6.1 Sol 与 GPT-6 Luna 都是 GPT，AA 智能指数差 20 分以上；Opus 5.5 与 Haiku 5.5 同理。\n表现：能做的事一样，做对的概率不一样。\n所以：别用最便宜的那一档去判断“AI 行不行”。",
"Claude 强长任务与代码；Gemini 强长文与中文；DeepSeek 便宜快；豆包顺着你。\n表现：同一个任务，有人 4 分钟做完，有人 34 分钟还套错。\n所以：别拿 A 的短板去否定 B 的长处。",
"能不能读图、改文件、跑脚本、连续干多久——这些不写在模型里，写在外壳里。\n表现：只靠 OCR 的配置零命中；能把 PDF 渲染成图的全部命中。\n所以：模型一样时，先换壳。",
]
bodies = [sh for sh in s.shapes if sh.has_text_frame and abs(sh.top/EMU-7.05) < 0.06]
bodies.sort(key=lambda x: x.left)
for b, txt in zip(bodies, new):
    settext(b, txt, size=13.5)
    for p in b.text_frame.paragraphs:
        p.line_spacing = 1.26
save(prs); print("ok", len(bodies))
