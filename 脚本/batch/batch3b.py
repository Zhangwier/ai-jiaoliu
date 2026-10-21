# -*- coding: utf-8 -*-
"""重排 s40：三段式卡（标题/正文/注意），字号调好"""
import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
EMU=360000
prs = openprs(); s = prs.slides[39]
for sh in list(s.shapes):
    if sh.has_text_frame:
        if sh.top/EMU > 4.3: sh._element.getparent().remove(sh._element)
    else:
        if sh.top/EMU > 4.3: sh._element.getparent().remove(sh._element)

def tb(x,y,w,h,text,size,color=NAVY,bold=False,anchor=None,space=1.25):
    b=s.shapes.add_textbox(Cm(x),Cm(y),Cm(w),Cm(h))
    tf=b.text_frame; tf.word_wrap=True
    if anchor is not None: tf.vertical_anchor=anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    for i,ln in enumerate(text.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.line_spacing=space
        r=p.add_run(); r.text=ln
        r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name="微软雅黑"
    return b

def rect(x,y,w,h,fill=None,name=None):
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Cm(x),Cm(y),Cm(w),Cm(h))
    if fill is None: r.fill.background()
    else: r.fill.solid(); r.fill.fore_color.rgb=fill
    r.line.fill.background(); r.shadow.inherit=False
    return r

cards=[
 ("01  打破信息茧房", TEAL,
  "你不知道有这个东西\n所以不知道可以要",
  "找工具、找产品、找公式、找规范、找先例。你知道要什么，但不知道有这些路子。",
  "一百多个文件里定位一份图纸；某个问题该查哪本规范；Excel 里这个效果用什么函数。",
  "它的价值是省掉搜索和试错，不是替你判断。",
  "先问“有哪些路子”，再让它把路子比成一张表"),
 ("02  量大门槛低、目标明确", NAVY,
  "一堆低级活\n人做会烦、会漏",
  "批量图片转表格、批量改名归档、批量抽字段、按固定模板灌数。",
  "千张巡检照片要转成 Excel；三百个合同要抽甲方和金额；一千个文件名要按规则改。",
  "这类不涉及复核：转完就是结果，不需要你再判断对错。",
  "一次跑一百个，先抽 3 个验格式和条数"),
 ("03  综合任务", ORANGE,
  "要判断、要取舍\n要跨材料对照",
  "写修编报告、审图纸、合同条文改写、多步工作流、要调用工具的长任务。",
  "四份修编报告要互相对得上；一张图纸要看图、看表、看说明再下结论。",
  "这类最容易出彩，也最容易出事：幻觉、迎合、上下文腐烂全在这里。",
  "拆成小步，每步留一个能查的结果，人签字"),
]
x=1.52; W=9.86; GAP=0.62
for i,(tag,color,head,body,case,note,howto) in enumerate(cards):
    xx=x+i*(W+GAP)
    rect(xx,4.50,W,11.05,fill=WHITE)
    rect(xx,4.50,W,1.02,fill=color)
    b=s.shapes.add_textbox(Cm(xx+0.30),Cm(4.50),Cm(W-0.60),Cm(1.02))
    tf=b.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    p=tf.paragraphs[0]; r=p.add_run(); r.text=tag
    r.font.size=Pt(15); r.font.bold=True; r.font.color.rgb=WHITE; r.font.name="微软雅黑"
    tb(xx+0.30,5.70,W-0.60,1.10,head,15.5,NAVY,True,space=1.25)
    tb(xx+0.30,6.90,W-0.60,1.60,body,12.5,NAVY,space=1.28)
    tb(xx+0.30,8.65,W-0.60,0.55,"典型场景",10.5,TEAL,True)
    tb(xx+0.30,9.15,W-0.60,1.70,case,12,RGBColor(0x2B,0x34,0x45),space=1.28)
    tb(xx+0.30,10.95,W-0.60,0.55,"要注意",10.5,ORANGE,True)
    tb(xx+0.30,11.45,W-0.60,1.60,note,12,RGBColor(0x2B,0x34,0x45),space=1.28)
    rect(xx+0.30,13.15,W-0.60,1.30,fill=CREAM)
    tb(xx+0.46,13.15,W-0.92,1.30,"怎么做："+howto,11.5,GREY,anchor=MSO_ANCHOR.MIDDLE,space=1.2)

tb(1.52,15.60,30.81,0.62,"顺序很重要：先用第一类建立信任，再把第二类交给它跑量，最后才把第三类整段交出去。",14,TEAL,True)
rect(1.52,16.20,30.81,1.15,fill=PEACH)
rect(1.52,16.20,0.23,1.15,fill=ORANGE)
tb(2.16,16.20,2.54,1.15,"结论",14,NAVY,anchor=MSO_ANCHOR.MIDDLE)
tb(4.45,16.20,27.26,1.15,"先按类分，再按类给不同的检查强度：第三类永远要人签字",18,NAVY,anchor=MSO_ANCHOR.MIDDLE)
save(prs); print("ok")
