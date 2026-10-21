# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm, Pt
EMU=360000
prs = openprs()

UPD = {
 54: ("A1–A4", "模型本身 · 怎么交代任务 · Agent 与工具 · 常见误解", 15),
 59: ("B1–B13", "模型线 · 可用入口 · 强项 · 边界；OpenAI、Anthropic、Meta（Muse）、DeepSeek、Qwen 另各有专题页；降智与封号紧跟在对应公司之后", 14),
 80: ("C1–C5", "四件事自查 · 照着做七步 · 地区与支付 · 注册 · 价格 · 省钱的路子 · 省流版；价格与地区政策以官方页面为准", 14),
 88: ("D1–D4", "社区看风向 · GitHub 看版本 · B 站看演示 · X 看一手", 15),
 93: ("E1–E4", "降智测试（附两个自查工具）· 中转站与代理池 · 护栏与破限 · 开源落地", 14),
 110: ("F1–F9", "发展史 · 神经网络 · token 与向量 · Transformer 与注意力 · 架构 · 规模定律 · 对齐与推理 · 多模态", 14),
 121: ("Z1–Z3", "Z1 按编号归档；Z2、Z3 可以直接点开的原链接；正文页脚都指到这里", 14),
}
for n,(code,body,sz) in UPD.items():
    s=prs.slides[n-1]
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        if sh.left/EMU<22 and 12.0 < sh.top/EMU < 14.6 and sh.height/EMU>1.0:
            settext(sh, f"{code}  {body}", size=sz)
            print(n, "|", f"{code}  {body}"[:50])
save(prs)
