import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm
EMU=360000
prs = openprs(); s = prs.slides[80]
# 重新排：所有 t>4.2 且 t<16.0 的形状按顺序压缩
shapes = [sh for sh in s.shapes if 4.2 < sh.top/EMU < 16.0]
shapes.sort(key=lambda x: (round(x.top/EMU,1), x.left))
# 分组按行起始 top
rows = {}
for sh in shapes:
    rows.setdefault(round(sh.top/EMU,2), []).append(sh)
print(sorted(rows))
newT = [4.45 + 1.52*i for i in range(7)]
# 行底矩形 top 是 y，圆点/文字 top 是 y+0.55 / y+0.20
keys = sorted(rows)
base = newT
for i,k in enumerate(keys):
    delta = base[i] - k
    for sh in rows[k]:
        sh.top = Cm(sh.top/EMU + delta)
# 高度
for sh in s.shapes:
    t = sh.top/EMU
    if 4.3 < t < 15.2:
        h = sh.height/EMU
        if h > 1.6: sh.height = Cm(1.52)
        elif abs(h-0.66) < 0.05: sh.height = Cm(0.62)
        else: sh.height = Cm(1.45)
save(prs); print("ok")
