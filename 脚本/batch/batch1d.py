# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
EMU=360000
prs = openprs(); s = prs.slides[6]

# 1) 删掉第5行区域所有形状（t 13.7~14.05）
kill = []
for sh in list(s.shapes):
    try: t = sh.top/EMU
    except: continue
    if 13.70 < t < 14.05:
        kill.append(sh)
print("killing", [(k.name, round(k.top/EMU,2)) for k in kill])
for k in kill: k._element.getparent().remove(k._element)

# 2) 重建
T5, H5 = 13.82, 1.85
z = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.52), Cm(T5), Cm(30.82), Cm(H5))
z.fill.solid(); z.fill.fore_color.rgb = CREAM; z.line.fill.background(); z.shadow.inherit=False
z.name = "Row5Zebra"

def tb(x, w, txt, size, bold=False):
    b = s.shapes.add_textbox(Cm(x), Cm(T5), Cm(w), Cm(H5))
    tf = b.text_frame; tf.word_wrap = True
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    p = tf.paragraphs[0]; r = p.add_run(); r.text = txt
    r.font.size = Pt(size); r.font.bold = bold; r.font.name="微软雅黑"
    r.font.color.rgb = ORANGE if bold else NAVY
    b.name = "Row5_" + str(x)
    return b

tb(1.83, 4.59, "奖励黑客", 16.5, True)
tb(6.92, 9.43, "绕规则、改测试、走捷径；护住自己、多要权限", 14.5)
tb(16.86, 7.13, "按结果给奖励，只对指标负责", 14.5)
tb(24.50, 7.64, "写清红线，过程也要查", 14.5)

save(prs); print("ok")
