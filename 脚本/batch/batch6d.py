# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
EMU=360000
prs = openprs(); s = prs.slides[85]

# 找三栏 x
xs = sorted({round(sh.left/EMU,2) for sh in s.shapes if abs(sh.top/EMU-4.55)<0.03 and sh.height/EMU<1.2})
print("tags", xs)
cardW = 9.86
# 1) body 区缩短
for sh in s.shapes:
    if sh.has_text_frame and abs(sh.top/EMU-6.40)<0.05:
        sh.top = Cm(6.15); sh.height = Cm(2.80)
        for p in sh.text_frame.paragraphs: p.line_spacing = 1.30
# 2) 分隔细线 -> 9.02
for sh in s.shapes:
    if not sh.has_text_frame and abs(sh.top/EMU-10.15)<0.06 and sh.height/EMU<0.1:
        sh.top = Cm(9.02)
# 3) note 文字 -> 移到灰底块内
for sh in s.shapes:
    if sh.has_text_frame and abs(sh.top/EMU-10.35)<0.06:
        sh.top = Cm(9.30); sh.height = Cm(2.60)
        for p in sh.text_frame.paragraphs: p.line_spacing = 1.34
# 4) 卡片高 -> 14.90 底
for sh in s.shapes:
    if not sh.has_text_frame and abs(sh.top/EMU-4.55)<0.03 and sh.height/EMU>10:
        sh.height = Cm(10.35)
save(prs); print("ok")
