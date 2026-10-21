# -*- coding: utf-8 -*-
"""修正 C1 四卡：按行分组定位"""
import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm
from pptx.enum.text import MSO_ANCHOR
prs = openprs(); s = prs.slides[79]
def g(n): return find(s, name=n)

ROW1 = {"big":["TextBox 9","TextBox 13"], "head":["TextBox 10","TextBox 14"], "body":["TextBox 11","TextBox 15"]}
ROW2 = {"big":["TextBox 17","TextBox 21"], "head":["TextBox 18","TextBox 22"], "body":["TextBox 19","TextBox 23"]}

def place(row, card_top):
    for nm in row["big"]:
        g(nm).top = Cm(card_top+0.55); g(nm).height = Cm(6.60)
        g(nm).text_frame.vertical_anchor = MSO_ANCHOR.TOP
    for nm in row["head"]:
        g(nm).top = Cm(card_top+0.55); g(nm).height = Cm(1.60)
    for nm in row["body"]:
        g(nm).top = Cm(card_top+2.20); g(nm).height = Cm(2.55)

place(ROW1, 4.95)
place(ROW2, 10.41)

# 大字母对齐：居中偏上
from pptx.enum.text import PP_ALIGN
for nm in ("TextBox 9","TextBox 13","TextBox 17","TextBox 21"):
    sh = g(nm)
    for p in sh.text_frame.paragraphs: p.alignment = PP_ALIGN.CENTER
save(prs); print("ok")
