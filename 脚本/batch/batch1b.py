# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

prs = openprs()
def S(n): return prs.slides[n-1]
def g(slide, name): return find(slide, name=name)

def newtb(slide, x, y, w, h, text, size, color=NAVY, bold=False, align=None, anchor=None, name=None):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame; tf.word_wrap = True
    if anchor is not None: tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align is not None: p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        r.font.name = "微软雅黑"
    if name: tb.name = name
    return tb

s = S(10)
newtb(s, 1.52, 3.81, 30.82, 1.02, "三步里只有第三步在教它“怎么说话才讨人喜欢”；这是迎合的动机来源，不是 bug", 17, GREY)
s = S(11)
settext(g(s,"TextBox 341"), "它只算“下一个字”，没有“回头查证”这一步；流畅不等于正确", size=18)
save(prs)
print("ok")
