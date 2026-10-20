# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[23]
# 删掉残留图
for sh in list(s.shapes):
    if sh.name=="Logo" and abs(sh.top/360000-12.27)<0.1:
        sh._element.getparent().remove(sh._element); print("removed stray pic")
# 补第3行第4个：moonshot，位置接在 5.77 / top 11.58
pth = r"D:\MG\program\ai交流\成品\_work29\logos4\moonshot.png"
SIZE=0.62; x=5.77; top=11.58
ov = s.shapes.add_shape(MSO_SHAPE.OVAL, Cm(x), Cm(top), Cm(SIZE), Cm(SIZE))
ov.fill.solid(); ov.fill.fore_color.rgb = RGBColor(0xFF,0xFF,0xFF)
ov.line.color.rgb = RGBColor(0x8F,0xC9,0xC3); ov.line.width = Pt(0.75)
ov.shadow.inherit = False; ov.name="LogoBadge"
pad=SIZE*0.17
s.shapes.add_picture(pth, Cm(x+pad), Cm(top+pad), Cm(SIZE-2*pad), Cm(SIZE-2*pad)).name="Logo"
prs.save(P); print("ok")
