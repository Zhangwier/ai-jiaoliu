import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm
prs = openprs(); s = prs.slides[40]
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("另有一张"):
        settext(sh, "另有一张「册汇总」，先看它决定翻哪一册", size=11)
    if sh.has_text_frame and "第二册 路线 · 26 个文件" in sh.text_frame.text:
        settext(sh, "第二册 路线 · 26 个文件 · 9 个数量表\n第三册一分册 路基 · 18 个 · 11 个数量表", size=12)
save(prs); print("ok")
