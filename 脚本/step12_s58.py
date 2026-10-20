# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
for sh in prs.slides[57].shapes:
    if sh.shape_id==4:
        ps = sh.text_frame.paragraphs
        print("before:", [p.text for p in ps])
        while len(ps) > 1:
            ps[-1]._p.getparent().remove(ps[-1]._p)
            ps = sh.text_frame.paragraphs
        print("after:", [p.text for p in sh.text_frame.paragraphs])
prs.save(P)
