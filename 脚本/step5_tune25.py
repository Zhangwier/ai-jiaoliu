# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Pt
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[24]; S={sh.shape_id:sh for sh in s.shapes}
def set_body(sh, txt, size=None):
    tf=sh.text_frame; p0=tf.paragraphs[0]; runs=list(p0.runs)
    if runs:
        runs[0].text=txt
        for r in runs[1:]: r._r.getparent().remove(r._r)
    else: tf.text=txt
    if size:
        for r in tf.paragraphs[0].runs: r.font.size=Pt(size)
set_body(S[34], "未进本次盲测，也无同口径分数；仅日常使用观感", size=14.5)
set_body(S[29], "出字最快、最便宜；但先来一段“雷霆大思考”，自问自答完才给结果", size=14.5)
set_body(S[24], "知识面最广，中文表达最顺；但几乎不反对你，先夸再顺着说", size=14.5)
set_body(S[19], "最听得懂人话，一句话就能办对；会主动查证，也爱反过来纠正你", size=14.5)
set_body(S[14], "写代码偏“防御式”，兜底写得很全；推理这一项最硬", size=14.5)
prs.save(P)
print("ok")
