# -*- coding: utf-8 -*-
"""企业落地重做为一条路线：选任务 → 补能力 → 分人机 → 定投入 → 看结果。
全程用修编报告案例串起来；开源平台页移到附录 D，改成对照表。"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from build_v3 import (put, rect, hline, vline, dot, section_bar, rowband,
                      L, W, C_TITLE, C_BODY, C_MUTED, C_TEAL, C_ORANGE,
                      C_HEAD, C_LINE, C_WHITE, C_WARM, C_COOL, C_ZEBRA)

ROUTE = ["① 选任务", "② 补能力", "③ 分人机", "④ 定投入", "⑤ 看结果"]


def clear(slide):
    spTree = slide.shapes._spTree
    for child in list(spTree):
        if child.tag in (qn("p:sp"), qn("p:cxnSp"), qn("p:pic"),
                         qn("p:graphicFrame"), qn("p:grpSp")):
            spTree.remove(child)


def set_bg(slide):
    from build_v3 import set_bg as sb
    sb(slide)


def shell(slide, page, step, title, source, notes, crumb=None):
    """页眉 + 标题 + 路线指示条 + 页脚。"""
    set_bg(slide)
    put(slide, L, 0.32, 6.0, 0.30, "04  怎么用得稳", size=12, bold=True,
        color=C_TEAL, anchor=MSO_ANCHOR.TOP)
    put(slide, 11.73, 0.32, 1.00, 0.30, str(page), size=12, color=C_MUTED,
        align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.TOP)
    rect(slide, L, 0.66, 0.50, 0.06, C_TEAL)
    put(slide, L, 0.76, W, 0.62, title, size=25, bold=True, color=C_TITLE,
        spacing=1.0)

    # 路线条
    x = L
    cw = 2.37
    for i, label in enumerate(ROUTE):
        cur = (i + 1 == step)
        box = rect(slide, x, 1.42, cw, 0.34,
                   C_TEAL if cur else C_ZEBRA,
                   MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
        box.line.fill.background()
        put(slide, x, 1.42, cw, 0.34, label, size=10.5 if cur else 10,
            bold=cur, color=C_WHITE if cur else C_MUTED,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += cw + 0.045
    if crumb:
        put(slide, L, 1.82, W, 0.22, crumb, size=10, color=C_MUTED,
            anchor=MSO_ANCHOR.TOP)

    hline(slide, L, 7.00, W)
    put(slide, L, 7.06, 10.50, 0.30, source, size=10.5, color=C_MUTED,
        anchor=MSO_ANCHOR.TOP)
    put(slide, 10.73, 7.06, 2.00, 0.30, "来源索引 ›", size=10.5, bold=True,
        color=C_TEAL, anchor=MSO_ANCHOR.TOP)
    slide.notes_slide.notes_text_frame.text = notes


def band(slide, y, label, body, kind="warn", h=0.50):
    bar = C_WARM if kind == "warn" else C_COOL
    edge = C_ORANGE if kind == "warn" else C_TEAL
    rect(slide, L, y, W, h, bar)
    rect(slide, L, y, 0.09, h, edge)
    put(slide, 0.85, y, 1.05, h, label, size=13.5, bold=True, color=edge)
    put(slide, 1.80, y, 10.70, h, body, size=15.5, bold=True, color=C_TITLE,
        spacing=1.0)


def table(slide, x, y, w, cols, rows, widths, rh=0.66, hdr=0.36,
          fs=10.5, hfs=11, zebra=True):
    rect(slide, x, y, w, hdr, C_HEAD)
    cx = x
    for label, frac in zip(cols, widths):
        cw = w * frac
        put(slide, cx + 0.14, y, cw - 0.20, hdr, label, size=hfs, bold=True,
            color=C_WHITE)
        cx += cw
    yy = y + hdr
    for i, row in enumerate(rows):
        if zebra:
            rowband(slide, x, yy, w, rh, i)
        else:
            rect(slide, x, yy, w, rh, C_WHITE)
        cx = x
        for j, (cell, frac) in enumerate(zip(row, widths)):
            cw = w * frac
            put(slide, cx + 0.14, yy + 0.04, cw - 0.22, rh - 0.08, cell,
                size=fs, color=C_BODY if j else C_TITLE,
                bold=(j == 0), spacing=1.16, anchor=MSO_ANCHOR.MIDDLE)
            cx += cw
        yy += rh
    return yy


# ============================================================ ① 选任务
def p1(slide, page):
    shell(slide, page, 1,
          "企业 AI 的起点：一个能验收的业务任务",
          "来源：本稿修编报告案例的执行与复核记录；两案例返工统计；本稿整理",
          "第 1 步 · 选任务\n\n"
          "这一页解决“第一步做什么”。不要从平台、模型、架构开始，从一个能验收的任务开始。\n\n"
          "为什么选修编报告：它有明确的输入（清单、函件、勘误表）、明确的产出（一份报告）、"
          "明确的验收人（编制人、复核人）。这三点齐了，才谈得上“做对没做对”。\n\n"
          "五问逐个过：\n"
          "一，痛点。清单核对、原因查找、模板修改反复耗时；一份报告两三天，返工多半花在找依据上。"
          "这是可观察的，不是猜的。\n"
          "二，哪部分交给 AI。提取变化项、定位函件、起草说明——这三件事的共同点是“有材料可查、"
          "产出是文字”，正好是模型擅长的。\n"
          "三，怎么判断做对了。金额能复算、原因有出处、模板无旧信息残留。这三条都能当场验，"
          "不靠感觉。\n"
          "四，谁接收。明确编制人与复核人。这一条最容易被跳过，但它是“能不能进流程”的关键。\n"
          "五，值不值得继续。比较总时间与总成本，而且必须把复核和返工算进去——只算生成那几分钟，"
          "一定得出“很划算”的错误结论。\n\n"
          "页底那句是本页落点：试点范围越具体，越容易判断投入有没有价值。"
          "反过来说，范围越大越像口号，越没法验收。",
          crumb="拿一份真实的修编报告走一遍：五问都有答案，才值得立项")

    table(slide, L, 2.12, W,
          ["判断问题", "修编报告场景怎么回答"],
          [("业务痛点是什么？", "清单核对、原因查找、模板修改反复耗时；一份报告两三天，返工多半花在找依据上"),
           ("哪部分适合交给 AI？", "提取变化项、定位函件、起草说明——有材料可查、产出是文字"),
           ("怎样判断做对了？", "金额能复算、原因有出处、模板无旧信息残留"),
           ("谁接收并负责？", "明确编制人与复核人；签字的人不变，AI 只是把草稿做快"),
           ("值不值得继续投入？", "比较包含复核和返工在内的总时间与成本，而不是只看生成那几分钟")],
          [0.28, 0.72], rh=0.70)

    band(slide, 6.30, "结论", "试点范围越具体，越容易判断投入有没有价值", kind="cool")


# ============================================================ ② 补能力
def p2(slide, page):
    shell(slide, page, 2,
          "同一份修编报告，进入单位流程还需要什么",
          "来源：本稿两案例的复核记录；Scalekit 多租户 Agent 访问控制（身份分层）；本稿整理",
          "第 2 步 · 补能力\n\n"
          "上一页确认了任务值得做；这一页说清楚：同样一份报告，为什么个人用得好好的，"
          "进了单位就不够。\n\n"
          "四个问题都是业务语言，不是技术名词：\n"
          "一，身份与权限——谁可以读这个项目？AI 用谁的身份去访问资料？个人用的时候这两个问题不存在，"
          "因为只有你一个人。\n"
          "二，资料与版本——用的是哪版清单？函件齐不齐？旧版作废没有？这一条我特意提到最前面，"
          "因为在工程业务里，用错版本比有没有漂亮的管理后台严重得多：后台难看只是难用，"
          "版本错了是结论错。\n"
          "三，审批与追溯——谁改了结论？谁复核的？能不能找到当时的依据？\n"
          "四，成本与运行——每份花多少？失败谁处理？能不能恢复？\n\n"
          "下半页那条链是本页的骨架：资料入库 → AI 编制草稿 → 程序复算 → 人工复核 → 审批归档。"
          "每一环只写两件事：谁负责、留下什么产物。不写技术组件，因为听众要记的是责任，不是架构。\n\n"
          "特别注意“程序复算”这一环是独立的，不并进 AI——数字要能独立复算，这是后面验收的基础。",
          crumb="个人用法缺的往往不是模型能力，是这四件事；其中“资料与版本”最要紧")

    table(slide, L, 2.12, W,
          ["要补的能力", "听众容易理解的具体问题"],
          [("身份与权限", "谁可以读取这个项目？AI 用谁的身份访问资料？"),
           ("资料与版本", "用的是哪版清单？函件是否齐全？旧版是否作废？"),
           ("审批与追溯", "谁修改了结论？谁复核？能否找到当时的依据？"),
           ("成本与运行", "每份报告花多少？失败后谁处理？能否恢复？")],
          [0.24, 0.76], rh=0.46)

    put(slide, L, 4.42, 6.20, 0.24, "一条业务链：每个环节只留“谁负责 + 什么产物”",
        size=10.5, bold=True, color=C_TEAL, anchor=MSO_ANCHOR.TOP)

    chain = [("资料入库", "业务人员", "文件清单、版本号"),
             ("AI 编制草稿", "编制人", "草稿、变更表、待确认项"),
             ("程序复算", "程序", "复算底稿、差额"),
             ("人工复核", "复核人", "复核记录、修改痕迹"),
             ("审批归档", "编制人 / 负责人", "签发件、留档材料")]
    x = L
    bw = 2.25
    for i, (t, who, out) in enumerate(chain):
        fill = C_COOL if i == 2 else C_WHITE
        box = rect(slide, x, 4.70, bw, 1.30, fill,
                   MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        box.line.color.rgb = __import__("pptx").dml.color.RGBColor.from_string(C_LINE)
        box.line.width = Pt(0.75)
        put(slide, x + 0.10, 4.74, bw - 0.20, 0.24, t, size=10.5, bold=True,
            color=C_TITLE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
        put(slide, x + 0.10, 5.02, bw - 0.20, 0.22, who, size=8.5, color=C_TEAL,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
        put(slide, x + 0.10, 5.26, bw - 0.20, 0.66, out, size=8, color=C_MUTED,
            align=PP_ALIGN.CENTER, spacing=1.12, anchor=MSO_ANCHOR.TOP)
        if i < 4:
            put(slide, x + bw, 5.16, 0.27, 0.24, "→", size=12, color=C_MUTED,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
        x += bw + 0.27

    band(slide, 6.30, "边界", "签字的人不变：AI 只把草稿做快，判断与签发仍在人手上", kind="warn")


# ============================================================ ③ 分人机
def p3(slide, page):
    shell(slide, page, 3,
          "修编报告试点：哪些步骤自动化，哪些由人负责",
          "来源：本稿修编报告案例的实际执行与复核记录；表格结构说明",
          "第 3 步 · 分人机\n\n"
          "这一页回答“谁干什么”。五个环节，三种承担者：业务人员、程序与 AI、编制人与复核人。\n\n"
          "为什么“确认材料与口径”必须是人：口径这件事没有唯一答案，它取决于这次结算按哪版规则、"
          "哪些函件算数。让 AI 定口径，等于把责任推给工具。\n\n"
          "为什么“提取清单变化”可以给程序与 AI：这一步是“把 A 和 B 对齐”，"
          "有明确的对照关系，产出是一张能逐行核对的变化表。\n\n"
          "为什么“匹配变更原因”要检索加 AI，而且必须交出“待确认项”：原因往往散在函件、"
          "勘误表、会议纪要里，检索负责找，AI 负责组织语言。找不到的，就老实列出来。\n\n"
          "为什么“复算与格式检查”给程序：数字用脚本重算，不靠模型心算。这一环是后面验收的地基。\n\n"
          "为什么“判断与签发”回到人：报告是要签字的，签字的人必须看过。\n\n"
          "页底那句边界是我最想强调的：找不到原因的子目进入待确认清单，不用通顺的套话填满。"
          "这是这类项目最容易出的问题——模型会写得很像样，但那是编的。",
          crumb="三种承担者：业务人员定口径，程序与 AI 做提取与起草，人做判断与签发")

    table(slide, L, 2.12, W,
          ["环节", "建议承担者", "必须交出的结果"],
          [("确认材料与口径", "业务人员", "文件清单、版本、金额口径"),
           ("提取清单变化", "程序与 AI", "可追溯的变化表（逐行能对回原文）"),
           ("匹配变更原因", "检索与 AI", "原因草稿、函件位置、待确认项"),
           ("复算与格式检查", "程序", "差额、异常与检查记录"),
           ("判断与签发", "编制人、复核人", "经确认的报告及留档材料")],
          [0.26, 0.20, 0.54], rh=0.68)

    rect(slide, L, 6.02, W, 0.24, C_ZEBRA)
    put(slide, 0.78, 6.02, W - 0.32, 0.24, [
        ("一条边界：", 9, True, C_ORANGE),
        ("找不到原因的子目进入待确认清单，不用通顺的套话填满——套话比空白更难发现", 9,
         False, C_BODY)], spacing=1.0)

    band(slide, 6.30, "分工", "把“谁签字”定下来，再谈自动化到哪一步", kind="cool")


# ============================================================ ④ 定投入
def p4(slide, page):
    shell(slide, page, 4,
          "部署、检索与微调，分别解决什么问题",
          "来源：本稿两案例实测；各厂部署与微调文档；公开的企业落地经验；本稿整理",
          "第 4 步 · 定投入\n\n"
          "这一页拆开三件常被混在一起的事：部署方式、检索、微调。它们解决的是完全不同的问题，"
          "钱也花在不同地方。\n\n"
          "资料不能离开指定环境——这是部署与访问控制的问题，不是训练的问题。"
          "很多单位一听到“数据不能出内网”，第一反应是自己训一个，其实先要解决的是“模型跑在哪、"
          "谁能访问”。\n\n"
          "模型不知道本单位资料——这是检索的问题：建一个有版本、有权限的资料库。"
          "微调解决不了这个：微调改的是说话方式和任务习惯，不是往里灌知识。\n\n"
          "数字和固定流程不稳定——这是工程问题：加程序计算、加校验、把步骤写成工作流。"
          "换成更强的模型往往也治不好，因为问题在流程不在智力。\n\n"
          "特定任务仍不达标——到这一步才谈微调，而且前提是有训练样本和评测依据，"
          "否则训完说不清好没好。\n\n"
          "想从零训练通用模型——这是另一件事，要单独论证数据、算力、团队和长期维护，"
          "不适合和上面的试点混在一起谈。\n\n"
          "页底那句是方法：先定位缺的是资料、流程还是任务能力，再决定投入在哪里。"
          "定位错了，钱一定花在错的地方。",
          crumb="四类问题对应四种投入；先定位，再花钱")

    table(slide, L, 2.12, W,
          ["遇到的问题", "优先考虑的方向"],
          [("资料不能离开指定环境", "选择满足要求的部署方式与访问控制（私有部署、专有实例、网关管控）"),
           ("模型不知道本单位资料", "建立可检索、有版本和权限的资料库；不靠微调灌知识"),
           ("数字和固定流程不稳定", "增加程序计算、校验与工作流；把关键量交给脚本复算"),
           ("特定任务表现仍不达标", "有训练样本和评测依据后，评估微调"),
           ("想从零训练通用模型", "单独论证数据、算力、团队与长期维护投入")],
          [0.30, 0.70], rh=0.72)

    band(slide, 6.30, "顺序", "先定位缺的是资料、流程还是任务能力，再决定投入在哪里", kind="cool")


# ============================================================ ⑤ 看结果
def p5(slide, page):
    shell(slide, page, 5,
          "试点是否有效，要看复核后的交付结果",
          "来源：本稿修编报告案例的返工与复核记录；本稿整理；对照表待试点填写",
          "第 5 步 · 看结果\n\n"
          "这一页是收口：怎么判断这个试点到底成不成。\n\n"
          "五项指标，都能记，不靠感觉：\n"
          "一，每份总工时——包括准备、生成、复核、返工。只记生成时间一定高估收益。\n"
          "二，首次验收通过情况——按固定验收清单记，通过就是通过，不通过写明缺什么。\n"
          "三，关键错误——金额、单位、版本、原因引用，分别记。混在一起记，就不知道该改哪。\n"
          "四，原因可追溯情况——能不能定位到有效的原始依据。这一条决定报告敢不敢签发。\n"
          "五，每份总成本——模型费用加人工与运行维护成本。\n\n"
          "右边那张对照表我特意留空：不填虚构的提升比例。试点没跑够样本之前，"
          "任何百分比都是编的。等跑完一批，再填。\n\n"
          "下半页是责任分工，三句话：业务负责人决定答案和验收标准；技术负责人处理系统、权限和运行；"
          "复核人检查交付物并记录错误。三条缺一条，试点就会变成“没人管”。\n\n"
          "页底那句是全节的落点：只有把人工复核和返工算进去，才能知道 AI 有没有真正节省工作。",
          crumb="五项指标都能记录；对照表留空，等试点数据填")

    table(slide, L, 2.12, 7.32,
          ["指标", "记录方式"],
          [("每份总工时", "准备、生成、复核、返工，分段计时后相加"),
           ("首次验收通过", "按固定验收清单逐项打勾，未通过写明缺项"),
           ("关键错误", "金额 / 单位 / 版本 / 原因引用，四类分开记"),
           ("原因可追溯", "能否定位到有效的原始依据（函件号、页码）"),
           ("每份总成本", "模型费用 + 人工工时折算 + 运行维护")],
          [0.26, 0.74], rh=0.58, fs=9.5)

    # 右侧空白对照表
    tx = 8.10
    tw = 4.73
    rect(slide, tx, 2.12, tw, 0.36, C_HEAD)
    put(slide, tx + 0.12, 2.12, 1.95, 0.36, "对照项", size=10, bold=True, color=C_WHITE)
    put(slide, tx + 2.02, 2.12, 1.28, 0.36, "人工流程", size=10, bold=True, color=C_WHITE)
    put(slide, tx + 3.30, 2.12, 1.30, 0.36, "AI 辅助", size=10, bold=True, color=C_WHITE)
    rows = ["每份总工时", "首次验收通过", "关键错误数", "原因可追溯", "每份总成本"]
    yy = 2.48
    for i, r in enumerate(rows):
        rowband(slide, tx, yy, tw, 0.58, i)
        put(slide, tx + 0.12, yy, 1.90, 0.58, r, size=9.5, bold=True, color=C_TITLE,
            anchor=MSO_ANCHOR.MIDDLE)
        put(slide, tx + 2.02, yy, 1.24, 0.58, "待测", size=9, color=C_MUTED,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        put(slide, tx + 3.30, yy, 1.26, 0.58, "待测", size=9, color=C_MUTED,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        yy += 0.58
    put(slide, tx, yy + 0.06, tw, 0.44,
        "试点未跑够样本，不填提升比例；跑完一批再填，并注明样本量",
        size=8.5, color=C_MUTED, spacing=1.12, anchor=MSO_ANCHOR.TOP)

    # 责任分工
    put(slide, L, 5.44, 5.0, 0.24, "维护责任：三条各归各的", size=10.5, bold=True,
        color=C_TEAL, anchor=MSO_ANCHOR.TOP)
    roles = [("业务负责人", "决定答案和验收标准"),
             ("技术负责人", "系统、权限与运行问题"),
             ("复核人", "检查交付物并记录错误")]
    x = L
    for t, v in roles:
        rect(slide, x, 5.70, 3.98, 0.46, C_ZEBRA)
        put(slide, x + 0.14, 5.70, 1.30, 0.46, t, size=9.5, bold=True, color=C_TITLE,
            anchor=MSO_ANCHOR.MIDDLE)
        put(slide, x + 1.50, 5.70, 2.34, 0.46, v, size=9, color=C_BODY,
            anchor=MSO_ANCHOR.MIDDLE)
        x += 3.98 + 0.19

    band(slide, 6.30, "落点", "把人工复核和返工算进去，才知道 AI 有没有真正省下工作")


# ============================================================ 附录 D4
def p_appendix(slide, page):
    from build_v3 import chrome
    chrome(
        slide, page,
        "D4  开源平台对照：先看它声称支持什么",
        "三档标注：文档声称 / 实测通过 / 还需开发；本稿只到第一档",
        "来源：github.com/himovo/movo（README、架构与能力表）；github.com/Stonewuu/agenteam（README）；"
        "本稿未做部署实测",
          "D4 开源平台对照\n\n"
          "这一页放在附录，因为它属于“选型参考”，不该出现在正文的业务主线里。\n\n"
          "上一版把这两个项目写成“能直接部署”，容易让人以为装上就能用于生产。"
          "这次改成分三档标注：文档声称支持、我们实际验证通过、还需要开发。"
          "后两档目前都是空的——我们只读过文档，没有部署实测，所以都标“未实测”。\n\n"
          "比较维度换成业务能判断的五项：权限、资料接入、审批、日志、部署维护。"
          "技术栈、容器数量这些放备注，正文不占篇幅。\n\n"
          "用法：把这张表当“提问清单”。任何平台说自己能做企业版，就让它逐行回答这五项，"
          "并说明是文档写的还是已经跑通的。",
    )

    table(slide, L, 1.90, W,
          ["维度", "MOVO（基于 DSH）", "AgenTeam（群策）", "我们的验证"],
          [("权限", "组织 / 用户 / 角色；模型、知识、能力分别授权", "角色与资源授权；按资源上架", "文档声称"),
           ("资料接入", "PDF/DOCX/XLSX/PPTX/CSV/MD 解析，含图表；检索服务", "知识库 + 数据源接入", "文档声称"),
           ("审批", "敏感动作审批、执行追踪", "关键步骤人工确认", "文档声称"),
           ("日志", "执行追踪、产物留存、审计记录", "执行记录、通知回执", "文档声称"),
           ("部署维护", "Docker Compose 多服务；Mongo / Redis / Weaviate", "Docker 部署；MySQL / Redis", "文档声称")],
          [0.14, 0.31, 0.28, 0.13], rh=0.68, fs=9.5)

    rect(slide, L, 5.92, W, 0.26, C_ZEBRA)
    put(slide, 0.78, 5.92, W - 0.32, 0.26, [
        ("怎么用这张表：", 9.5, True, C_TEAL),
        ("让平台逐行回答这五项，并说明“文档写的”还是“已跑通的”；本稿只读过文档，未做部署实测", 9.5,
         False, C_BODY)], spacing=1.0)

    band(slide, 6.30, "提醒", "“开源可部署”不等于“装上就能用于生产”，中间差的是这五项的实测", kind="cool")


# ============================================================ main
def main(src, dst):
    prs = Presentation(src)
    jobs = [(48, p1, 50), (49, p2, 51), (50, p3, 52), (51, p4, 53), (52, p5, 54)]
    for idx, fn, page in jobs:
        clear(prs.slides[idx])
        fn(prs.slides[idx], page)
        print("rebuilt", idx + 1, "->", page)

    # 附录页：追加后移到 Z 分隔页之前
    blank = next(l for l in prs.slide_layouts if l.name == "Blank")
    s = prs.slides.add_slide(blank)
    p_appendix(s, 0)
    new = len(prs.slides) - 1

    zidx = None
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip() == "Z":
                zidx = i
    if zidx is None:
        raise SystemExit("Z divider not found")

    lst = prs.slides._sldIdLst
    items = list(lst)
    for el in items:
        lst.remove(el)
    order = items[:zidx] + [items[new]] + items[zidx:new]
    for el in order:
        lst.append(el)

    # 页码：插入点之后 +1（附录页自身给插入位置的号）
    base = zidx + 1  # 1-based 物理页号
    for i, sl in enumerate(prs.slides):
        if i < zidx:
            continue
        for sh in sl.shapes:
            if not sh.has_text_frame or sh.left is None:
                continue
            if abs(sh.left / 914400 - 11.73) > 0.05 or sh.width / 914400 > 1.25:
                continue
            p = sh.text_frame.paragraphs[0]
            t = p.runs[0].text.strip() if p.runs else ""
            if not t.isdigit():
                continue
            if i == zidx:
                p.runs[0].text = str(base)
            else:
                p.runs[0].text = str(int(t) + 1)

    # D 分隔页页数与说明
    for sl in prs.slides:
        for sh in sl.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if "共 15 页" in r.text:
                        r.text = r.text.replace(
                            "共 15 页：D1 案例口径 · D2 中转站与代理池 · D3 护栏与破限",
                            "共 16 页：D1 案例口径 · D2 中转站与代理池 · D3 护栏与破限 · D4 开源平台对照")

    prs.save(dst)
    print("saved", dst, len(prs.slides))
    for i, sl in enumerate(prs.slides, 1):
        if 47 <= i <= 56 or 86 <= i <= 92:
            t = ""
            for sh in sl.shapes:
                if sh.has_text_frame and sh.name == "TextBox 4":
                    t = sh.text_frame.text.split("\n")[0]
            n = None
            for sh in sl.shapes:
                if sh.has_text_frame and sh.left is not None and \
                        abs(sh.left / 914400 - 11.73) < 0.05 and sh.width / 914400 < 1.25:
                    n = sh.text_frame.text
            print(f"  {i:3d} [{n}] {t[:46]}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
