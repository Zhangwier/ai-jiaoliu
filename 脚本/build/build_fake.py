# -*- coding: utf-8 -*-
"""在“个人 vs 企业”之前插一页：什么情况下企业 Agent 是伪需求。"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

from build_v3 import (put, rect, hline, vline, dot, chrome, section_bar,
                      callout, rowband, style_run, L, W,
                      C_TITLE, C_BODY, C_MUTED, C_TEAL, C_ORANGE, C_HEAD,
                      C_LINE, C_WHITE, C_WARM, C_COOL, C_ZEBRA)


def page_fake(slide, page):
    chrome(
        slide, page,
        "先说清楚：什么情况下，企业 Agent 是伪需求",
        "如果它只是“个人版加一层登录和审计”，那确实是——没人会为更慢的同一个东西付费",
        "来源：本稿对企业落地的观察；MOVO / AgenTeam 的能力边界；公开的企业 Agent 失败复盘；本稿整理",
        notes=(
            "讲解与会后阅读\n\n"
            "这一页是主动破题。现场一定有人心里想“企业 Agent 不就是给个人版套个壳吗”，"
            "与其等人问，不如自己先说。\n\n"
            "先说伪需求的三种典型，都是真实见过的死法：\n"
            "一，把通用聊天框搬进内网，加登录、加审计。能力一点没变，但多了限制、多了审批、"
            "还比个人版慢——用户为什么不掏手机用自己的？这种项目上线三个月就没人打开。\n"
            "二，一上来做“全公司知识库”。范围大到没人能验收，最后变成一个没人搜的搜索框。\n"
            "三，什么都交给 Agent。没有责任人、没有验收标准，出错没人认，业务方不敢用。\n\n"
            "再看真需求的四个特征，这是判断标准，不是概念：\n"
            "一，碰得到个人碰不到的东西——内部系统、图纸库、造价软件、财务数据。这是硬壁垒，"
            "个人版永远做不到。\n"
            "二，有量。一年几百上千份，人撑不住，才值得做成可重复的。一个人一天做三份，"
            "做成系统是浪费。\n"
            "三，产出可复算。数字对不对能验，否则没人敢签字，也就进不了正式流程。\n"
            "四，有明确责任人。谁签这个字，出事找谁——这条决定业务方敢不敢用。\n\n"
            "四条里缺两条以上，基本就是伪需求。可以直接拿它当筛子过一遍手上的场景。\n\n"
            "最后那句是全页落点，也是我自己的判断：企业 Agent 的增量不在“更聪明”，在“能进流程”。"
            "个人 Agent 的产出是“参考”，看一眼、信不信由你；企业 Agent 的产出是“交付物”，"
            "要签字、要归档、要被审计。身份、边界、留痕、账这四件，不是成本，是准入条件——"
            "没有它们，AI 的产出永远停在“参考”，价值从一百块掉到零。\n\n"
            "反过来说一句更狠的：如果一件事的产出不需要签字、不需要归档、错了也没人追，"
            "那它本来就该用个人版——这不是伪需求，是选错了工具。"
        ))

    # ---- 左：伪需求
    section_bar(slide, L, 1.86, 6.05, "三种伪需求：换汤不换药", color=C_ORANGE, h=0.36, size=13)
    fakes = [
        ("搬聊天框", "能力没变，只多了登录、审批、限速；用户为什么不掏手机用自己的", C_ORANGE),
        ("全公司知识库", "范围大到没人能验收，最后变成一个没人搜的搜索框", C_ORANGE),
        ("什么都交给 Agent", "没有责任人、没有验收标准，出错没人认，业务不敢用", C_ORANGE),
    ]
    y = 2.30
    for k, v, col in fakes:
        rect(slide, L, y, 6.05, 0.66, C_WHITE)
        rect(slide, L, y, 0.05, 0.66, col)
        put(slide, 0.76, y + 0.06, 1.70, 0.22, k, size=10.5, bold=True, color=C_TITLE,
            anchor=MSO_ANCHOR.TOP)
        put(slide, 2.54, y + 0.06, 3.92, 0.56, v, size=9, color=C_BODY, spacing=1.18,
            anchor=MSO_ANCHOR.TOP)
        y += 0.76

    rect(slide, L, 4.64, 6.05, 0.46, C_ZEBRA)
    put(slide, 0.72, 4.64, 5.72, 0.46,
        "共同点：都只是把个人用法搬到公司网络里，没有碰到任何“企业才有”的东西。",
        size=9, color=C_MUTED, spacing=1.15, anchor=MSO_ANCHOR.MIDDLE)

    rect(slide, L, 5.26, 6.05, 0.60, C_WARM)
    put(slide, 0.74, 5.30, 5.68, 0.22, "所以真正的分界线只有一条", size=9.5, bold=True,
        color=C_ORANGE, anchor=MSO_ANCHOR.TOP)
    put(slide, 0.74, 5.54, 5.68, 0.28,
        "产出要不要签字、要不要归档、错了要不要有人负责", size=9.5, color=C_BODY,
        anchor=MSO_ANCHOR.TOP)

    # ---- 右：真需求
    section_bar(slide, 6.95, 1.86, 5.78, "四个真需求：缺两条以上，就别做", color=C_TEAL,
                h=0.36, size=13)
    reals = [
        ("碰得到别人碰不到的", "内部系统、图纸库、造价软件、财务数据——硬壁垒，个人永远做不到"),
        ("有量", "一年几百上千份，人撑不住，才值得做成可重复的"),
        ("产出可复算", "数字对不对能验，否则没人敢签字，进不了正式流程"),
        ("有明确责任人", "谁签这个字、出事找谁——这条决定业务方敢不敢用"),
    ]
    y = 2.30
    for k, v in reals:
        dot(slide, 7.12, y + 0.12, 0.10, C_TEAL)
        put(slide, 7.30, y, 2.05, 0.22, k, size=10, bold=True, color=C_TEAL,
            anchor=MSO_ANCHOR.TOP)
        put(slide, 9.40, y - 0.01, 3.18, 0.48, v, size=9, color=C_BODY, spacing=1.15,
            anchor=MSO_ANCHOR.TOP)
        y += 0.56

    rect(slide, 6.95, 4.72, 5.78, 0.90, C_COOL)
    put(slide, 7.14, 4.76, 5.42, 0.22, "反过来判断：什么时候该用个人版", size=9.5,
        bold=True, color=C_TEAL, anchor=MSO_ANCHOR.TOP)
    put(slide, 7.14, 5.00, 5.42, 0.56,
        "产出不用签字、不用归档、错了也没人追——那就该用个人版。这不是伪需求，是选错了工具。",
        size=9, color=C_BODY, spacing=1.15, anchor=MSO_ANCHOR.TOP)

    callout(slide, "落点", "企业 Agent 的增量不在“更聪明”，在“能进流程”：从参考变成交付物")


def main(src, dst):
    prs = Presentation(src)
    blank = next(l for l in prs.slide_layouts if l.name == "Blank")

    s = prs.slides.add_slide(blank)
    page_fake(s, 50)
    new = len(prs.slides) - 1

    lst = prs.slides._sldIdLst
    items = list(lst)
    for el in items:
        lst.remove(el)
    order = items[:49] + [items[new]] + items[49:new]
    for el in order:
        lst.append(el)

    # index 49 之后（原 50 起）页码 +1
    for i, sl in enumerate(prs.slides):
        if i <= 49:
            continue
        for sh in sl.shapes:
            if not sh.has_text_frame or sh.left is None:
                continue
            if abs(sh.left / 914400 - 11.73) > 0.05 or sh.width / 914400 > 1.25:
                continue
            p = sh.text_frame.paragraphs[0]
            if p.runs and p.runs[0].text.strip().isdigit():
                p.runs[0].text = str(int(p.runs[0].text.strip()) + 1)

    prs.save(dst)
    print("saved", dst, len(prs.slides))
    for i, sl in enumerate(prs.slides, 1):
        if 47 <= i <= 56:
            n = None
            for sh in sl.shapes:
                if sh.has_text_frame and sh.left is not None and \
                        abs(sh.left / 914400 - 11.73) < 0.05 and sh.width / 914400 < 1.25:
                    n = sh.text_frame.text
            t = ""
            for sh in sl.shapes:
                if sh.has_text_frame and sh.name == "TextBox 4":
                    t = sh.text_frame.text.split("\n")[0]
            print(f"  {i:3d} [{n}] {t[:50]}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
