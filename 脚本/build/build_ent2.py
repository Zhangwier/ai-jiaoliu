# -*- coding: utf-8 -*-
"""企业落地三页重做：
   50 个人 Agent vs 企业 Agent（替换原来“四条路线”那页的位置）
   51 MOVO / AgenTeam 是什么，缺了哪几件
   52 工作流与 RAG 怎么养
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from build_v3 import (put, rect, hline, vline, dot, chrome, section_bar,
                      callout, rowband, node, style_run, L, W,
                      C_TITLE, C_BODY, C_MUTED, C_TEAL, C_ORANGE, C_HEAD,
                      C_LINE, C_WHITE, C_WARM, C_COOL, C_ZEBRA)


def clear(slide):
    spTree = slide.shapes._spTree
    for child in list(spTree):
        if child.tag in (qn("p:sp"), qn("p:cxnSp"), qn("p:pic"),
                         qn("p:graphicFrame"), qn("p:grpSp")):
            spTree.remove(child)


# ================================================================ 50 差异
def page_diff(slide, page):
    chrome(
        slide, page,
        "个人 Agent 和企业 Agent，差在哪四件事",
        "同一套模型、同一套提示词，换到公司里要重做的不是能力，是身份、边界、留痕和账",
        "来源：Scalekit《Access Control for Multi-Tenant AI Agents》（五种身份）；"
        "AWS 企业 Agentic AI 实践；MOVO / AgenTeam 产品机制；本稿整理",
        notes=(
            "讲解与会后阅读\n\n"
            "这一页是全节的入口，回答一个最常被跳过的问题：公司里用的 Agent，和个人用的到底差在哪。\n\n"
            "先说结论：能力上几乎一样，差在四件事。\n\n"
            "一，身份。个人用的时候，“我”就是全部——我登录、我授权、我负责。"
            "公司里一件事至少牵扯五个身份：谁触发的、用谁的凭据执行、谁授权的、属于哪个组织、"
            "最后在对方系统里记在谁名下。Scalekit 那篇文章把这五个身份列了出来，"
            "他们举的例子很典型：一家金融科技公司让 Agent 跨四个团队往 GitHub 建 issue，"
            "因为四个团队共用一个 OAuth token，结果 issue 全落到别的团队的仓库里，三个月才发现。"
            "这不是攻击，是“身份没建对”。\n\n"
            "二，边界。个人的边界是“我自己愿意给什么”。公司的边界是“这个人这个岗位能看哪些数据、"
            "能干哪些动作”，而且要能证明越不过去。所以企业版一定要有组织、角色、数据分级、"
            "工具白名单；缺一个，第一次出事的成本就够做十次。\n\n"
            "三，留痕。个人用完就完了。公司要能回答“上周三这个数字是谁让 AI 算的、依据哪份文件、"
            "中间改了哪一步”。这就是审批、执行追踪、产物留存、版本。\n\n"
            "四，账。个人按月订阅，超了就慢一点。公司要按人、按项目、按场景算成本和额度，"
            "还要能对上采购和发票。\n\n"
            "右边这张表可以直接对着讲：左边是“个人够用”，右边是“公司必须有”。"
            "注意最后一行的措辞——个人版不是“做得差”，是“不负责这几件事”。\n\n"
            "落到我们自己的工作：一张图纸、一份修编报告，如果只有一个人用 AI 做完然后自己签字，"
            "那还是个人用法；要变成公司用法，就得有第二个人能复核、有记录能追、"
            "有边界保证它没把别的项目资料带进来。"
        ))

    # 顶部一句话
    rect(slide, L, 1.86, W, 0.38, C_COOL)
    put(slide, 0.80, 1.86, 11.80, 0.38, [
        ("能力几乎一样　", 11.5, True, C_TEAL),
        ("差的不是模型，是“谁能用、能碰什么、出了事怎么查、钱怎么算”", 11.5, False, C_BODY)],
        spacing=1.0)

    # 四行对照
    rect(slide, L, 2.32, W, 0.38, C_HEAD)
    for x, w, t in [(0.72, 1.55, "差在哪"), (2.34, 4.40, "个人：够用就行"),
                    (6.94, 5.10, "公司：必须有")]:
        put(slide, x, 2.32, w, 0.38, t, size=12, bold=True, color=C_WHITE)

    rows = [
        ("身份", "一个人：我登录、我授权、我负责",
         "五个身份要分开：谁触发 / 用谁凭据执行 / 谁授权 / 属哪个组织 / 记在谁名下", C_ORANGE),
        ("边界", "我愿意给什么就给什么，越界了也只是我自己吃亏",
         "组织、角色、数据分级、工具白名单；要能证明“越不过去”，不是靠自觉", C_ORANGE),
        ("留痕", "用完就完了，答错重问一次",
         "审批、执行追踪、产物留存、版本：要能回答“谁让 AI 算的、依据哪份、改过哪步”", C_TEAL),
        ("账", "按月订阅，超了就慢一点",
         "按人 / 项目 / 场景算额度与成本；能对采购、发票和预算", C_TEAL),
    ]
    y = 2.70
    for i, (k, a, b, col) in enumerate(rows):
        rowband(slide, L, y, W, 0.86, i)
        rect(slide, L, y, 0.06, 0.86, col)
        put(slide, 0.74, y, 1.50, 0.86, k, size=14, bold=True, color=C_TITLE)
        put(slide, 2.34, y, 4.30, 0.86, a, size=10.5, color=C_MUTED, spacing=1.18)
        put(slide, 6.94, y, 5.05, 0.86, b, size=10.5, color=C_BODY, spacing=1.18)
        y += 0.86

    # 底部：一个真实教训
    rect(slide, L, 6.16, W, 0.24, C_WARM)
    put(slide, 0.80, 6.16, W - 0.34, 0.24, [
        ("真实教训：", 9, True, C_ORANGE),
        ("四个团队共用一个 token 往 GitHub 建 issue，因为身份没分，issue 全落错仓库，三个月才发现——这不是攻击，是没建身份",
         9, False, C_BODY)], spacing=1.0)

    callout(slide, "一句话", "个人版不负责“身份、边界、留痕、账”——这四件恰好是企业版的主要工作量")


# ================================================================ 51 两个项目是什么
def page_projects(slide, page):
    chrome(
        slide, page,
        "两个能直接部署的开源项目：先看清它们各自是什么",
        "MOVO 是“给 Agent 配公司”；AgenTeam 是“给公司配数字员工”——方向相反，补的件不一样",
        "来源：github.com/himovo/movo（README、架构图、能力表）；github.com/Stonewuu/agenteam（README）",
        notes=(
            "讲解与会后阅读\n\n"
            "这两个项目很容易被当成“差不多的东西”，其实出发点相反，看懂了就能判断自己该抄哪一层。\n\n"
            "MOVO：它的定位是“DSH 负责跑 Agent，MOVO 负责把它搬进企业”。所以它补的是**环境**："
            "12 个容器（网关、用户台、管理台、对话 API、管理 API、文档服务、DSH Runtime、"
            "Mongo、Redis、Weaviate、文档 worker、密钥引导），"
            "能力上把知识、文档、生成、Skill/Tool/MCP、审批审计、组织与角色都做进去。"
            "它假设“已经有一个 Agent 要跑”，问题是怎么让它合规地跑在单位里。\n\n"
            "AgenTeam（群策）：它的定位是“数字员工”——把模型、技能、知识库、工具组合成一个岗位，"
            "发布、上架、按资源授权，让团队成员“雇佣”它干活。"
            "所以它补的是**组织与节奏**：项目、待办、定时任务、企业微信 / 飞书通知、逐人回执、"
            "版本与上架。技术栈是 Java 21 + Spring Boot + AgentScope 2，MySQL 存业务、Redis 存会话，"
            "另有一个带浏览器与 Git 的沙箱。\n\n"
            "一句话对比：MOVO 回答“这个 Agent 能不能在单位里安全地跑”；"
            "AgenTeam 回答“这个岗位的活，能不能交给一个数字同事，并且有人跟进”。\n\n"
            "对我们更有用的是下半页那张“缺哪件”的表：不管选哪个，单位自己还得补四样——"
            "数据分级、验收标准、评测题集、以及“谁负责养”。这四样任何开源项目都给不了。"
        ))

    # ---- 左：MOVO
    section_bar(slide, L, 1.86, 6.05, "MOVO：给 Agent 配一间合规的办公室", h=0.36, size=13)
    put(slide, L, 2.28, 6.05, 0.46, [
        ("一句话　", 9.5, True, C_TEAL),
        ("DSH 负责跑 Agent，MOVO 负责把它搬进企业", 9.5, False, C_BODY)],
        spacing=1.1, anchor=MSO_ANCHOR.TOP)
    put(slide, L, 2.76, 6.05, 0.44,
        "12 个容器：网关 · 用户台 · 管理台 · 对话 API · 管理 API · 文档服务 · DSH Runtime · "
        "Mongo / Redis / Weaviate",
        size=8.5, color=C_MUTED, spacing=1.15, anchor=MSO_ANCHOR.TOP)

    movo = [
        ("补的是环境", "身份、权限、审批、审计、交付物——让 Agent 能合规地跑在单位里"),
        ("能力面", "知识检索与多轮研究、文档解析（含图表）、报告/PPTX/表格生成"),
        ("治理面", "敏感动作审批、执行追踪、产物留存、组织/角色/模型/知识/审计"),
        ("适合谁", "已有 Agent 或想用现成 Agent 能力，缺的是“能上生产的那一层”"),
    ]
    y = 3.28
    for i, (k, v) in enumerate(movo):
        rowband(slide, L, y, 6.05, 0.66, i)
        put(slide, 0.74, y, 1.40, 0.66, k, size=10.5, bold=True, color=C_TEAL,
            spacing=1.05)
        put(slide, 2.18, y, 4.30, 0.66, v, size=9.5, color=C_BODY, spacing=1.18)
        y += 0.66

    # ---- 右：AgenTeam
    section_bar(slide, 6.95, 1.86, 5.78, "AgenTeam：给公司配一群数字员工", color=C_ORANGE,
                h=0.36, size=13)
    put(slide, 6.95, 2.28, 5.78, 0.46, [
        ("一句话　", 9.5, True, C_ORANGE),
        ("数字员工 ＝ 模型＋技能＋知识库＋工具，发布、上架、按资源授权", 9.5, False, C_BODY)],
        spacing=1.1, anchor=MSO_ANCHOR.TOP)
    put(slide, 6.95, 2.76, 5.78, 0.44,
        "Java 21 + Spring Boot + AgentScope 2；MySQL 存业务、Redis 存会话；带浏览器与 Git 的沙箱",
        size=8.5, color=C_MUTED, spacing=1.15, anchor=MSO_ANCHOR.TOP)

    agt = [
        ("补的是组织", "岗位化、发布上架、资源授权、版本——把 Agent 变成“同事”"),
        ("节奏", "按项目组织；一次性 / 每日 / 每周 / 每月定时任务"),
        ("接入", "企业微信、飞书：账号绑定、业务通知、逐人回执"),
        ("适合谁", "想把具体岗位的活交出去，并且要有人跟进、能查历史"),
    ]
    y = 3.28
    for i, (k, v) in enumerate(agt):
        rowband(slide, 6.95, y, 5.78, 0.66, i)
        put(slide, 7.12, y, 1.40, 0.66, k, size=10.5, bold=True, color=C_ORANGE,
            spacing=1.05)
        put(slide, 8.56, y, 4.02, 0.66, v, size=9.5, color=C_BODY, spacing=1.18)
        y += 0.66

    # ---- 底：自己还得补的四样
    rect(slide, L, 6.06, W, 0.34, C_WARM)
    put(slide, 0.80, 6.06, W - 0.34, 0.34, [
        ("两个都装上也补不齐的四样：", 9, True, C_ORANGE),
        ("① 数据分级（哪些能出单位）　② 验收标准（什么算做完）　③ 评测题集（改了会不会变差）　"
         "④ 谁负责养（每周看返工与检索日志）", 9, False, C_BODY)], spacing=1.0)

    callout(slide, "怎么选", "缺合规与治理 → 抄 MOVO 那层；缺协同与跟进 → 抄 AgenTeam 那层", kind="cool")


# ================================================================ 52 怎么养
def page_ops(slide, page):
    chrome(
        slide, page,
        "真正难的不是搭起来，是让它越用越准",
        "工作流和 RAG 都得“养”：把日常返工和失败变成规则，否则上线三个月就僵住",
        "来源：本稿两案例的返工记录；MOVO / AgenTeam 的版本与审计机制；检索评测常用做法；本稿整理",
        notes=(
            "讲解与会后阅读\n\n"
            "这一页是这三页里最实操的，也是我自己的判断：企业 AI 项目很少死在“搭不起来”，"
            "多数死在“没人养”。\n\n"
            "先讲工作流这半边，四个动作构成一个环：\n"
            "① 记录：每次人工返工记一条——哪一步、改了什么、为什么。最容易做，也最容易漏；\n"
            "② 提炼：同一类返工出现三次，就升级成一条标准或一个校验脚本（比如“合计必须与分项一致”）；\n"
            "③ 清理：每季度过一遍旧规则，删掉为老模型打的补丁——不删就变成天花板；\n"
            "④ 版本化：工作流要有版本号和变更记录，谁改的、为什么改、影响哪些场景。没版本就没法回滚。\n\n"
            "再讲 RAG 这半边，关键是“把失败分类”，别笼统说“答得不好”：\n"
            "A 没检索到 → 动切分粒度、补别名同义词、加元数据过滤；\n"
            "B 检索到但没用 → 加 rerank、调 top-k、把片段压短；\n"
            "C 用了但答错 → 改提示词约束、要求引用原文、加二次核对。\n"
            "三类混在一起，就永远不知道该改哪。\n\n"
            "然后是评测：沉淀 50–100 道真实问题，每次改切分、换模型、动提示词都重跑一遍。"
            "没有这套题，所有优化都只是“感觉好一点”。\n\n"
            "最后是组织：得有人负责这件事——不是 IT 一个部门，而是业务里最会用 AI 的那位，"
            "每周花两小时看返工记录和检索日志。这一条比任何工具都重要。\n\n"
            "收尾：把使用变成数据，把数据变成规则。听着朴素，但九成单位没做。"
        ))

    section_bar(slide, L, 1.86, 6.05, "工作流怎么长：四个动作，一个环", h=0.36, size=13)
    loop = [
        ("① 记录", "每次人工返工记一条：哪一步、改了什么、为什么", C_TEAL),
        ("② 提炼", "同类返工出现三次 → 升级成一条标准或一个校验脚本", C_TEAL),
        ("③ 清理", "每季度删旧规则，尤其“当年绕坑”的补丁", C_ORANGE),
        ("④ 版本化", "带版本号与变更记录：谁改、为何改、影响谁", C_ORANGE),
    ]
    vline(slide, 0.72, 2.34, 2.84, C_LINE, 1.0)
    y = 2.30
    for k, v, col in loop:
        dot(slide, 0.72, y + 0.13, 0.12, col)
        put(slide, 0.98, y, 5.45, 0.24, k, size=11.5, bold=True, color=col,
            anchor=MSO_ANCHOR.TOP)
        put(slide, 0.98, y + 0.25, 5.45, 0.46, v, size=9.5, color=C_BODY,
            spacing=1.15, anchor=MSO_ANCHOR.TOP)
        y += 0.72

    put(slide, 0.98, 5.24, 5.60, 0.54,
        "判断标准：一条规则半年没人碰过、也没引发过返工，就该删——它只是在拖慢模型。",
        size=9, color=C_MUTED, spacing=1.15, anchor=MSO_ANCHOR.TOP)

    section_bar(slide, 6.95, 1.86, 5.78, "RAG 怎么调：先把失败分成三类", color=C_ORANGE,
                h=0.36, size=13)
    fails = [
        ("A 没检索到", "召回问题", "切分粒度、补别名同义词、加元数据过滤", C_ORANGE),
        ("B 检索到没用", "排序问题", "加 rerank、调 top-k、把片段压短", C_TEAL),
        ("C 用了但答错", "生成问题", "改约束、要求引原文、加二次核对", C_HEAD),
    ]
    y = 2.30
    for k, tag, fix, col in fails:
        rect(slide, 6.95, y, 5.78, 0.78, C_WHITE)
        rect(slide, 6.95, y, 0.05, 0.78, col)
        put(slide, 7.10, y + 0.07, 2.05, 0.22, k, size=10.5, bold=True, color=C_TITLE,
            anchor=MSO_ANCHOR.TOP)
        put(slide, 9.20, y + 0.07, 0.95, 0.22, tag, size=8, color=C_MUTED,
            align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.TOP)
        put(slide, 7.10, y + 0.31, 5.45, 0.42, "→ " + fix, size=9, color=C_BODY,
            spacing=1.15, anchor=MSO_ANCHOR.TOP)
        y += 0.86

    rect(slide, 6.95, 4.94, 5.78, 0.86, C_COOL)
    put(slide, 7.12, 4.98, 5.48, 0.22, "评测集：把优化从“感觉”变成“数”", size=9.5,
        bold=True, color=C_TEAL, anchor=MSO_ANCHOR.TOP)
    put(slide, 7.12, 5.22, 5.48, 0.52,
        "沉淀 50–100 道真实问题；每次改切分、换模型、动提示词都重跑一遍，分数掉了就回退",
        size=9, color=C_BODY, spacing=1.15, anchor=MSO_ANCHOR.TOP)

    rect(slide, L, 5.90, W, 0.34, C_WARM)
    put(slide, 0.80, 5.90, W - 0.34, 0.34, [
        ("谁来做：", 9.5, True, C_ORANGE),
        ("得有一个“养”的人——不是 IT 一个部门，而是业务里最会用 AI 的那位，每周两小时看返工记录与检索日志",
         9.5, False, C_BODY)], spacing=1.0)

    callout(slide, "一句话", "把“使用”变成数据，把“数据”变成规则——九成单位没做这一步", kind="cool")


# ================================================================ main
def main(src, dst):
    prs = Presentation(src)

    for idx, fn in ((48, page_diff), (49, page_projects), (50, page_ops)):
        footer = None
        for sh in prs.slides[idx].shapes:
            if sh.has_text_frame and sh.left is not None:
                if abs(sh.left / 914400 - 11.73) < 0.05 and sh.width / 914400 < 1.25:
                    footer = sh.text_frame.text.strip()
        clear(prs.slides[idx])
        fn(prs.slides[idx], footer or (idx + 1))
        print("rebuilt", idx + 1, "footer", footer)

    # 04 扉页副标题
    s = prs.slides[38]
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith("为什么要用 RAG"):
            tf = sh.text_frame
            for para in tf.paragraphs:
                for r in list(para.runs):
                    r._r.getparent().remove(r._r)
            run = tf.paragraphs[0].add_run()
            run.text = ("为什么要用 RAG、怎么整理成工作流、多 Agent 何时值、"
                        "个人与企业 Agent 差在哪、企业实际怎么养——五个问题。")
            rPr = run._r.get_or_add_rPr()
            rPr.set("sz", "1400")
            rPr.set("b", "0")
            rPr.append(parse_xml(
                '<a:solidFill %s><a:srgbClr val="E6ECF5"/></a:solidFill>' % nsdecls("a")))
            for tag in ("a:latin", "a:ea", "a:cs"):
                rPr.append(parse_xml(
                    '<%s %s typeface="微软雅黑"/>' % (tag, nsdecls("a"))))
            tf.paragraphs[0].line_spacing = 1.2

    # 右侧目录：数据分级那行改成“个人 vs 企业”
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("数据分级"):
            sh.text_frame.paragraphs[0].runs[0].text = "个人 vs 企业 Agent"

    prs.save(dst)
    print("saved", dst, len(prs.slides))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
