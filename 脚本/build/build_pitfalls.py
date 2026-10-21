# -*- coding: utf-8 -*-
"""新增一页：工作流与 Agent 编排的三个坑（第 47 页，紧跟多 Agent 那页）"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

from build_v3 import (put, rect, hline, vline, dot, chrome, section_bar,
                      callout, rowband, L, W, C_TITLE, C_BODY, C_MUTED,
                      C_TEAL, C_ORANGE, C_HEAD, C_LINE, C_WHITE, C_WARM,
                      C_COOL, C_ZEBRA)


def page_pitfalls(slide, page):
    chrome(
        slide, page,
        "工作流和 Agent 编排，最容易踩的三个坑",
        "前面讲的是怎么搭；这一页讲它什么时候会害你——三条都是本稿踩过的",
        "来源：本稿两案例的执行与返工记录；多 Agent 协作记录；本稿整理",
        notes=(
            "讲解与会后阅读\n\n"
            "这一页是全稿里最“反自己”的一页：前面花了很大力气讲工作流和多 Agent，"
            "这里要说清它们什么时候会害人。三条都是本稿真实踩过的，不是转述别人的说法。\n\n"
            "第一条：低级模型加一堆工作流，不如高级模型一句话。\n"
            "原因是：工作流不会增加模型的理解力。它只是把任务切小、把要求写细。"
            "如果模型本来就听不懂“这一步到底要什么”，切成五步只是把同一个误解重复五遍，"
            "而且每一步都多一次出错机会。实测里最典型的是图纸审查：配置差的那几组，"
            "不是步骤不够细，是根本没看懂图。反过来，Opus 那一档一句话下去，"
            "三处关键问题全中——它不需要被拆成五步。\n"
            "所以顺序应该是：先确认模型够不够用，再谈工作流。"
            "模型不够，工作流只能让你更快地得到错的东西。\n\n"
            "第二条：工作流如果不逐步人为确认，很容易花大量时间产出一堆垃圾，"
            "而且事后看不出是哪一步坏的。\n"
            "这条我体会最深。一条五步流水线跑完，交出来的东西不对，你想回头查："
            "是第一步材料读错了？第二步切分切坏了？还是第三步判断偏了？"
            "中间产物没留，就查不出来。于是只能整条重跑，时间全耗在重跑上。\n"
            "所以工作流的第一个设计原则不是“自动”，是“每步留一个能查的产物”。"
            "这也是我不喜欢“Agent 编排”的原因：编排把步骤藏在黑箱里，"
            "中间态不落地，出问题只能靠猜。\n\n"
            "第三条：低级 AI 理解不了你的意思，换另一个 AI 来看也理解不了。\n"
            "很多人以为“让另一个模型复核”能救。但如果复核者拿到的还是同一段模糊描述，"
            "它只会给出另一个同样模糊的判断——两个都不懂，加起来还是不懂。"
            "复核能起作用的前提是：被复核的东西本身有据可查（原文、页码、算式），"
            "复核者能回到材料上逐条对。否则“两个模型都说对”只是两次猜测。\n\n"
            "三条合起来是一句话：模型的理解力是地基，工作流是脚手架。"
            "地基不够，脚手架搭得再漂亮也没用。"
        ))

    # 三条坑，纵向
    pits = [
        ("坑一", "低级模型 ＋ 一堆工作流 ＜ 高级模型一句话",
         "工作流不增加理解力，只把同一个误解重复五遍，还多四次出错机会。"
         "实测：差的那几组不是步骤不够细，是没看懂图；Opus 一句话三处全中。",
         "先确认模型够不够用，再谈工作流", C_ORANGE),
        ("坑二", "不逐步确认，跑完一堆垃圾，还查不出哪步坏的",
         "五步跑完结果不对，中间产物没留，就只能整条重跑。"
         "时间全耗在重跑，而不是定位。",
         "每步必须留一个能查的产物——这也是“不喜欢 Agent 编排”的根源", C_ORANGE),
        ("坑三", "这个 AI 听不懂，换另一个 AI 看也听不懂",
         "复核者拿到的还是同一段模糊描述，只会给出另一个同样模糊的判断。"
         "两个都不懂，加起来还是不懂。",
         "复核的前提是：材料本身有据可查，能回到原文逐条对", C_HEAD),
    ]

    y = 1.92
    for tag, title, why, fix, col in pits:
        rect(slide, L, y, W, 1.30, C_WHITE)
        rect(slide, L, y, 0.06, 1.30, col)
        put(slide, 0.76, y + 0.08, 0.72, 0.24, tag, size=10, bold=True, color=col,
            anchor=MSO_ANCHOR.TOP)
        put(slide, 1.54, y + 0.06, 6.20, 0.30, title, size=13, bold=True,
            color=C_TITLE, spacing=1.05, anchor=MSO_ANCHOR.TOP)
        put(slide, 1.54, y + 0.44, 6.20, 0.78, why, size=9.5, color=C_BODY,
            spacing=1.2, anchor=MSO_ANCHOR.TOP)
        # 右侧对策
        rect(slide, 8.10, y + 0.12, 4.55, 1.06, C_ZEBRA)
        put(slide, 8.24, y + 0.18, 4.30, 0.20, "怎么避", size=8.5, bold=True,
            color=C_TEAL, anchor=MSO_ANCHOR.TOP)
        put(slide, 8.24, y + 0.40, 4.30, 0.72, fix, size=9, color=C_BODY,
            spacing=1.16, anchor=MSO_ANCHOR.TOP)
        y += 1.40

    rect(slide, L, 6.14, W, 0.26, C_COOL)
    put(slide, 0.78, 6.14, W - 0.34, 0.26,
        "顺序不能反：先确认模型够不够用，再决定要不要拆步骤、拆几个 Agent",
        size=9.5, color=C_BODY, spacing=1.05)

    callout(slide, "一句", "模型的理解力是地基，工作流是脚手架——地基不够，脚手架再漂亮也没用")


def main(src, dst):
    prs = Presentation(src)
    blank = next(l for l in prs.slide_layouts if l.name == "Blank")

    def title(i):
        for sh in prs.slides[i].shapes:
            if sh.has_text_frame and sh.name == "TextBox 4":
                return sh.text_frame.text.split("\n")[0]
        return ""

    s = prs.slides.add_slide(blank)
    page_pitfalls(s, 0)

    # 插到「多个 Agent 协作」之后
    anchor = next(i for i in range(len(prs.slides))
                  if title(i).startswith("多个 Agent 协作"))
    lst = prs.slides._sldIdLst
    items = list(lst)
    el = items[-1]
    lst.remove(el)
    items = list(lst)
    items[anchor + 1].addprevious(el)
    print("inserted after", title(anchor)[:24])

    # 页码重排（隐藏页不参与）
    page = 0
    for sl in prs.slides:
        if sl._element.get("show") != "0":
            page += 1
        for sh in sl.shapes:
            if not sh.has_text_frame or sh.left is None:
                continue
            if abs(sh.left / 914400 - 11.73) > 0.05 or sh.width / 914400 > 1.3:
                continue
            p = sh.text_frame.paragraphs[0]
            if p.runs and p.runs[0].text.strip().isdigit():
                p.runs[0].text = str(page)

    # 04 扉页索引补一行
    for sl in prs.slides:
        big = [sh for sh in sl.shapes if sh.has_text_frame and
               sh.text_frame.text.strip() == "04" and sh.width and sh.width / 914400 > 3]
        if not big:
            continue
        rows = [sh for sh in sl.shapes if sh.has_text_frame and sh.left and
                sh.left / 914400 > 9 and sh.text_frame.text.strip() and
                sh.text_frame.text.strip() != "这一部分"]
        rows.sort(key=lambda x: x.top)
        import copy
        anchor_row = next((x for x in rows if "工作流" in x.text_frame.text), None)
        if anchor_row is not None:
            new_el = copy.deepcopy(anchor_row._element)
            anchor_row._element.addnext(new_el)
            new_sh = [x for x in sl.shapes if x._element is new_el][0]
            new_sh.text_frame.paragraphs[0].runs[0].text = "工作流的三个坑"
            rows = [sh for sh in sl.shapes if sh.has_text_frame and sh.left and
                    sh.left / 914400 > 9 and sh.text_frame.text.strip() and
                    sh.text_frame.text.strip() != "这一部分"]
            rows.sort(key=lambda x: x.top)
            base = rows[0].top
            for i, x in enumerate(rows):
                x.top = base + int(i * 0.56 * 914400)
            print("rail rows:", [x.text_frame.text[:14] for x in rows])
        break

    prs.save(dst)
    print("saved", dst, len(prs.slides))
    for i, sl in enumerate(prs.slides, 1):
        if 44 <= i <= 50:
            print(f"  {i:3d} {title(i-1)[:50]}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
