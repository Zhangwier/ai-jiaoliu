# -*- coding: utf-8 -*-
"""
重做第 40 页：从三个例子提炼"什么样的活适合交给 AI"的判据；
并给第 41 页脱敏（去掉项目名与真实文件名）。

第 40 页旧内容整页重建；第 41 页只做文本替换。
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

F = r"D:\MG\program\ai交流\AI交流分享-成品26.pptx"
FONT = "微软雅黑"

C_TITLE = "14213D"
C_TXT = "2B3445"
C_MUTED = "6B7385"
C_TEAL = "0F766E"
C_ORANGE = "D9622B"
C_BG_SOFT = "E1F0EE"
C_BG_WARN = "FBE9DE"
C_DARK = "12203A"
C_WHITE = "FFFFFF"
C_PAGE_BG = "F6F5F1"
C_LINE = "DAD7CE"

NSMAP = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
         'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
         'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')


def rect(s, x, y, w, h, color, border=None, bw=0.8):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)
    if border:
        sh.line.color.rgb = RGBColor.from_string(border)
        sh.line.width = Pt(bw)
    else:
        sh.line.fill.background()
    return sh


def tb(s, x, y, w, h, text, size=12, bold=False, color=C_TXT,
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.2, mono=False):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = line
        f = r.font
        f.name = "Consolas" if mono else FONT
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = RGBColor.from_string(color)
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            if rPr.find(qn(tag)) is None:
                rPr.append(parse_xml('<%s %s typeface="%s"/>' % (tag, NSMAP, FONT)))
    return box


prs = Presentation(F)

# ======================================================= 重建第 40 页
s = prs.slides[39]

# 清空形状（保留 nvGrpSpPr / grpSpPr）
spTree = s.shapes._spTree
for child in list(spTree):
    if child.tag.endswith("}nvGrpSpPr") or child.tag.endswith("}grpSpPr"):
        continue
    spTree.remove(child)

# 背景沿用（若已丢失则补上）
cSld = s._element.find(qn("p:cSld"))
if cSld.find(qn("p:bg")) is None:
    cSld.insert(0, parse_xml(
        '<p:bg %s><p:bgPr><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
        '<a:effectLst/></p:bgPr></p:bg>' % (NSMAP, C_PAGE_BG)))

# 页头
tb(s, 0.6, 0.32, 6.0, 0.3, "04  怎么用得稳", 12, True, C_TEAL)
tb(s, 11.73, 0.32, 1.0, 0.3, "40", 12, False, C_MUTED, align=PP_ALIGN.RIGHT)
rect(s, 0.6, 0.66, 0.5, 0.06, C_TEAL)
ln = s.shapes.add_connector(1, Inches(0.6), Inches(7.0), Inches(12.73), Inches(7.0))
ln.line.color.rgb = RGBColor.from_string("E4E7EB")
ln.line.width = Pt(0.5)
tb(s, 0.6, 7.06, 10.0, 0.3, "来源：本人日常任务记录（2026）", 8.5, False, "94A3B8")
tb(s, 10.73, 7.06, 2.0, 0.3, "来源索引 ›", 8.5, False, "94A3B8")

# 标题
tb(s, 0.6, 0.74, 12.13, 0.58, "什么样的活，适合交给 AI", 27, True, C_TITLE,
   anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
tb(s, 0.6, 1.34, 12.13, 0.36,
   "从我手上这些活里抽出来的五条——符合就交出去，不符合就自己干", 16.5, False, C_MUTED)

# 五条判据
rules = [
    ("01", "格式是死的", "章节、列、单元格都已定好，它只负责往里填"),
    ("02", "答案唯一", "同一份材料，谁做都该是这个数；有争议的判断不交"),
    ("03", "对错能当场验", "合计、条目、位置，打开就能看；验得比做得还快"),
    ("04", "依据在材料里", "答案写在文件上，它只负责找出来、搬过去"),
    ("05", "重复到人会烦", "几百个文件、上百条子目，人做容易漏、也容易烦"),
]
y = 1.84
for num, name, desc in rules:
    rect(s, 0.6, y, 12.13, 0.64, C_WHITE, border=C_LINE)
    rect(s, 0.6, y, 0.06, 0.64, C_TEAL)
    tb(s, 0.78, y, 0.45, 0.64, num, 10.5, True, C_TEAL, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 1.28, y, 2.4, 0.64, name, 14, True, C_TITLE, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 3.85, y, 8.7, 0.64, desc, 11.5, False, C_TXT, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.72

# 例子条（降级为佐证，不写项目名与文件名）
rect(s, 0.6, 5.5, 12.13, 0.36, C_BG_SOFT)
rect(s, 0.6, 5.5, 0.06, 0.36, C_TEAL)
tb(s, 0.78, 5.5, 0.7, 0.36, "举例", 9, True, C_TEAL, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 1.5, 5.5, 11.1, 0.36,
   "表格里的公式与重算 · 一堆文件里定位某一份 · 按模板把定稿数据填进报告",
   10, False, C_TXT, anchor=MSO_ANCHOR.MIDDLE)

# 结论
rect(s, 0.6, 6.3, 12.13, 0.55, C_BG_WARN)
rect(s, 0.6, 6.3, 0.09, 0.55, C_ORANGE)
tb(s, 0.85, 6.3, 1.05, 0.55, "一句", 14, True, C_ORANGE, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 1.78, 6.3, 10.7, 0.55, "别照着例子找活，照着判据筛活", 20, True, C_TITLE,
   anchor=MSO_ANCHOR.MIDDLE)

s.notes_slide.notes_text_frame.text = """讲解与会后阅读

什么样的活，适合交给 AI

这一页不列清单，列判据。前面举过的那些活只是恰好符合下面几条；换一份工作、换一个项目，照样用这五条去筛。

1. 格式是死的。章节、表格列、单元格位置都事先定好，它只负责往里填。要不要改结构、该分几章，仍然是人的判断——这部分交给它，往往就发散。

2. 答案唯一。同一份材料，不同人做出来应该是一个数。凡是"看情况""要商量"的判断，交给它只会得到一个看起来很确定的错答案。

3. 对错能当场验。合计对不对、条目少没少、位置对不对，打开文件就能看出来。验证成本低于重做成本，这类活才值得交出去。

4. 依据在材料里。答案写在给它的文件上，它只负责找出来、搬过去。不能靠它"知道"——它不知道的时候也会往下写。

5. 重复到人会烦。几百个文件里找一份、上百条子目逐条核对，机械重复、量大，人做容易漏也容易烦；这类活它不累。

用法：五条全中，直接交；中三四条，交出去但加强核对；只中一两条，自己干更快。

注意分寸：这五条讲的是"这活适不适合"，不是"这个模型行不行"。模型强弱决定它做得多细，判据决定该不该交给它。

依据：本人日常任务记录（2026）。"""

# ======================================================= 第 41 页脱敏
s41 = prs.slides[40]
REPL = {
    "TJ1 标段 101 个文件散在十几个分册里；先让它扫成一张表，之后每次找图只要筛一行":
        "一个标段一百多个文件散在十几个分册里；先让它扫成一张表，之后每次找图只要筛一行",
    "来源：_TJ1文件索引.xlsx（2026-10-14 生成）；make_file_index.py；南湛高速 TJ1 标段施工图":
        "来源：本地文件索引脚本与生成结果（2026-10）",
    "005_控制成果表.xlsx": "控制成果表.xlsx",
    "010_逐桩坐标表.xlsx": "逐桩坐标表.xlsx",
    "013_砍树挖根数量表.xls": "砍树挖根数量表.xls",
    "2.0 第二册 路线": "第二册 路线",
    "2.0 第二册 路线 · 26 个文件 · 0.98 MB · 9 个数量表      3.1 第三册 第一分册 路基 · 18 个 · 4.12 MB · 11 个数量表":
        "第二册 路线 · 26 个文件 · 9 个数量表        第三册第一分册 路基 · 18 个 · 11 个数量表",
    "实际表共 9 列：另含子目录、大小、修改时间、相对路径（可点击直接打开）":
        "实际表共 9 列，另含子目录、大小、修改时间，和可点击的相对路径",
}
n = 0
for sh in s41.shapes:
    if not sh.has_text_frame:
        continue
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if r.text in REPL:
                r.text = REPL[r.text]
                n += 1
print("slide41 runs replaced:", n)

s41.notes_slide.notes_text_frame.text = """讲解与会后阅读

拿"找图纸"举例：先建索引，再检索

1. 痛点：一个标段一百多个图纸和数量表文件，散在十几个分册目录里。找一张数量表，要么翻目录文件，要么按文件名猜，实际常常要花十几分钟。

2. 做法：先扫一遍根目录，生成一张索引表。表里九列：序号、文件名、类型、出图编号、所在册/分册、子目录、大小、修改时间、相对路径。备份文件和临时文件跳过；出图编号来自各册目录，目录没提供的就留空。

3. 另一张「册汇总」：按册统计文件数、体量、各类文件各几个。先看它能判断该翻哪一册，不用一个个点开。

4. 用法：按册筛、按出图编号定位、按类型找同类。表里没收录的，直接问它——它会去翻目录和分册表，再告诉你大概在哪一本。

5. 这就是最简单的 RAG 思路：先把资料整理成可检索的结构，再让模型在命中的片段上回答，而不是把整个文件夹塞进去。区别在于这里检索的是文件名和出图编号，不是向量。

6. 可复现：换项目只改一个根目录参数重跑脚本。这也是它的价值——索引能重生成、能核对，不是一次性的死表。

依据：本地文件索引脚本与生成结果（2026-10）。"""

prs.save(F)
print("saved")
