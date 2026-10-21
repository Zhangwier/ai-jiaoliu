# -*- coding: utf-8 -*-
"""更新第 39 页（04 章节过渡页）的导语与目录项，反映新增的两页。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

F = r"D:\MG\program\ai交流\AI交流分享-成品26.pptx"
FONT = "微软雅黑"
NSMAP = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
         'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
         'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')


def style(run, size, bold, color):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        if rPr.find(qn(tag)) is None:
            rPr.append(parse_xml('<%s %s typeface="%s"/>' % (tag, NSMAP, FONT)))


def add_tb(slide, x, y, w, h, text, size, bold, color):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 1.15
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    r = p.add_run()
    r.text = text
    style(r, size, bold, color)
    return box


prs = Presentation(F)
s = prs.slides[38]

# 1) 导语
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("为什么要用 RAG"):
        r = sh.text_frame.paragraphs[0].runs[0]
        r.text = ("先说我日常拿它干的三件活——公式、查图、填报告；"
                  "再看为什么要用 RAG、怎么整理成工作流、多 Agent 何时值、企业怎么从试点走到交付。")

# 2) 目录项：收集右侧那一列（x≈9.1）的文本框，按 y 排序
items = []
for sh in s.shapes:
    if not sh.has_text_frame or sh.left is None:
        continue
    if abs(sh.left / 914400 - 9.1) < 0.05 and sh.top / 914400 > 1.9:
        items.append(sh)
items.sort(key=lambda sh: sh.top)
print("found items:", [sh.text_frame.text for sh in items])

new_texts = [
    "我日常在做什么：三件活",
    "找图纸：先建索引再检索",
    "方法总览：对应哪个问题",
    "提示词：写成任务书",
    "上下文管理",
    "资料库（RAG）",
    "工作流 / Agent / 多 Agent",
    "工作流的三个坑",
    "企业：从试点到交付",
]

# 现有的 7 个塞进前 7 条；多出来的 2 条新建
for sh, txt in zip(items, new_texts):
    runs = sh.text_frame.paragraphs[0].runs
    runs[0].text = txt
    for r in runs[1:]:
        r._r.getparent().remove(r._r)

# 重排 y：起始 2.02，间距 0.475，共 9 条
y0, step, h = 2.02, 0.475, 0.42
for i, sh in enumerate(items):
    sh.top = int(Inches(y0 + i * step))

extra = new_texts[len(items):]
for j, txt in enumerate(extra):
    idx = len(items) + j
    add_tb(s, 9.1, y0 + idx * step, 3.8, h, txt, 17, False, "E6ECF5")

prs.save(F)
print("divider updated; items now", len(new_texts))
