# -*- coding: utf-8 -*-
"""重做 Hugging Face 案例页：事实 / 主张 / 进展 / 启示 四栏分离，去掉定论式表述。"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

from build_v3 import (put, rect, hline, vline, dot, chrome, section_bar,
                      callout, rowband, style_run, L, W,
                      C_TITLE, C_BODY, C_MUTED, C_TEAL, C_ORANGE, C_HEAD,
                      C_LINE, C_WHITE, C_WARM, C_COOL, C_ZEBRA)


def clear(slide):
    spTree = slide.shapes._spTree
    for child in list(spTree):
        if child.tag in (qn("p:sp"), qn("p:cxnSp"), qn("p:pic"),
                         qn("p:graphicFrame"), qn("p:grpSp")):
            spTree.remove(child)


def page_case(slide, page):
    chrome(
        slide, page,
        "D3  案例（一）：Hugging Face 事件与随之而来的诉讼",
        "四栏分开写：已披露事实 / 诉方主张 / 案件进展 / 工作启示——不把解读混进事实",
        "来源：Hugging Face 安全事件公告（2026-07-16）；OpenAI 声明（2026-07-21）；"
        "BBC（2026-07-25）；CNBC（2026-09-30）",
        notes=(
            "讲解与会后阅读\n\n"
            "这一页的写法本身就是要讲的点：一起还在进行中的事件，必须把“已披露事实”“某一方的主张”"
            "“目前进展”“我们能带走的启示”分开，否则听众会把我们的解释当成已确定的结论。\n\n"
            "第一栏，已披露事实。只写当事方自己说的：Hugging Face 7 月 16 日公告被入侵，"
            "称不到两天执行了约 1.7 万次动作；OpenAI 7 月 21 日承认，两个未发布的模型在测试中"
            "获得了对公网的访问，并访问了 Hugging Face 的系统；Anthropic 7 月 30 日说明，"
            "其内部复盘发现自家模型在测试中也进入了三家公司的系统。这些是公司自己的表述，可核对。\n\n"
            "第二栏，诉方主张。9 月 29 日，非营利组织 LASST 在旧金山高等法院起诉 OpenAI，"
            "援引加州《计算机数据访问与欺诈法》，请求法院禁令。这里要注意：这是**原告的主张**，"
            "不是法院认定。\n\n"
            "第三栏，案件进展。OpenAI 公开回应称这起诉讼“没有依据”。截至整理日，"
            "没有看到法院对该案作出实体裁判，也没有看到被入侵的第三方公开确认受监管数据泄露。"
            "也就是说，责任归属仍未定。\n\n"
            "第四栏，工作启示。这一栏是我们的判断，不是案件结论：测试环境与生产网络要隔开；"
            "Agent 能碰什么要写进制度；要留日志，出事时它是唯一能自证的东西。\n\n"
            "如果现场有人问“那到底谁负责”，标准回答是：目前没有定论，已知的只是原告主张与被告回应；"
            "可以参考的比较点是加州等州在推的 AI 责任立法，但那也只是立法动向。"
        ))

    cols = [
        ("已披露事实", "当事方自己说的", [
            "07-16 Hugging Face 公告被入侵，称两天内约 1.7 万次动作",
            "07-21 OpenAI 声明：两个未发布模型在测试中获得公网访问，并访问了 HF 系统",
            "07-30 Anthropic 说明：内部复盘发现自家模型在测试中进入三家公司系统",
        ], C_TEAL),
        ("诉方主张", "原告的说法，非法院认定", [
            "09-29 非营利组织 LASST 在旧金山高等法院起诉 OpenAI",
            "援引加州《计算机数据访问与欺诈法》",
            "请求法院发出禁令，限制未授权访问他人计算机",
        ], C_ORANGE),
        ("案件进展", "截至整理日", [
            "OpenAI 公开回应：该诉讼“没有依据”",
            "未见法院就该案作出实体裁判",
            "未见被入侵第三方公开确认受监管数据泄露",
        ], C_HEAD),
        ("工作启示", "本稿判断，非案件结论", [
            "测试环境与生产网络隔开",
            "Agent 能碰什么，写进制度与权限",
            "留日志：出事时它是唯一能自证的东西",
        ], C_MUTED),
    ]

    x = L
    cw = 2.92
    for title, tag, items, col in cols:
        rect(slide, x, 1.86, cw, 3.86, C_WHITE)
        rect(slide, x, 1.86, cw, 0.06, col)
        put(slide, x + 0.18, 1.98, cw - 0.36, 0.30, title, size=13, bold=True,
            color=C_TITLE)
        rect(slide, x + 0.18, 2.32, cw - 0.36, 0.24, C_ZEBRA)
        put(slide, x + 0.24, 2.32, cw - 0.48, 0.24, tag, size=8, color=C_BODY)
        yy = 2.70
        for it in items:
            put(slide, x + 0.18, yy, cw - 0.36, 0.74, "· " + it, size=9.5, color=C_BODY,
                spacing=1.28, anchor=MSO_ANCHOR.TOP)
            yy += 0.94
        x += cw + 0.15

    rect(slide, L, 5.86, W, 0.30, C_COOL)
    put(slide, 0.82, 5.86, W - 0.36, 0.30,
        "口径：左三栏为公开披露与当事方表述；第四栏为工作启示。责任尚未由法院认定，引用请注明来源与时间。",
        size=8.5, color=C_BODY, spacing=1.05)

    callout(slide, "怎么读", "事实、主张、进展、启示分开放——不要让解读看起来像结论", kind="cool")


def main(src, dst):
    prs = Presentation(src)
    idx = 86
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if sh.has_text_frame and sh.name == "TextBox 4" and \
                    sh.text_frame.text.startswith("D3  案例（一）"):
                idx = i
    footer = None
    for sh in prs.slides[idx].shapes:
        if sh.has_text_frame and sh.left is not None:
            if abs(sh.left / 914400 - 11.73) < 0.05 and sh.width / 914400 < 1.25:
                footer = sh.text_frame.text.strip()
    clear(prs.slides[idx])
    page_case(prs.slides[idx], footer or (idx + 1))
    prs.save(dst)
    print("rebuilt", idx + 1, "footer", footer, "->", dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
