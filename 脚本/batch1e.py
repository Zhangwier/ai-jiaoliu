import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Pt
from pptx.dml.color import RGBColor
prs = openprs(); s = prs.slides[6]
def g(n): return find(s, name=n)
def style(nm, size, bold, color):
    sh = g(nm)
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name="微软雅黑"
style("Row5_1.83", 17, True, ORANGE)
style("Row5_6.92", 15.5, False, RGBColor(0x2B,0x34,0x45))
style("Row5_16.86", 15.5, True, TEAL)
style("Row5_24.5", 15.5, True, TEAL)
# 行名/单元格高度对齐：文本垂直居中
from pptx.enum.text import MSO_ANCHOR
for nm in ("Row5_1.83","Row5_6.92","Row5_16.86","Row5_24.5"):
    g(nm).text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
save(prs); print("ok")
