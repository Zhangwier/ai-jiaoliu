# -*- coding: utf-8 -*-
from pptx import Presentation
from lxml import etree
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[66]
for sh in s.shapes:
    if sh.name == "TextBox 2":
        print(etree.tostring(sh._element, pretty_print=True).decode()[:2000])
