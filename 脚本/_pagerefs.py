# -*- coding: utf-8 -*-
import re, sys
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
pat = re.compile(r"(P\s?\d+(?:\s?[–\-—~]\s?P?\s?\d+)?|第\s?\d+\s?页)")
for i, s in enumerate(prs.slides, 1):
    hits=[]
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        for m in pat.finditer(t):
            hits.append(m.group(0))
    if hits:
        print(i, len(hits), " | ".join(hits[:14]))
