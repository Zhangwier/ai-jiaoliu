# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Pt
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
s = prs.slides[24]; S={sh.shape_id:sh for sh in s.shapes}
def set_body(sh, txt, size=None):
    tf=sh.text_frame; p0=tf.paragraphs[0]; runs=list(p0.runs)
    if runs:
        runs[0].text=txt
        for r in runs[1:]: r._r.getparent().remove(r._r)
    else: tf.text=txt
    if size:
        for r in tf.paragraphs[0].runs: r.font.size=Pt(size)
set_body(S[34], "彻底顺着你：你说 A 对它就夸 A 对，你说 A 错它立刻改口", size=14.5)
set_body(S[35], "未进本次盲测，也无同口径分数；只按日常使用记录", size=14.5)
prs.save(P); print("fixed", S[34].text_frame.text, "|", S[35].text_frame.text)
