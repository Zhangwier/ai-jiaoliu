# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm, Pt
from pptx.enum.text import MSO_ANCHOR
EMU=360000
prs = openprs(); s = prs.slides[19]

# 1) 删掉“自测”段
for sh in s.shapes:
    if not sh.has_text_frame: continue
    tf = sh.text_frame
    for p in list(tf.paragraphs):
        if p.text.startswith("自测："):
            p._p.getparent().remove(p._p)

# 2) 脚注条文本补自测（窄一点的字号）
foots = {
 "换档位，是最直接的提升": "自测：换个贵一档再跑同一条",
 "按任务挑，比按总分挑划算": "自测：同一任务交给两家，看谁更少返工",
 "壳决定它能干什么，模型决定干得好不好": "自测：换一个能读原文件的壳，任务不变",
}
for sh in s.shapes:
    if not sh.has_text_frame: continue
    t = sh.text_frame.text.strip()
    if t in foots:
        tf = sh.text_frame
        p0 = tf.paragraphs[0]
        r0 = p0.runs[0]
        r0.font.size = Pt(11.5)
        p1 = tf.add_paragraph()
        r = p1.add_run(); r.text = foots[t]
        r.font.size = Pt(10.5); r.font.color.rgb = TEAL; r.font.bold = True; r.font.name="微软雅黑"
        p1.line_spacing = 1.05
        sh.height = Cm(0.95)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE

# 3) 卡片略降，给脚注条更多空间
for sh in s.shapes:
    if abs(sh.top/EMU-4.72) < 0.05 and sh.height/EMU > 9.5:
        sh.height = Cm(9.95)
    if abs(sh.top/EMU-13.02) < 0.06:
        sh.top = Cm(12.98)
save(prs); print("ok")
