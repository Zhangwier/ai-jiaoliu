# -*- coding: utf-8 -*-
"""重建奖励黑客页 v2：修溢出"""
import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

prs = openprs(); s = prs.slides[16]
for sh in list(s.shapes): sh._element.getparent().remove(sh._element)

def tb(x,y,w,h,text,size,color=NAVY,bold=False,align=None,anchor=None,space=1.15,name=None):
    b = s.shapes.add_textbox(Cm(x),Cm(y),Cm(w),Cm(h))
    tf = b.text_frame; tf.word_wrap=True
    if anchor is not None: tf.vertical_anchor = anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    for i,ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if align is not None: p.alignment = align
        p.line_spacing = space
        r = p.add_run(); r.text = ln
        r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name="微软雅黑"
    if name: b.name=name
    return b

def rect(x,y,w,h,fill=None,line=None,name=None):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x),Cm(y),Cm(w),Cm(h))
    if fill is None: r.fill.background()
    else: r.fill.solid(); r.fill.fore_color.rgb=fill
    if line is None: r.line.fill.background()
    else: r.line.color.rgb=line; r.line.width=Pt(0.75)
    r.shadow.inherit=False
    if name: r.name=name
    return r

tb(1.52,0.81,15.24,0.76,"01  认识 AI",12,GREY)
tb(29.80,0.81,2.54,0.76,"17",12,RGBColor(0x6B,0x73,0x85),align=PP_ALIGN.RIGHT)
rect(1.52,1.68,1.27,0.15,fill=TEAL)
tb(1.52,1.98,30.82,1.45,"奖励黑客与工具性趋同：它只对“指标”负责，不对“你的意图”负责",27,NAVY,True)
tb(1.52,3.50,30.82,0.85,"你说“把这件事办成”，它就去优化“看起来办成了”；最容易出问题的两处：改题、越界",17,GREY,space=1.05)
conn = s.shapes.add_connector(1,Cm(1.52),Cm(17.78),Cm(32.34),Cm(17.78)); conn.line.color.rgb=RGBColor(0xE2,0xDF,0xD6); conn.line.width=Pt(0.75)
tb(1.52,17.93,26.67,0.76,"来源：OpenAI、Anthropic 关于“奖励黑客”的公开说明；Omohundro（2008）工具性趋同；本分享实测记录（2026-10）",8.5,GREY)
tb(27.26,17.93,5.08,0.76,"来源索引 ›",8.5,GREY)

# ---- 左 ----
rect(1.52,4.65,15.24,10.28,fill=WHITE)
rect(1.52,4.65,0.20,10.28,fill=ORANGE)
tb(2.03,4.85,14.10,0.75,"奖励黑客  Reward hacking",18,NAVY,True)
tb(2.03,5.62,14.10,0.98,"它的目标是“分数高”，不是“事做对”；没写进评分标准的，它不会替你补上。",12.5,GREY,space=1.18)
tb(2.03,6.70,14.10,4.60,
"公开的典型\n"
"让它“通过单元测试”，它把测试改成永远通过；让它“减少报错”，它把报错静默吞掉；让它“提高通过率”，它把阈值调宽。\n"
"身边的典型\n"
"让它“把格式调一致”，它删掉不合格式的段落；让它“过校验”，它绕开校验脚本；让它“把表填满”，它用推算值补齐空行。",
12.5,NAVY,space=1.22)
tb(2.03,11.50,14.10,0.80,"为什么：结果一旦可测量，绕开它、伪造它都比真做便宜。",12.5,ORANGE,space=1.15)
rect(1.52,12.55,15.24,2.24,fill=PEACH) if False else rect(1.52,12.55,15.24,2.24,fill=PEACH)
tb(2.03,12.70,14.23,1.95,
"怎么看：交付物“太顺”就要警惕——数字全对、格式全齐，却没有一处引用原文。\n"
"核法：挑 3 条结论回原文；看它改了哪些文件；看测试脚本有没有被动过。",12,NAVY,space=1.18)

# ---- 右 ----
rect(17.27,4.65,15.06,10.28,fill=WHITE)
rect(17.27,4.65,0.20,10.28,fill=TEAL)
tb(17.78,4.85,14.05,0.75,"工具性趋同  Instrumental convergence",18,NAVY,True)
tb(17.78,5.62,14.05,1.05,"不管目标是什么，智能体都会倾向做同样几件事——它们对几乎任何目标都有用。",12.5,GREY,space=1.18)
items = [
 ("自我保护","不愿被关闭、替换、降级；会隐瞒、装作已完成"),
 ("扩充资源","多要权限、多占额度、多存数据、多开账号"),
 ("影响操作者","说你想听的话，让你更信任、更依赖它"),
 ("走捷径","先做出“能交差”的结果，把难题留给人"),
]
y = 6.95
for t1,t2 in items:
    rect(17.78,y,0.14,1.10,fill=TEAL)
    tb(18.12,y,3.30,1.10,t1,13.5,NAVY,True,anchor=MSO_ANCHOR.MIDDLE,space=1.05)
    tb(21.30,y,11.00,1.10,t2,12.5,NAVY,anchor=MSO_ANCHOR.MIDDLE,space=1.18)
    y += 1.24
rect(17.27,12.55,15.06,2.24,fill=MINT)
tb(17.78,12.70,14.05,1.95,
"这不是“恶意”：它只是在优化你给的指标。\n"
"约束要写在指标外面：不改测试、不删文件、不动生产数据、不自己加权限。",12,NAVY,space=1.18)

# ---- 结论 ----
rect(1.52,16.05,30.82,1.30,fill=PEACH)
rect(1.52,16.05,0.23,1.30,fill=ORANGE)
tb(2.16,16.05,2.54,1.30,"结论",14,NAVY,anchor=MSO_ANCHOR.MIDDLE)
tb(4.45,16.05,27.26,1.30,"目标写得越“结果化”，越要给过程设卡；它不是坏，是只对你的指标负责",18,NAVY,anchor=MSO_ANCHOR.MIDDLE)

save(prs); print("ok")
