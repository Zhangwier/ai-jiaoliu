# -*- coding: utf-8 -*-
"""Muse 页右栏完全重建"""
import sys
sys.path.insert(0, r"D:\MG\program\ai交流\脚本")
from _lib import *
from pptx.util import Cm, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
EMU=360000
prs = openprs(); s = prs.slides[68]
# 删掉右栏全部（left>=9.0，t 在 4.5~16.1），保留结论条与页脚
for sh in list(s.shapes):
    l = sh.left/EMU; t = sh.top/EMU
    if l >= 9.0 and 4.5 < t < 16.05:
        sh._element.getparent().remove(sh._element)

def tb(x,y,w,h,text,size,color=NAVY,bold=False,space=1.28):
    b=s.shapes.add_textbox(Cm(x),Cm(y),Cm(w),Cm(h))
    tf=b.text_frame; tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    for i,ln in enumerate(text.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.line_spacing=space
        r=p.add_run(); r.text=ln
        r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name="微软雅黑"
    return b

rows = [
 ("它是什么", "Meta Superintelligence Labs 的第一款前沿模型（Muse Wonder / Wonder Full，2026-10-12）。AA 智能指数 52；LMArena 文字榜曾列第二。前身 Muse Spark（2026-04）现在仍在支撑 Meta AI。", 2.10),
 ("两个招牌能力", "Active Memory：跨对话记住你的偏好、行程、家人信息，下次直接接着用。\nActive Focus：长任务不跑偏——订多程行程、反复改方案时不丢前面的条件。", 1.95),
 ("入口与形态", "Muse App（muse.ai，网页＋手机）；Muse Code（订阅者的编程入口）；Meta Model API；Facebook / Instagram / WhatsApp 里的 Meta AI。", 1.80),
 ("短板与边界", "只对美国、加拿大开放（第三方核查口径）；免费档有用量上限，Power $20、Maximum $100。\n企业级权限与审计弱于 OpenAI、Anthropic；编程不是它的主场。", 1.95),
 ("什么时候值得试", "要一个“记得住你”的私人助理：行程、比价、定制、排座位、反复修改的方案。\n要交付工程文件、要接公司系统，还是用别的。", 1.85),
]
y = 4.83
for i,(k,v,rh) in enumerate(rows):
    if i > 0:
        c=s.shapes.add_connector(1,Cm(9.19),Cm(y),Cm(32.33),Cm(y))
        c.line.color.rgb=RGBColor(0xE2,0xDF,0xD6); c.line.width=Pt(0.75)
    tb(9.19, y+0.06, 2.79, 1.30, k, 11.5, TEAL, True)
    tb(12.14, y, 20.19, rh-0.05, v, 11.5, NAVY)
    y += rh + 0.10
save(prs); print("ok", y)
