# -*- coding: utf-8 -*-
from pptx import Presentation
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
prs = Presentation(P)
for i,s in enumerate(prs.slides,1):
    if i>119: break
    best=""
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        sz=0
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size: sz=max(sz,r.font.size.pt)
        if sz>=22 and 1.5 < sh.top/360000 < 4.6:
            t=sh.text_frame.text.replace("\n"," / ")
            if t: best=t; break
    print(f"{i}\t{best}")
