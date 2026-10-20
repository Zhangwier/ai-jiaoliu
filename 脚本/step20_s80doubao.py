# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
P = r"D:\MG\program\ai交流\成品\_work29\base.pptx"
DIRS = [r"D:\MG\program\ai交流\成品\_work29\logos3", r"D:\MG\program\ai交流\成品\_work29\logos4", r"D:\MG\program\ai交流\成品\_work29\logos2", r"D:\MG\program\ai交流\成品\_work29\logos"]
def pick(n):
    for d in DIRS:
        p=os.path.join(d,n+".png")
        if os.path.exists(p): return p
prs = Presentation(P)
s = prs.slides[79]
# 删除该行已加的 badge/logo（按位置判定）
kill=[]
for sh in s.shapes:
    if sh.name in ("LogoBadge","Logo"):
        if 1.7 <= sh.left/360000 <= 3.2 and 10.3 <= sh.top/360000 <= 11.9:
            kill.append(sh)
for sh in kill:
    sh._element.getparent().remove(sh._element)
print("removed", len(kill))
SIGHT = 0.50
def badge(sl, pth, x, top):
    ov = sl.shapes.add_shape(MSO_SHAPE.OVAL, Cm(x), Cm(top), Cm(SIGHT), Cm(SIGHT))
    ov.fill.solid(); ov.fill.fore_color.rgb = RGBColor(0xFF,0xFF,0xFF)
    ov.line.color.rgb = RGBColor(0xE2,0xDF,0xD6); ov.line.width = Pt(0.75)
    ov.shadow.inherit=False; ov.name="LogoBadge"
    pad=SIGHT*0.17
    sl.shapes.add_picture(pth, Cm(x+pad), Cm(top+pad), Cm(SIGHT-2*pad), Cm(SIGHT-2*pad)).name="Logo"
S = {sh.shape_id: sh for sh in s.shapes}
sh = S[32]
top = (sh.top + (sh.height - Cm(SIGHT))/2)/360000
x=1.84
for k in ("doubao","qwen"):
    badge(s, pick(k), x, top); x += SIGHT+0.05
sh.left = Cm(x+0.04); sh.width = Cm(5.51-(x+0.04))
sh.text_frame.text = "豆包*/千问*"
for p in sh.text_frame.paragraphs:
    p.alignment = None
    for r in p.runs: r.font.size = Pt(10)
prs.save(P)
print("row left", round(sh.left/360000,2), "w", round(sh.width/360000,2))
